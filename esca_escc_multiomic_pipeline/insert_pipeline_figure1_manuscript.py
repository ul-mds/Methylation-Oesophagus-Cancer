"""
Insert Plots/00_analysis_pipeline_flowchart.png as Fig. 1 in a manuscript .docx,
renumber existing figures (1→2 … 4→5), cite Fig. 1 in Methods, and audit citations.
"""
from __future__ import annotations

import argparse
import re
import shutil
import zipfile
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph

STUDY_ROOT = Path(__file__).resolve().parents[1]
W_T = qn("w:t")
W_R = qn("w:r")
DEFAULT_DOCX = STUDY_ROOT / "Manuscript.docx"
PLOT_DIR = STUDY_ROOT / "Plots"

REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
IMAGE_REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"

FIG1_CAP = (
    "Fig. 1. Analysis pipeline for the TCGA ESCA (Firehose Legacy) cohort. "
    "Flowchart of data sources (A), pre-processing and OS60 landmark endpoints (B), "
    "penalized Cox modelling and external GEO checks (C), fast-death methylation and "
    "survival analyses (D), and exploratory somatic and mediation layers (E). "
    "Arrows indicate the primary analysis sequence; dashed arrows mark key cross-block links."
)

# After renumbering: Fig. N → PNG
FIGURE_SPECS: list[tuple[str, str, float]] = [
    ("Fig. 1.", "00_analysis_pipeline_flowchart.png", 6.4),
    ("Fig. 2.", "01_cross_apply_cindex.png", 6.2),
    ("Fig. 3.", "02_fastdeath_methylation_volcano_panel.png", 6.6),
    ("Fig. 4.", "04_KM_methylation_60mo_panel.png", 6.6),
    ("Fig. 5.", "05_somatic_combined_A_heatmap_B_interaction.png", 6.2),
]

# Longer / compound patterns first when bumping figure numbers (+1)
FIG_BUMP_ORDER = [
    ("Fig. 4A–B", "Fig. 5A–B"),
    ("Fig. 4B", "Fig. 5B"),
    ("Fig. 4A", "Fig. 5A"),
    ("Fig. 4", "Fig. 5"),
    ("Fig. 3d–f", "Fig. 4d–f"),
    ("Fig. 3a–c", "Fig. 4a–c"),
    ("Fig. 3", "Fig. 4"),
    ("Fig. 2A–B", "Fig. 3A–B"),
    ("Fig. 2B", "Fig. 3B"),
    ("Fig. 2A", "Fig. 3A"),
    ("Fig. 2", "Fig. 3"),
    ("Fig. 1", "Fig. 2"),
]


def _insert_paragraph_before(paragraph: Paragraph, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    return new_para


def _insert_paragraph_after(paragraph: Paragraph, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    return new_para


def _p_has_image(p: Paragraph) -> bool:
    return bool(p._p.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"))


def _ensure_image_before_caption(caption_para: Paragraph, png: Path, width: Inches) -> None:
    """Place PNG in the paragraph immediately above the caption (this manuscript's layout)."""
    cap_el = caption_para._p
    prev = cap_el.getprevious()
    if prev is not None and prev.tag.endswith("}p"):
        img_para = Paragraph(prev, caption_para._parent)
        for r in list(img_para.runs):
            r._r.getparent().remove(r._r)
        img_para.add_run().add_picture(str(png), width=width)
        return
    pic_p = _insert_paragraph_before(caption_para, "")
    pic_p.add_run().add_picture(str(png), width=width)


def _find_caption(doc: Document, prefix: str) -> Paragraph:
    for p in doc.paragraphs:
        if (p.text or "").strip().startswith(prefix):
            return p
    raise ValueError(f"Missing caption: {prefix!r}")


def _apply_text_fn_to_document(doc: Document, text_fn) -> None:
    """Replace text in every w:t node, preserving run properties (highlights, bold, etc.)."""
    for t_el in doc.element.body.iter(W_T):
        if t_el.text:
            new = text_fn(t_el.text)
            if new != t_el.text:
                t_el.text = new


def _prepend_run(paragraph: Paragraph, text: str) -> None:
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    paragraph._p.insert(0, run)


def _replace_in_paragraph(paragraph: Paragraph, old: str, new: str) -> bool:
    """Replace old→new in w:t nodes; preserves run properties (highlights)."""
    for t_el in paragraph._p.iter(W_T):
        if t_el.text and old in t_el.text:
            t_el.text = t_el.text.replace(old, new, 1)
            return True
    nodes = list(paragraph._p.iter(W_T))
    full = "".join(n.text or "" for n in nodes)
    if old not in full:
        return False
    new_full = full.replace(old, new, 1)
    if nodes:
        nodes[0].text = new_full
        for n in nodes[1:]:
            n.text = ""
    return True


def _set_paragraph_text_preserve_first_run_rpr(paragraph: Paragraph, text: str) -> None:
    nodes = list(paragraph._p.iter(W_T))
    if nodes:
        nodes[0].text = text
        for n in nodes[1:]:
            n.text = ""


def _already_has_pipeline_fig1(doc: Document) -> bool:
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith("Fig. 1.") and "Analysis pipeline" in t:
            return True
    return False


def _bump_figure_refs(text: str) -> str:
    placeholders: dict[str, str] = {}

    def protect(m: re.Match) -> str:
        key = f"@@PROT{len(placeholders)}@@"
        placeholders[key] = m.group(0)
        return key

    for pat in (
        r"Supplementary Fig\.?\s*S\d+(?:[–\-—]\s*S\d+)?",
        r"Supplementary Figs\.?\s*S\d+(?:[–\-—]\s*S\d+)?",
        r"Fig\.?\s*images",
        r"Fig\.?\s*legends",
        r"combined Fig\.s",
    ):
        text = re.sub(pat, protect, text, flags=re.IGNORECASE)
    for old, new in FIG_BUMP_ORDER:
        text = text.replace(old, new)
    for key, val in placeholders.items():
        text = text.replace(key, val)
    return text


def _prune_orphan_image_parts(docx_path: Path) -> None:
    with zipfile.ZipFile(docx_path, "r") as zin:
        doc_xml = zin.read("word/document.xml").decode("utf-8")
        used_rids = set(re.findall(r'r:embed="(rId\d+)"', doc_xml))
        rels_data = zin.read("word/_rels/document.xml.rels")
        names = zin.namelist()

    root = ET.fromstring(rels_data)
    remove_rels = []
    keep_media: set[str] = set()
    for rel in root.findall(f"{{{REL_NS}}}Relationship"):
        if rel.get("Type") != IMAGE_REL_TYPE:
            continue
        rid = rel.get("Id")
        tgt = rel.get("Target", "").replace("\\", "/")
        if rid in used_rids:
            keep_media.add(tgt)
        else:
            remove_rels.append(rel)
    for rel in remove_rels:
        root.remove(rel)
    new_rels = ET.tostring(root, encoding="utf-8", xml_declaration=True)

    buf = BytesIO()
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            name = item.filename
            if name == "word/_rels/document.xml.rels":
                zout.writestr(item, new_rels)
                continue
            if name.startswith("word/media/") and name.endswith(".png"):
                rel_path = name.removeprefix("word/")
                if rel_path not in keep_media:
                    continue
            zout.writestr(item, zin.read(name))
    docx_path.write_bytes(buf.getvalue())


def _audit(doc: Document) -> list[str]:
    issues: list[str] = []
    fig_caps = set()
    tbl_caps = set()
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        m = re.match(r"^Fig\.?\s*(\d+)", t)
        if m:
            fig_caps.add(int(m.group(1)))
        m = re.match(r"^Table\s+(\d+)", t)
        if m:
            tbl_caps.add(int(m.group(1)))

    for i, p in enumerate(doc.paragraphs):
        t = p.text or ""
        if re.match(r"^Fig\.?\s+\d", t.strip()):
            continue
        for m in re.finditer(r"(?<!Supplementary )Fig\.?\s*(\d+)", t):
            n = int(m.group(1))
            if n not in fig_caps:
                issues.append(f"para {i}: cites Fig. {n} but captions are {sorted(fig_caps)}")
        for m in re.finditer(r"(?<!Supplementary )Table\s+(\d+)", t):
            n = int(m.group(1))
            if n not in tbl_caps and not t.strip().startswith("Table"):
                issues.append(f"para {i}: cites Table {n} but captions are {sorted(tbl_caps)}")

    methods_has_fig1 = False
    in_methods = False
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t == "Methods":
            in_methods = True
            continue
        if in_methods and t == "References":
            break
        if in_methods and re.search(r"Fig\.?\s*1\b", t):
            methods_has_fig1 = True
    if not methods_has_fig1:
        issues.append("Methods section does not cite Fig. 1")

    return issues


def _audit_references(doc: Document) -> list[str]:
    issues: list[str] = []
    ref_nums: set[int] = set()
    in_refs = False
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t == "References":
            in_refs = True
            continue
        if in_refs:
            if t.startswith("Acknowledgements") or t.startswith("Author contributions"):
                break
            m = re.match(r"^(\d+)\.", t)
            if m:
                ref_nums.add(int(m.group(1)))

    cited: set[int] = set()
    for p in doc.paragraphs:
        for m in re.finditer(r"\[(\d+)\]", p.text or ""):
            cited.add(int(m.group(1)))

    for n in sorted(cited):
        if n not in ref_nums:
            issues.append(f"Reference [{n}] cited in text but missing from References list")

    return issues


def patch_manuscript(docx_path: Path | None = None) -> Path:
    DOCX_PATH = docx_path or DEFAULT_DOCX
    for _, stem, _ in FIGURE_SPECS:
        png = PLOT_DIR / stem
        if not png.is_file():
            raise FileNotFoundError(f"Missing plot: {png}")

    backup = DOCX_PATH.with_name(DOCX_PATH.stem + ".docx.bak")
    shutil.copy2(DOCX_PATH, backup)

    doc = Document(str(DOCX_PATH))

    if _already_has_pipeline_fig1(doc):
        print(f"Already contains pipeline Fig. 1 — re-embedding images only.")
    else:
        def _patch_text(text: str) -> str:
            text = text.replace("nominal p in Table 7", "nominal p in Supplementary Table S7")
            return _bump_figure_refs(text)

        _apply_text_fn_to_document(doc, _patch_text)

        # Insert pipeline figure before current Fig. 2 block (image then caption pattern)
        cap2 = _find_caption(doc, "Fig. 2.")
        # Image paragraph is immediately before caption in this manuscript
        img_anchor = cap2
        prev = cap2._p.getprevious()
        if prev is not None and prev.tag.endswith("}p"):
            img_anchor = Paragraph(prev, cap2._parent)

        new_img = _insert_paragraph_before(img_anchor, "")
        new_img.add_run().add_picture(str(PLOT_DIR / "00_analysis_pipeline_flowchart.png"), width=Inches(6.4))
        _insert_paragraph_after(new_img, FIG1_CAP)

        # Methods: cite Fig. 1 at cohort assembly (prepend; preserve yellow highlights in body)
        methods_lead = (
            "The end-to-end analytical workflow—from cohort assembly through survival modelling, "
            "fast-death methylation discovery, and exploratory somatic layers—is summarized in Fig. 1. "
        )
        for p in doc.paragraphs:
            if (p.text or "").startswith("Primary tumor data were retrieved from the TCGA"):
                if "Fig. 1" not in (p.text or ""):
                    _prepend_run(p, methods_lead)
                break

        # Results opening: name pipeline overview as Fig. 1
        for p in doc.paragraphs:
            t = p.text or ""
            if t.startswith("We interrogated the TCGA ESCA") and "multi-omics pipeline (Fig. 1)" not in t:
                _replace_in_paragraph(
                    p,
                    "through a systematic multi-omics pipeline, structuring",
                    "through a systematic multi-omics pipeline (Fig. 1), structuring",
                )
                break

        # Update display-items note (preserve any paragraph formatting)
        display_note = (
            "Scientific Reports allows at most eight combined Fig.s and tables in the main article. "
            "This manuscript currently includes five Fig.s and four tables (nine display items). "
            "Before submission, move at least one item to Supplementary Information or merge panels, "
            "and cite each in the main text using the word Supplementary "
            "(e.g. Supplementary Table S1; Supplementary Fig. S1)."
        )
        for p in doc.paragraphs:
            if "Scientific Reports allows at most eight combined" in (p.text or ""):
                _set_paragraph_text_preserve_first_run_rpr(p, display_note)
                break

        doc.save(str(DOCX_PATH))

    # Re-embed all figure images in paragraphs above captions
    doc = Document(str(DOCX_PATH))
    for prefix, stem, w in FIGURE_SPECS:
        cap = _find_caption(doc, prefix)
        _ensure_image_before_caption(cap, PLOT_DIR / stem, Inches(w))
    doc.save(str(DOCX_PATH))
    _prune_orphan_image_parts(DOCX_PATH)

    issues = _audit(Document(str(DOCX_PATH)))
    issues.extend(_audit_references(Document(str(DOCX_PATH))))
    print(f"Saved: {DOCX_PATH}")
    print(f"Backup: {backup}")
    if issues:
        print("Audit issues:")
        for x in issues:
            print(" ", x)
    else:
        print("Audit: all figure/table citations OK; Methods cites Fig. 1; references OK.")
    return DOCX_PATH


def main() -> None:
    ap = argparse.ArgumentParser(description="Insert pipeline flowchart as Fig. 1")
    ap.add_argument(
        "--docx",
        type=Path,
        default=DEFAULT_DOCX,
        help="Target manuscript .docx (default: Manuscript.docx)",
    )
    args = ap.parse_args()
    patch_manuscript(args.docx.resolve())


if __name__ == "__main__":
    main()
