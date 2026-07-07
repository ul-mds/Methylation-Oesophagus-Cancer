"""
Move main-text Table 1 to Supplementary Table S1:
  - Remove Table 1 caption + embedded Word table from Results
  - Attach table grid after Supplementary Table S1 caption (if not already present)
  - Update all in-text 'Table 1' → 'Supplementary Table S1'
  - Update REMARK checklist cells
  - Set display-items note to 5 figures + 3 tables = 8
"""
from __future__ import annotations

import argparse
import copy
import re
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

STUDY_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DOCX = STUDY_ROOT / "Manuscript First.docx"
W_T = qn("w:t")
W_TBL = qn("w:tbl")
W_P = qn("w:p")


def _ptext(el) -> str:
    return "".join(t.text or "" for t in el.iter(W_T)).strip()


def _replace_in_paragraph(p: Paragraph, replacements: list[tuple[str, str]]) -> None:
    for t_el in p._p.iter(W_T):
        if not t_el.text:
            continue
        new = t_el.text
        for old, rep in replacements:
            if old in new:
                new = new.replace(old, rep)
        if new != t_el.text:
            t_el.text = new


def _replace_in_all_paragraphs(doc: Document, replacements: list[tuple[str, str]]) -> None:
    for p in doc.paragraphs:
        _replace_in_paragraph(p, replacements)


def _replace_in_table_cells(doc: Document, replacements: list[tuple[str, str]]) -> None:
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace_in_paragraph(p, replacements)


def _find_table1_block(body_children: list) -> tuple[int, object | None, object | None]:
    """Return (caption_idx, caption_el, tbl_el) for main Table 1."""
    cap_el = tbl_el = None
    cap_idx = -1
    for idx, el in enumerate(body_children):
        if el.tag != W_P:
            continue
        t = _ptext(el)
        if t.startswith("Table 1."):
            cap_el = el
            cap_idx = idx
            if idx + 1 < len(body_children) and body_children[idx + 1].tag == W_TBL:
                tbl_el = body_children[idx + 1]
            break
    return cap_idx, cap_el, tbl_el


def _s1_has_following_table(s1_cap_el) -> bool:
    nxt = s1_cap_el.getnext()
    return nxt is not None and nxt.tag == W_TBL


def move_table1(docx_path: Path) -> None:
    backup = docx_path.with_name(docx_path.stem + ".pre_table1_move.docx.bak")
    shutil.copy2(docx_path, backup)

    doc = Document(str(docx_path))
    body = doc.element.body
    children = [c for c in body if not c.tag.endswith("}sectPr")]

    cap_idx, cap_el, tbl_el = _find_table1_block(children)
    if cap_el is None:
        raise RuntimeError("Main-text Table 1 caption not found.")

    # Detach Table 1 caption and table from body
    cap_el.getparent().remove(cap_el)
    if tbl_el is not None:
        tbl_el.getparent().remove(tbl_el)

    # Attach table data after Supplementary Table S1 caption
    s1_cap = None
    for p in doc.paragraphs:
        if (p.text or "").strip().startswith("Supplementary Table S1."):
            s1_cap = p._p
            break
    if s1_cap is None:
        raise RuntimeError("Supplementary Table S1 caption not found.")
    if tbl_el is not None and not _s1_has_following_table(s1_cap):
        s1_cap.addnext(copy.deepcopy(tbl_el))

    # Text replacements (order matters: specific phrases first)
    reps = [
        ("provided in Table 1 and Supplementary Table S1", "provided in Supplementary Table S1"),
        ("(Fig. 2; Table 1)", "(Fig. 2; Supplementary Table S1)"),
        ("Fig. 2; Table 1", "Fig. 2; Supplementary Table S1"),
        ("Table 1 and Supplementary Table S1", "Supplementary Table S1"),
        ("Table 1 / S1", "Supplementary Table S1"),
        ("Table 1 / Fig. 2", "Supplementary Table S1 / Fig. 2"),
        ("Table 1 — within / cross / pooled", "Supplementary Table S1 — within / cross / pooled"),
        ("Results; Table 1 / S1", "Results; Supplementary Table S1"),
        ("C-index; Table 1 / Fig. 2", "C-index; Supplementary Table S1 / Fig. 2"),
    ]
    _replace_in_all_paragraphs(doc, reps)
    _replace_in_table_cells(doc, reps)

    # Catch remaining standalone Table 1 (not Table 10-19)
    for p in doc.paragraphs:
        t = p.text or ""
        if re.search(r"(?<!Supplementary )Table 1\b", t) and not t.strip().startswith("Supplementary"):
            _replace_in_paragraph(p, [("Table 1", "Supplementary Table S1")])

    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                if re.search(r"Table 1\b", cell.text) and "Supplementary Table S1" not in cell.text:
                    for p in cell.paragraphs:
                        _replace_in_paragraph(p, [("Table 1", "Supplementary Table S1")])

    # Display-items note
    note = (
        "Scientific Reports allows at most eight combined figures and tables in the main article. "
        "This manuscript currently includes five figures and three tables (eight display items). "
        "Concordance-index summaries formerly in main-text Table 1 appear as Supplementary Table S1. "
        "Before submission, cite each supplementary item in the main text using the word Supplementary "
        "(e.g. Supplementary Table S1; Supplementary Fig. S1)."
    )
    for p in doc.paragraphs:
        if "Scientific Reports allows at most eight combined" in (p.text or ""):
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
            p.add_run(note)
            break

    doc.save(str(docx_path))
    print(f"Saved: {docx_path}")
    print(f"Backup: {backup}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--docx", type=Path, default=DEFAULT_DOCX)
    args = ap.parse_args()
    move_table1(args.docx.resolve())


if __name__ == "__main__":
    main()
