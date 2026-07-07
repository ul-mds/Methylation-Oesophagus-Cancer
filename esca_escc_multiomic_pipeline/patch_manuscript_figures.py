"""
Rebuild the Figures section of the landmark60 manuscript: combined panels + renumbering.

Requires: Plots/01_cross_apply_cindex.png … 06_interaction_screen.png
Run after: make_plots.main() and make_methylation_os_plots.main()
"""
from __future__ import annotations

import re
import zipfile
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph


def _insert_paragraph_after(paragraph: Paragraph, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    return new_para


def _remove_paragraphs_between(doc: Document, start_idx: int, end_idx_exclusive: int) -> None:
    """Remove paragraph elements by index (inclusive start, exclusive end)."""
    for i in range(end_idx_exclusive - 1, start_idx - 1, -1):
        p = doc.paragraphs[i]
        el = p._element
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)


REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
IMAGE_REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"


def _prune_orphan_image_parts(docx_path: Path) -> None:
    """Remove unused image relationships and media/*.png after python-docx leaves orphans."""
    with zipfile.ZipFile(docx_path, "r") as zin:
        names = zin.namelist()
        doc_xml = zin.read("word/document.xml").decode("utf-8")
        used_rids = set(re.findall(r'r:embed="(rId\d+)"', doc_xml))
        rels_data = zin.read("word/_rels/document.xml.rels")

    root = ET.fromstring(rels_data)
    remove_rels = []
    keep_media_targets: set[str] = set()
    for rel in root.findall(f"{{{REL_NS}}}Relationship"):
        if rel.get("Type") != IMAGE_REL_TYPE:
            continue
        rid = rel.get("Id")
        tgt = rel.get("Target", "")
        if rid in used_rids:
            keep_media_targets.add(tgt.replace("\\", "/"))
        else:
            remove_rels.append(rel)
    for rel in remove_rels:
        root.remove(rel)

    new_rels = ET.tostring(root, encoding="utf-8", xml_declaration=True)

    buf = BytesIO()
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(
        buf, "w", compression=zipfile.ZIP_DEFLATED
    ) as zout:
        for item in zin.infolist():
            name = item.filename
            if name == "word/_rels/document.xml.rels":
                zout.writestr(item, new_rels)
                continue
            if name.startswith("word/media/") and name.endswith(".png"):
                rel_path = name.removeprefix("word/")
                if rel_path not in keep_media_targets:
                    continue
            data = zin.read(name)
            zout.writestr(item, data)

    docx_path.write_bytes(buf.getvalue())


def patch_manuscript(study_root: Path | None = None) -> Path:
    study_root = study_root or Path(__file__).resolve().parents[1]
    docx_path = study_root / "TCGA_ESCA_ESCC_landmark60_full_scientific_manuscript.docx"
    plot_dir = study_root / "Plots"

    figures = [
        (
            "Figure 1. Cross-apply performance (C-index) using 60-month landmark OS. "
            "Histology-specific models discriminate within histology but do not transfer across histologies.",
            plot_dir / "01_cross_apply_cindex.png",
            Inches(6.2),
        ),
        (
            "Figure 2. Fast-death differential methylation by histology (volcano). "
            "Left: ESCA/EAC; right: ESCC. Each point is a locus with Δβ = mean(β_fast)−mean(β_other) "
            "and y-axis = −log10(q_diff).",
            plot_dir / "02_fastdeath_methylation_volcano_panel.png",
            Inches(6.6),
        ),
        (
            "Figure 3. Top fast-death hypo- and hypermethylated loci (Δβ) in ESCA/EAC (left) and ESCC (right).",
            plot_dir / "03_fastdeath_top_hypo_hyper_panel.png",
            Inches(6.6),
        ),
        (
            "Figure 4. Kaplan–Meier curves (60-month landmark OS) stratified by median methylation. "
            "Upper row: ESCA/EAC, three probes with strongest methylation–OS signal (univariate Cox on OS60). "
            "Lower row: ESCC, three probes with strongest fast-death differential methylation (q_diff).",
            plot_dir / "04_KM_methylation_60mo_panel.png",
            Inches(6.6),
        ),
        (
            "Figure 5. Exploratory mechanistic follow-up: mutation-associated shifts in expression-derived "
            "module scores (Δ=mut−WT) stratified by histology.",
            plot_dir / "05_mutation_module_heatmap.png",
            Inches(6.0),
        ),
        (
            "Figure 6. Mutation×histology interaction screen on 60-month landmark OS (nominal p-values).",
            plot_dir / "06_interaction_screen.png",
            Inches(6.0),
        ),
    ]

    for _, png, _ in figures:
        if not png.is_file():
            raise FileNotFoundError(f"Missing plot (run plotting first): {png}")

    doc = Document(str(docx_path))

    # --- Results text: figure cross-refs ---
    replacements = [
        (
            "In EAC, fast-death was characterized predominantly by hypomethylation across top-ranked loci (Figure 2A), with leading candidates including SIGLEC14, EPHB1, ZNF214, and RAB11FIP3 (Table 2A). In ESCC, the fast-death signal was anchored by a particularly strong hypomethylation at SHISA3 (Figure 2B; Table 2B), alongside additional loci such as C7orf13, PRKG1, and SEMA3E. Bar summaries of the top hypo- and hypermethylated loci highlight the different directionality and composition of fast-death programs across histologies (Figure 2C–D).",
            "In EAC, fast-death was characterized predominantly by hypomethylation across top-ranked loci (Figure 2, left), with leading candidates including SIGLEC14, EPHB1, ZNF214, and RAB11FIP3 (Table 2A). In ESCC, the fast-death signal was anchored by a particularly strong hypomethylation at SHISA3 (Figure 2, right; Table 2B), alongside additional loci such as C7orf13, PRKG1, and SEMA3E. Bar summaries of the top hypo- and hypermethylated loci highlight the different directionality and composition of fast-death programs across histologies (Figure 3).",
        ),
        (
            "When dichotomized by median methylation, these loci stratified OS60 in Kaplan–Meier analyses (Figure 3A–C). For ESCC, we visualized 60-month Kaplan–Meier curves for leading fast-death loci, including SHISA3, PRKG1, and C7orf13 (Figure 3D–F), providing an interpretable mapping between fast-death differential methylation and time-to-event differences within a bounded clinical horizon.",
            "When dichotomized by median methylation, these loci stratified OS60 in Kaplan–Meier analyses (Figure 4, upper row). For ESCC, we visualized 60-month Kaplan–Meier curves for leading fast-death loci, including SHISA3, PRKG1, and C7orf13 (Figure 4, lower row), providing an interpretable mapping between fast-death differential methylation and time-to-event differences within a bounded clinical horizon.",
        ),
        (
            "For example, several mutations showed associations with EMT or immune/IFN-γ module shifts in a histology-dependent manner (Figure 4A; Table 4). The interaction screen and module follow-up are presented as hypothesis-generating evidence for downstream biological context rather than definitive interaction discoveries (Figure 4B).",
            "For example, several mutations showed associations with EMT or immune/IFN-γ module shifts in a histology-dependent manner (Figure 5; Table 4). The interaction screen and module follow-up are presented as hypothesis-generating evidence for downstream biological context rather than definitive interaction discoveries (Figure 6).",
        ),
    ]
    for old, new in replacements:
        for p in doc.paragraphs:
            if old in p.text:
                p.text = p.text.replace(old, new)

    for p in doc.paragraphs:
        if "Figure 4A" in p.text or "Figure 4B" in p.text:
            p.text = p.text.replace("Figure 4A", "Figure 5").replace("Figure 4B", "Figure 6")

    # Locate Figures section
    fig_idx = None
    disc_idx = None
    for i, p in enumerate(doc.paragraphs):
        t = (p.text or "").strip()
        if t == "Figures":
            fig_idx = i
        if fig_idx is not None and t == "Discussion":
            disc_idx = i
            break
    if fig_idx is None or disc_idx is None:
        raise RuntimeError("Could not find 'Figures' / 'Discussion' headings in manuscript.")

    already_patched = any(
        (p.text or "").startswith("Figure 2. Fast-death differential methylation by histology")
        for p in doc.paragraphs
    )
    if not already_patched:
        _remove_paragraphs_between(doc, fig_idx + 1, disc_idx)
        anchor = doc.paragraphs[fig_idx]
        cur = anchor
        for cap, png, width in figures:
            cur = _insert_paragraph_after(cur, cap)
            pic_p = _insert_paragraph_after(cur, "")
            pic_p.add_run().add_picture(str(png), width=width)
            cur = pic_p

    out_path = docx_path
    doc.save(str(out_path))
    _prune_orphan_image_parts(out_path)
    return out_path


def main():
    p = patch_manuscript()
    print("Saved:", p)


if __name__ == "__main__":
    main()
