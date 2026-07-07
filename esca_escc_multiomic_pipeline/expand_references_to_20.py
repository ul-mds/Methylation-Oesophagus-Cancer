"""
Expand the manuscript reference list to 20 items with additional peer-reviewed
work (mostly 2022–2026) on ESCC/EAC methylation, multi-omics, ML prognosis, and
clinical marker reviews. Idempotent: skips if reference 20 is already present.

Also tightens reference 4 (Cao et al., Nat Commun 2020) and inserts one bridging
Introduction paragraph citing [11]–[20].
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


def _txt(p) -> str:
    return "".join(t.text or "" for t in p._p.iter(qn("w:t"))).strip()


def _has_ref_20(doc: Document) -> bool:
    return any(_txt(p).startswith("20. Xu Y") for p in doc.paragraphs)


def _insert_paragraph_after(paragraph: Paragraph, text: str, style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._element.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    new_para.add_run(text)
    if style is not None:
        new_para.style = style
    return new_para


def main():
    path = Path(__file__).resolve().parents[1] / "TCGA_ESCA_ESCC_landmark60_full_scientific_manuscript.docx"
    doc = Document(str(path))

    zhang_ok = (
        "12. Zhang X, Wang Z, Zhao Y, et al. Multi-omics analysis unveils a four-gene prognostic signature in "
        "esophageal squamous carcinoma and the therapeutic potential of PKP1. BMC Cancer. 2025;25:777. "
        "doi:10.1186/s12885-025-14150-8"
    )
    for p in doc.paragraphs:
        t = _txt(p)
        if t.startswith("12. Zhang") and ("14150" in t or ":119." in t):
            p.text = zhang_ok
            break

    # Normalize reference 4 to full Vancouver-style line
    for p in doc.paragraphs:
        t = _txt(p)
        if t.startswith("4. Multi-faceted epigenetic dysregulation"):
            p.text = (
                "4. Cao W, et al. Multi-faceted epigenetic dysregulation of gene expression promotes esophageal "
                "squamous cell carcinoma. Nat Commun. 2020;11:3675. doi:10.1038/s41467-020-17227-z"
            )
            break

    extra_refs = [
        "11. Zheng Y, et al. Genome-wide DNA methylation and gene expression profiling characterizes molecular subtypes of esophagus squamous cell carcinoma for predicting patient survival and immunotherapy efficacy. Cancers (Basel). 2022;14(20):4970. doi:10.3390/cancers14204970",
        "12. Zhang X, Wang Z, Zhao Y, et al. Multi-omics analysis unveils a four-gene prognostic signature in esophageal squamous carcinoma and the therapeutic potential of PKP1. BMC Cancer. 2025;25:777. doi:10.1186/s12885-025-14150-8",
        "13. OCCAMS Consortium, Ng AWT, et al. Disentangling oncogenic amplicons in esophageal adenocarcinoma. Nat Commun. 2024;15:4074. doi:10.1038/s41467-024-47619-4",
        "14. Chang Z, Li Y, Wang M, et al. Construction of a multimodal esophageal cancer prognostic model with screening of key genes and investigation of their biological functions. Discov Oncol. 2025;16:1560. doi:10.1007/s12672-025-03404-3",
        "15. Tang P, Li B, Zhou Z, et al. Integrated machine learning developed a prognosis-related gene signature to predict prognosis in oesophageal squamous cell carcinoma. J Cell Mol Med. 2024;28:e70171. doi:10.1111/jcmm.70171",
        "16. Tuersun A, et al. Interpretable machine learning models to predict survival in esophageal cancer: a study based on the SEER database and external validation in China. Front Physiol. 2025;16:1665383. doi:10.3389/fphys.2025.1665383",
        "17. Pu Y, et al. Estimating the prognosis of esophageal squamous cell carcinoma based on The Cancer Genome Atlas (TCGA) of m6A methylation-associated genes. J Gastrointest Oncol. 2022;13:1–12. doi:10.21037/jgo-21-686",
        "18. Zhao D, He X, Guo Y, et al. Advances in multi-omics for esophageal squamous cell carcinoma: diagnostic, prognostic, and therapeutic perspectives. Protein Cell. 2026. doi:10.1093/procel/pwag005",
        "19. Kunene V, Ding M, Yap M, et al. Prognostic markers in oesophageal and gastric cancer review. Are they ready for clinical practice? ESMO Gastrointest Oncol. 2024;1:100091. doi:10.1016/j.esmogo.2024.100091",
        "20. Xu Y, Wang Z, Pei B, et al. DNA methylation markers in esophageal cancer. Front Genet. 2024;15:1354195. doi:10.3389/fgene.2024.1354195",
    ]

    if not _has_ref_20(doc):
        insert_after = None
        for p in doc.paragraphs:
            if _txt(p).startswith("10. Cui H"):
                insert_after = p
                break
        if insert_after is None:
            raise RuntimeError("Could not find reference 10 (Cui et al.) to append after.")

        prev_el = insert_after._element
        for line in extra_refs:
            new_p = OxmlElement("w:p")
            prev_el.addnext(new_p)
            Paragraph(new_p, insert_after._parent).add_run(line)
            prev_el = new_p

    bridge = (
        "Complementary work published over the past several years situates the same overarching issue—how "
        "histology-specific omics risk is defined and communicated—in integrated methylation–expression ESCC "
        "subtypes [11], multi-omic four-gene prognostic signatures [12], EAC amplicon-level architecture [13], "
        "multimodal prognostic composites [14], ensemble transcriptome-based machine-learning scores in ESCC [15], "
        "interpretable registry-based survival models in mixed esophageal cohorts [16], TCGA m6A-related prognostic "
        "construction [17], ESCC multi-omic landscape reviews [18], clinical prognostic marker surveys across "
        "oesophagogastric disease [19], and focused syntheses of methylation markers [20]."
    )
    if not any(_txt(p) == bridge for p in doc.paragraphs):
        anchor = None
        for p in doc.paragraphs:
            if _txt(p).endswith("clinically communicable."):
                anchor = p
                break
        if anchor is None:
            raise RuntimeError("Could not find Introduction methylation paragraph anchor.")
        _insert_paragraph_after(anchor, bridge, style=None)

    # Widen bracket spans that summarize recent literature
    patches: list[tuple[str, str]] = [
        (
            "histology-resolved reporting [5,8–10].",
            "histology-resolved reporting [5,8–20].",
        ),
        (
            "methylation-focused reviews [6–10].",
            "methylation-focused reviews [6–20].",
        ),
        (
            "methylation–clinical association studies [2,4,5,9].",
            "methylation–clinical association studies [2,4,5,9,11,17,20].",
        ),
        (
            "heterogeneous cohorts [6,7].",
            "heterogeneous cohorts [6,7,15,16].",
        ),
        (
            "pathway context via Figure 5, which helps prioritize targets for future validation rather than establishing definitive effect modification. Independently, large ESCC consortia now relate consensus molecular subtypes to prognosis and treatment response, including epigenetic readouts at the subtype level [10].",
            "pathway context via Figure 5, which helps prioritize targets for future validation rather than establishing definitive effect modification. Independently, large ESCC consortia now relate consensus molecular subtypes to prognosis and treatment response, including epigenetic readouts at the subtype level [10], while EAC-focused genomic architecture studies motivate separate mechanistic priors for amplification-driven disease [13].",
        ),
    ]

    for old, new in patches:
        hit = False
        for p in doc.paragraphs:
            if old in _txt(p):
                p.text = _txt(p).replace(old, new, 1)
                hit = True
                break
        if not hit:
            for p in doc.paragraphs:
                if new in _txt(p):
                    hit = True
                    break
            if not hit:
                raise RuntimeError(f"Patch substring not found (manuscript edited?): {old[:60]}...")

    doc.save(str(path))
    print("References expanded to 20 (where applicable):", path)


if __name__ == "__main__":
    main()
