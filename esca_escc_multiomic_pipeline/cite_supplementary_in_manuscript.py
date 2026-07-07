"""
Insert in-text citations to Supplementary Tables S1–S7 and Supplementary Figs. S1–S2 in Methods and Results.

Idempotent: skips replacements if supplementary citation markers are already present.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document

DOC_NAME = "TCGA_ESCA_ESCC_landmark60_full_scientific_manuscript.docx"


def _apply_if(p, predicate, new_text: str) -> bool:
    t = p.text or ""
    if not predicate(t):
        return False
    if t == new_text:
        return False
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.add_run(new_text)
    return True


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    path = root / DOC_NAME
    doc = Document(str(path))
    n = 0

    for p in doc.paragraphs:
        t = p.text or ""

        if t.startswith("We analysed the TCGA esophageal carcinoma") and "Supplementary Tables S1–S7" not in t:
            nt = t.replace(
                "Subsections below follow this order. ",
                "Subsections below follow this order. Supplementary Tables S1–S7 and Supplementary Figs. S1–S2 "
                "provide file-backed tabular exports and mediation summary panels aligned with each empirical block. ",
                1,
            )
            if nt != t and _apply_if(p, lambda x: x == t, nt):
                n += 1

        if (
            "Table 1 reports the full model×evaluation matrix" in t
            and "secondary reanalysis." in t
            and "Supplementary Table S1" not in t
        ):
            nt = t.replace(
                "secondary reanalysis.",
                "secondary reanalysis. Supplementary Table S1 provides tabular concordance summaries with "
                "source-file pointers for replication.",
                1,
            )
            if _apply_if(p, lambda x: x == t, nt):
                n += 1

        if (
            "Genome-wide contrasts revealed divergent fast-death methylation" in t
            and "than in adenocarcinoma." in t
            and "Supplementary Tables S2" not in t
        ):
            nt = t.replace(
                "than in adenocarcinoma.",
                "than in adenocarcinoma. Supplementary Tables S2–S5 report probe-level means, Δβ, and q_diff for "
                "ranked hypo- and hypermethylation contrasts by histology.",
                1,
            )
            if _apply_if(p, lambda x: x == t, nt):
                n += 1

        if (
            "non-transferable prognostic geometry documented in our opening analysis." in t
            and "Supplementary Table S6" not in t
        ):
            nt = t.replace(
                "non-transferable prognostic geometry documented in our opening analysis.",
                "non-transferable prognostic geometry documented in our opening analysis. Supplementary Table S6 "
                "documents extended Cox fields for the EAC methylation–outcome screen underlying Table 6.",
                1,
            )
            if _apply_if(p, lambda x: x == t, nt):
                n += 1

        if (
            "Fig. 6a prioritizes genomic candidates for follow-up." in t
            and "Supplementary Table S7" not in t
        ):
            nt = t.replace(
                "Fig. 6a prioritizes genomic candidates for follow-up.",
                "Fig. 6a prioritizes genomic candidates for follow-up. Supplementary Table S7 ranks mutation-to-"
                "expression and mutation-to-module associations for the mechanistic screen, and Supplementary Figs. "
                "S1–S2 visualise mediation-style contrasts complementary to the primary figures.",
                1,
            )
            if _apply_if(p, lambda x: x == t, nt):
                n += 1

        if (
            t.startswith("We performed an exploratory mutation×histology interaction screen")
            and "definitive interaction claims." in t
            and "Supplementary Tables S1–S7" not in t
        ):
            nt = t.replace(
                "definitive interaction claims.",
                "definitive interaction claims. Supplementary Tables S1–S7 and Supplementary Figs. S1–S2 catalogue "
                "reproducibility-oriented exports referenced in the Results.",
                1,
            )
            if _apply_if(p, lambda x: x == t, nt):
                n += 1

    doc.save(str(path))
    print(f"Applied supplementary citations ({n} paragraph update(s)). Saved {path}")


if __name__ == "__main__":
    main()
