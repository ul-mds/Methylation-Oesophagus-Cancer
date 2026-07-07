"""Patch manuscript: GEO table numbering S8/S9, Discussion/Conclusions emphasis, SI list."""
from __future__ import annotations

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

DOC_PATH = (
    "/Users/senol/Desktop/AML/Esophageal Carcinoma (TCGA, Firehose Legacy)/"
    "TCGA_ESCA_ESCC_landmark60_full_scientific_manuscript.docx"
)


def insert_paragraph_after(paragraph: Paragraph, style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        try:
            new_para.style = style
        except Exception:
            pass
    return new_para


def main() -> None:
    doc = Document(DOC_PATH)

    # Captions in Results: resolve clash with existing Supplementary Table S1–S7
    for para in doc.paragraphs:
        t = para.text or ""
        if t.startswith("Table S1. GSE161533"):
            para.text = (
                "Supplementary Table S8. GSE161533: TCGA pre-specified ESCC genes "
                "overlapping top-1200 tumor–normal contrasts."
            )
        elif t.startswith("Table S2. GSE53624"):
            para.text = (
                "Supplementary Table S9. GSE53624 / GSE53622: landmark (60-month) "
                "penalized Cox on lncRNA features."
            )

    # Discussion: limitations paragraph — remove false claim; cite S8–S9
    for para in doc.paragraphs:
        if (para.text or "").startswith("Limitations include reliance"):
            txt = para.text
            txt = txt.replace(
                "the lack of external validation cohorts in the current analysis environment, ",
                "the absence of assay-matched independent HM450 cohorts with harmonized OS for "
                "direct multivariate replication (notwithstanding modular GEO checks in GSE161533 and "
                "GSE53624/GSE53622 reported here; Supplementary Tables S8–S9), ",
            )
            para.text = txt
            break

    # Discussion: insert contribution paragraph after second contribution (epigenetic analysis)
    anchor = None
    for para in doc.paragraphs:
        if (para.text or "").startswith(
            "A second contribution is the clinically anchored epigenetic analysis."
        ):
            anchor = para
            break
    if anchor is None:
        raise SystemExit("Discussion anchor paragraph not found")
    # Avoid duplicate if script re-run
    already = False
    for para in doc.paragraphs:
        if (para.text or "").startswith("Complementing the TCGA-centric core, the GEO modules"):
            already = True
            break
    if not already:
        ins = insert_paragraph_after(anchor, "Normal")
        ins.add_run(
            "Complementing the TCGA-centric core, the GEO modules add interpretable external context: "
            "(i) tumor–normal expression concordance for genes pre-specified by ESCC-only TCGA analyses, "
            "reducing post hoc selection bias in gene prioritisation (Supplementary Table S8); "
            "(ii) confirmation that the landmark penalized Cox workflow can be executed on independent "
            "array data with prespecified feature selection; and (iii) an empirical readout of "
            "cross-cohort prognostic transport for a fixed lncRNA feature space, separating pipeline "
            "feasibility from signature generalisability (Supplementary Table S9). These modules do not "
            "recapitulate TCGA assay geometry or mediation structure, but they sharpen claims about "
            "lineage- and platform-specificity ahead of prospective validation."
        )

    # Discussion: strengthen closing GEO paragraph (yellow block) — add explicit S8–S9 pointer
    for para in doc.paragraphs:
        if (para.text or "").startswith("External GEO modules complement"):
            txt = para.text
            if "Key numeric summaries are provided as Supplementary Tables S8" not in txt:
                txt += (
                    " Key numeric summaries are provided as Supplementary Tables S8–S9 for editorial "
                    "consistency with the existing Supplementary Table series (S1–S7)."
                )
                para.text = txt
            break

    # Conclusions: new paragraph after first conclusions paragraph
    c_anchor = None
    for para in doc.paragraphs:
        if (para.text or "").startswith("Taken together, this work supports"):
            c_anchor = para
            break
    if c_anchor is None:
        raise SystemExit("Conclusions anchor not found")
    dup = False
    for para in doc.paragraphs:
        if (para.text or "").startswith("External GEO validation, reported as Supplementary Tables S8–S9"):
            dup = True
            break
    if not dup:
        ins2 = insert_paragraph_after(c_anchor, "Normal")
        ins2.add_run(
            "External GEO validation, reported as Supplementary Tables S8–S9, supports two practical "
            "points for translation: TCGA-derived ESCC gene priorities show limited but non-zero "
            "concordance with independent tumor–normal contrasts when genes are chosen before inspecting "
            "GSE161533, and prognostic models trained on a custom lncRNA platform can fit strongly "
            "in-cohort yet lose discrimination when transported across series—reinforcing that "
            "clinical uptake will require histology-matched, assay-matched replication rather than "
            "implicit cross-platform extrapolation."
        )

    # Supplementary Materials blurb
    for para in doc.paragraphs:
        if (para.text or "").startswith(
            "Raw and processed outputs are stored alongside this manuscript"
        ):
            t = para.text
            if "S8" not in t:
                t = t.replace(
                    "main-text Figures 1–6 and Tables 1–7 are cited in the Results.",
                    "main-text Figures 1–6 and Tables 1–7 are cited in the Results. "
                    "Supplementary Tables S8–S9 summarise modular GEO validation analyses described in the Results.",
                )
                para.text = t
            break

    # Supplementary Tables list: insert S8, S9 before the "Tables S2–S6 are derived" paragraph
    insert_before = None
    for para in doc.paragraphs:
        if (para.text or "").startswith("Tables S2–S6 are derived"):
            insert_before = para
            break
    if insert_before is None:
        raise SystemExit("Supplementary summary paragraph not found")

    def has_s8_list() -> bool:
        for para in doc.paragraphs:
            if (para.text or "").startswith(
                "Supplementary Table S8. GSE161533 (GEO GPL570)"
            ):
                return True
        return False

    if not has_s8_list():
        new_p = OxmlElement("w:p")
        insert_before._p.addprevious(new_p)
        p9 = Paragraph(new_p, insert_before._parent)
        p9.style = "Normal"
        p9.add_run(
            "Supplementary Table S9. GSE53624 / GSE53622: landmark (60-month) penalized Cox on "
            "high-variance lncRNA features (train GSE53624; hold-out score on GSE53622; in-cohort "
            "GSE53622 fit for context). Source: GEO series matrices / supplementary tables as described "
            "in Methods (External validation)."
        )
        new_p2 = OxmlElement("w:p")
        new_p.addprevious(new_p2)
        p8 = Paragraph(new_p2, insert_before._parent)
        p8.style = "Normal"
        p8.add_run(
            "Supplementary Table S8. GSE161533 (GEO GPL570): TCGA pre-specified ESCC gene union "
            "(60 Cox expression features ∪ top-80 mediation-priority genes; 110 unique genes) "
            "overlapping gene-level summaries of the top 1200 probes ranked by |mean paired "
            "tumor–normal log2 fold change| (n = 28 patients). Source: GEO series matrix; mapping via "
            "mygene.info as described in Methods."
        )

    # Extend "Tables S2–S6" summary line
    for para in doc.paragraphs:
        if (para.text or "").startswith("Tables S2–S6 are derived"):
            t = para.text
            if "S8" not in t:
                t += (
                    " Supplementary Tables S8–S9 provide external GEO validation summaries "
                    "(tumor–normal overlap; lncRNA landmark Cox transport)."
                )
                para.text = t
            break

    doc.save(DOC_PATH)
    print("Patched:", DOC_PATH)


if __name__ == "__main__":
    main()
