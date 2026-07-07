"""
Insert novelty positioning, literature contrast, and reviewer-scope paragraphs
into Manuscript.docx. All new text is RED highlight only; existing runs unchanged.
"""
from __future__ import annotations

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

DOC_PATH = (
    "/Users/senol/Desktop/AML/Esophageal Carcinoma (TCGA, Firehose Legacy)/"
    "Manuscript.docx"
)

NOVELTY_AFTER_TRANSFER = (
    "Relation to prior literature and primary novelty. Molecular distinctness between "
    "EAC and ESCC is established by the TCGA ESCA consortium and subsequent "
    "pan-cancer cell-of-origin analyses, yet most esophageal prognostic studies train "
    "and validate models within a single histology without cross-applying fixed "
    "coefficients to the complementary lineage. To our knowledge, this report is "
    "among the first to quantify that within-lineage penalized Cox models combining "
    "RNA expression and HM450 methylation (C-index ≈0.89–0.90) collapse to "
    "concordance near random chance when applied across histology (C-index "
    "≈0.45–0.47), and that pooled cohort performance is recoverable only when "
    "baseline hazards are explicitly stratified by histology (C-index ≈0.80). We "
    "therefore treat high within-lineage discrimination as a reproducibility "
    "benchmark—not a standalone performance claim—and emphasise cross-lineage "
    "transport failure as the principal substantive contribution. This framing "
    "differs from pan-cancer machine-learning reports in which shared structure can "
    "aid small cohorts; here, fixed EAC and ESCC feature spaces are not "
    "interchangeable prognostic substrates."
)

FAST_DEATH_SHISA3 = (
    "Fast-death methylation in context. Contrasting the shortest-survival tail among "
    "OS60 decedents within each histology parallels short- versus long-term survivor "
    "methylome designs applied in other malignancies, but—to our knowledge—has not "
    "been used previously as a histology-resolved discovery framework in esophageal "
    "carcinoma under a pre-specified 60-month landmark. In ESCC, SHISA3 ranked first "
    "by genome-wide q_diff across 10%, 20%, and 30% death-tail thresholds "
    "(Supplementary Table S10). This signal is strong hypomethylation (Δβ ≈−0.28) in "
    "the aggressive-survival tail, which contrasts with promoter-hypermethylation "
    "and tumour-suppressor silencing reports for SHISA3 in other squamous "
    "malignancies. The discrepancy may reflect probe context (gene-body versus "
    "classical promoter CpG islands) and the fast-death contrast (versus "
    "tumour–normal or pan-cohort designs). Leading EAC fast-death loci—including "
    "SIGLEC14, EPHB1, ZNF214, and RAB11FIP3—are reported here as discovery-grade "
    "candidates without established EAC methylation–prognosis precedents in the "
    "literature; EPHB1 protein expression has been linked to favourable outcome in "
    "esophageal cohorts, underscoring that methylation in the fast-death tail and "
    "expression-based associations need not align directionally. We do not present "
    "these loci as clinic-ready biomarkers."
)

REVIEWER_SCOPE = (
    "Anticipated interpretations and scope boundaries. For transparency we state "
    "five boundaries relevant to external evaluation. (1) Within-lineage C-index "
    "values near 0.89–0.90 are consistent with prior TCGA EAC and ESCC risk models "
    "and are not claimed as performance advances. (2) Molecular distinctness does "
    "not by itself demonstrate prognostic non-transportability; our contribution is "
    "the formal cross-application of fixed models, which prior esophageal signature "
    "papers typically omit. (3) Modular GEO analyses do not replicate TCGA HM450 "
    "fast-death methylation on overall survival and cannot validate the RNA+HM450 "
    "Cox scores out of sample—constraints already reflected in our external-data "
    "scope and Limitations. (4) The fast-death framework is conceptually related to, "
    "but not identical with, short-/long-term survivor comparisons in other cancers; "
    "its novelty here is histology-resolved application under OS60 with "
    "pre-specified sensitivity analyses (Supplementary Table S10). (5) Mutation-layer "
    "and module-shift results are exploratory prioritizations under modest event "
    "counts and are not interpreted as definitive effect modification. This "
    "conservative, pre-specified reporting aligns with secondary analysis of a "
    "public legacy cohort and prioritises technical soundness and reproducibility "
    "over translational or clinical-utility claims—criteria emphasised by journals "
    "that evaluate methodological validity independently of perceived impact."
)

CONCLUSION_NOVELTY = (
    "Collectively, the study supplies a methodologically explicit, histology-aware "
    "cautionary framework for esophageal omics reporting—complementary to, not "
    "replacing, histology-specific signature development—and a biologically anchored "
    "catalogue of lineage-resolved fast-death methylation programmes for "
    "assay-matched follow-up. Assay-matched external validation of HM450-based "
    "programmes and region-resolved SHISA3 mapping remain the translational "
    "priorities implied by these data."
)


def red_run(paragraph: Paragraph, text: str) -> None:
    r = paragraph.add_run(text)
    r.font.highlight_color = WD_COLOR_INDEX.RED


def insert_red_after(paragraph: Paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    red_run(new_para, text)
    return new_para


def find_paragraph(doc: Document, prefix: str) -> Paragraph:
    for p in doc.paragraphs:
        if (p.text or "").startswith(prefix):
            return p
    raise RuntimeError(f"Not found: {prefix[:70]!r}")


def find_paragraph_contains(doc: Document, needle: str) -> Paragraph:
    for p in doc.paragraphs:
        if needle in (p.text or ""):
            return p
    raise RuntimeError(f"Not found containing: {needle[:70]!r}")


def already_present(doc: Document, needle: str) -> bool:
    return any(needle in (p.text or "") for p in doc.paragraphs)


def main() -> None:
    doc = Document(DOC_PATH)

    if not already_present(doc, "Relation to prior literature and primary novelty"):
        p_transfer = find_paragraph(
            doc,
            "The primary conclusion of this study is established through a rigorous transportability test",
        )
        insert_red_after(p_transfer, NOVELTY_AFTER_TRANSFER)

    if not already_present(doc, "Fast-death methylation in context"):
        p_fast = find_paragraph(
            doc,
            "A secondary contribution of this work is the deployment of an outcome-anchored epigenetic framework",
        )
        insert_red_after(p_fast, FAST_DEATH_SHISA3)

    if not already_present(doc, "Anticipated interpretations and scope boundaries"):
        p_landmark = find_paragraph(
            doc,
            "Finally, we acknowledge the implications of our chosen 60-month landmark horizon",
        )
        insert_red_after(p_landmark, REVIEWER_SCOPE)

    if not already_present(doc, "Collectively, the study supplies a methodologically explicit"):
        p_conc = find_paragraph(
            doc,
            "Our findings establish a pragmatic imperative for esophageal cancer research",
        )
        insert_red_after(p_conc, CONCLUSION_NOVELTY)

    doc.save(DOC_PATH)
    print(f"Saved novelty/reviewer-scope inserts (red) to:\n  {DOC_PATH}")


if __name__ == "__main__":
    main()
