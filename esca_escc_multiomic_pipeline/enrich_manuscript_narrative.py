"""
Expand manuscript narrative: every figure/table cited in flowing Results prose;
Abstract, Methods bridge, and Discussion cross-references.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document


def main():
    study = Path(__file__).resolve().parents[1]
    path = study / "TCGA_ESCA_ESCC_landmark60_full_scientific_manuscript.docx"
    doc = Document(str(path))

    # Index -> full replacement text (non-empty paragraphs only; indices from stable structure)
    updates = {
        25: (
            "Histology-specific prognostic signal is not transferable across EAC and ESCC "
            "(Figure 1; Table 1). "
            "The Results that follow mirror the analysis pipeline: transferability, fast-death methylation "
            "(Figures 2–3; Tables 2A–2D), 60-month methylation stratification (Figure 4; Table 3), and "
            "exploratory mutation context (Figures 5–6; Table 4)."
        ),
        27: (
            "Table 1. Model performance (C-index) on OS_60mo_landmark. "
            "Referenced in the first Results subsection together with Figure 1."
        ),
        29: "Fast-death methylation programs differ by histology (Figures 2–3; Tables 2A–2D).",
        31: (
            "Table 2A. ESCA/EAC fast-death hypomethylated loci (top by q_diff). "
            "Discussed with Figure 2 (left) and Figure 3 (left panel hypo bars)."
        ),
        33: (
            "Table 2B. ESCC fast-death hypomethylated loci (top by q_diff). "
            "Discussed with Figure 2 (right) and Figure 3 (right panel hypo bars); includes SHISA3."
        ),
        35: (
            "Table 2C. ESCA/EAC fast-death hypermethylated loci (top by q_diff). "
            "Discussed with Figure 3 (left panel hyper bars)."
        ),
        37: (
            "Table 2D. ESCC fast-death hypermethylated loci (top by q_diff). "
            "Discussed with Figure 3 (right panel hyper bars)."
        ),
        39: "Selected methylation loci stratify 60-month survival within histology (Figure 4; Table 3).",
        41: (
            "Table 3. ESCA/EAC methylation~OS Cox (top by cox_q). "
            "Referenced with Figure 4 (upper row) and the KM subsection of Results."
        ),
        43: (
            "Exploratory mutation×histology interactions suggest histology-dependent downstream context "
            "(Figures 5–6; Table 4)."
        ),
        45: (
            "Table 4. Top mechanistic associations (mutation → expression/module; lowest p-values). "
            "Referenced with Figure 5 (module heatmap), Figure 6 (interaction screen), and the mutation subsection of Results."
        ),
        4: (
            "Methods: We analyzed the TCGA ESCA Firehose Legacy study package, assembling a sample-level dataset "
            "with clinical outcomes, HM450 methylation, and RNA expression. Overall survival (OS) was evaluated "
            "using a 60-month landmark (follow-up capped at 60 months; deaths after 60 months censored). We trained "
            "histology-specific Cox models (EAC-only; ESCC-only) and assessed both within-histology performance and "
            "cross-histology transferability using the concordance index. To connect epigenetics to clinically aggressive "
            "disease, we defined a fast-death phenotype among decedents within each histology and identified differential "
            "methylation in fast-death versus others, followed by 60-month Kaplan–Meier visualizations for selected loci. "
            "We additionally performed exploratory mutation×histology interaction screening and mechanistic follow-up "
            "using expression-derived module scores. Primary numeric summaries appear in Tables 1–4; integrated graphics "
            "appear in Figures 1–6."
        ),
        5: (
            "Results: Histology-specific models discriminated well within histology (EAC-only C-index=0.894; "
            "ESCC-only C-index=0.901) but failed to transfer across histology (EAC→ESCC C-index=0.469; "
            "ESCC→EAC C-index=0.450), as summarized graphically in Figure 1 and in full in Table 1. A pooled model "
            "with histology-stratified baseline improved overall performance (C-index=0.803; Table 1), showing that "
            "pooling requires explicit histology handling. Fast-death methylation differed sharply by histology: "
            "genome-wide contrasts (Figure 2) and ranked hypo-/hypermethylation bar panels (Figure 3) aligned with "
            "Tables 2A–2D—EAC fast-death featured hypomethylation at loci including SIGLEC14, EPHB1, ZNF214, and "
            "RAB11FIP3 (Table 2A) with complementary hypermethylated ranks (Table 2C), whereas ESCC fast-death was "
            "dominated by strong SHISA3 hypomethylation (Table 2B) alongside C7orf13, PRKG1, and SEMA3E, with "
            "hypermethylated ranks in Table 2D. Methylation–OS associations in EAC (Table 3) and survival separation "
            "by median methylation across histologies are shown in Figure 4. Finally, mutation-linked module shifts "
            "(Figure 5; Table 4) and the nominal mutation×histology interaction landscape (Figure 6) provide exploratory "
            "downstream context."
        ),
        6: (
            "Conclusions: Prognostic signal in TCGA esophageal carcinoma is strongly histology-specific and largely "
            "non-transferable between EAC and ESCC (Figure 1; Table 1). Histology-resolved fast-death epigenetic programs "
            "(Figures 2–3; Tables 2A–2D) and 60-month survival stratification (Figure 4; Table 3) provide a clinically "
            "anchored view of aggressive biology within a bounded horizon, while mutation-layer follow-up (Figures 5–6; "
            "Table 4) motivates external validation."
        ),
        9: (
            "The absence of an explicit transferability test creates an interpretability gap. A prognostic model that "
            "performs well within one histology may simply be learning biology that is specific to that histology, "
            "making cross-histology deployment inappropriate. Conversely, a model that truly reflects shared aggressiveness "
            "biology might retain performance when applied across histologies. We therefore foregrounded transferability "
            "as a primary endpoint: we trained EAC-only and ESCC-only prognostic models and measured their performance "
            "both within histology and when applied to the other histology, with concordance summarized in Figure 1 and "
            "Table 1."
        ),
        10: (
            "In parallel, DNA methylation is attractive for prognosis research because it is stable, captures both "
            "cell-of-origin and tumor state, and can be measured robustly. Numerous methylation-based prognostic "
            "signatures have been reported for esophageal cancer, particularly ESCC. Rather than pursuing another "
            "signature for its own sake, we anchored methylation discovery to an intuitive clinical phenotype—fast death "
            "among decedents—and assessed methylation programs separately in EAC and ESCC within a 60-month landmark OS "
            "window, visualized as paired histology volcanoes and bar panels (Figures 2–3) with full locus-level ranks "
            "in Tables 2A–2D. We then linked prioritized loci to 60-month Kaplan–Meier separation (Figure 4) and, in EAC, "
            "to univariate Cox rankings (Table 3). This framing aims to produce results that are both biologically "
            "interpretable and clinically communicable."
        ),
        26: (
            "We first evaluated whether prognostic models trained within one histology generalize to the other. Using "
            "60-month landmark OS, both histology-specific models showed strong within-histology discrimination (EAC-only: "
            "C-index=0.894 in EAC; ESCC-only: C-index=0.901 in ESCC). Cross-application, however, collapsed to near-random "
            "discrimination (EAC→ESCC: C-index=0.469; ESCC→EAC: C-index=0.450), so the risk ordering learned in one "
            "histology does not preserve ranking utility in the other—this transferability failure is the central message "
            "of Figure 1. A pooled Cox model with histology-stratified baseline achieved a higher combined C-index "
            "(0.803), illustrating that joint modeling can recover prognostic signal when baseline risk is allowed to "
            "differ by histology. Together, the bar display in Figure 1 and the complete model×test matrix in Table 1 "
            "support the same conclusion: prognostic structure is histology-bound in this cohort."
        ),
        30: (
            "To connect epigenetic variation to clinically aggressive disease, we defined a fast-death phenotype among "
            "decedents within each histology and contrasted HM450 beta values between fast-death cases and all other "
            "samples in the same histology. Figure 2 tells the genome-wide story: the left volcano (ESCA/EAC) shows a "
            "broad excess of statistically supported negative Δβ (hypomethylation in fast-death), whereas the right "
            "volcano (ESCC) is dominated by a prominent SHISA3 signal on the hypomethylated side, indicating that the "
            "most statistically extreme fast-death epigenetic shift differs in both extent and focal gene emphasis "
            "between histologies. Figure 3 compresses these landscapes into interpretable bar summaries of the leading "
            "hypo- and hypermethylated loci side-by-side, making directionality and locus composition easy to compare. "
            "The underlying ranks are enumerated in Tables 2A–2D: Tables 2A and 2B list top hypomethylated loci in "
            "EAC and ESCC, respectively—including SIGLEC14, EPHB1, ZNF214, and RAB11FIP3 in EAC (Table 2A) and SHISA3, "
            "C7orf13, PRKG1, and SEMA3E among ESCC leaders (Table 2B)—while Tables 2C and 2D provide the parallel "
            "hypermethylated ranks used to construct the upward-oriented bars in Figure 3. In short, Figures 2–3 and "
            "Tables 2A–2D jointly show that “fast-death” methylation is not a single shared program but a "
            "histology-specific epigenetic portrait."
        ),
        40: (
            "We next asked whether methylation at prioritized loci translates into clinically visible 60-month survival "
            "separation when patients are split at the median probe beta. In ESCA/EAC, univariate Cox screening on OS60 "
            "highlighted CDC14C, VNN3, and GXYLT2 among the strongest methylation–OS associations (Table 3); Figure 4 "
            "(upper row) shows that median-dichotomized methylation at these probes yields discernible Kaplan–Meier "
            "separation within the landmark window, linking the statistical ranking in Table 3 to interpretable survival "
            "curves. For ESCC, we plotted the three probes with the strongest fast-death differential evidence (q_diff "
            "ranking) that also supported KM visualization—SHISA3, PRKG1, and C7orf13—so the lower row of Figure 4 "
            "directly extends the fast-death discovery panel (Figures 2–3; Tables 2A–2B) into a time-to-event view. "
            "Taken together, Figure 4 and Table 3 demonstrate that histology-appropriate probe selection yields "
            "methylation strata with distinct 60-month OS profiles, while underscoring that ESCC and EAC differ in "
            "which probes drive both association testing and visualization."
        ),
        44: (
            "Finally, we performed an exploratory mutation×histology interaction screen on OS60. Because interaction "
            "tests are power-hungry and events are sparse, few relationships survive strict FDR control; Figure 6 "
            "therefore presents the nominal interaction landscape (β versus −log10 p) as a hypothesis map rather than "
            "a confirmatory catalog, with labeled genes indicating where histology appears to modify mutation–outcome "
            "slopes. To ground those candidates biologically, Figure 5 summarizes histology-stratified shifts in "
            "expression-derived module scores (Δ=mutant−wild-type), highlighting, for example, EMT and immune/IFN-γ "
            "axes that move differently by histology; Table 4 lists the strongest mutation→module or mutation→expression "
            "associations feeding that heatmap. Read sequentially, Figure 6 motivates which genomic alterations merit "
            "follow-up, while Figure 5 and Table 4 supply plausible transcriptional mechanisms subject to independent "
            "replication."
        ),
        61: (
            "The primary finding is straightforward but often left implicit: in TCGA esophageal carcinoma, prognostic "
            "signal is not shared uniformly across histologies. Models trained within EAC and ESCC each performed well "
            "within their own histology, yet cross-application failed—Figure 1 and Table 1 make this collapse explicit "
            "for readers who prefer graphical or tabular summaries. This transferability failure is not a technical "
            "nuisance; it is a biological result that argues against treating EAC and ESCC as a single prognostic "
            "entity and helps explain why many published signatures do not replicate when moved across clinically "
            "heterogeneous cohorts."
        ),
        62: (
            "A second contribution is the clinically anchored epigenetic analysis. Rather than optimizing a "
            "high-dimensional signature, we asked what methylation programs distinguish the most aggressive clinical "
            "courses—fast death among decedents—within each histology. Figures 2–3, together with Tables 2A–2D, show "
            "that the leading hypo- and hypermethylation candidates differ sharply between histologies, consistent with "
            "distinct cell-of-origin and pathway wiring. Figure 4 then demonstrates that carefully chosen probes can "
            "still stratify 60-month OS within each histology, with EAC priorities guided by Table 3. These programs "
            "can be communicated as testable hypotheses (e.g., SHISA3-associated epigenetic state in aggressive ESCC) "
            "and evaluated in external cohorts with methylation assays."
        ),
        63: (
            "Our use of a 60-month landmark OS window serves two purposes: it matches common clinical decision horizons "
            "and reduces sensitivity to late follow-up differences. That said, landmarking can reduce power for late "
            "events, and ESCC contributes fewer OS events within the window than EAC, which limits precision for "
            "Kaplan–Meier displays such as Figure 4 (lower row) even when fast-death methylation contrasts are strong "
            "(Figures 2–3; Tables 2A–2B). Table 3 summarizes EAC-only methylation–OS Cox screening on OS60. "
            "The mutation×histology interaction results should be interpreted cautiously; interaction testing is "
            "power-intensive, and we deliberately positioned Figure 6 and Table 4 as exploratory. Their value lies in "
            "connecting candidate alterations to downstream pathway context via Figure 5, which helps prioritize targets "
            "for future validation rather than establishing definitive effect modification."
        ),
        64: (
            "Limitations include reliance on a cBioPortal study package without a full HM450 probe annotation table "
            "within the package (gene tokens reflect probe naming), the lack of external validation cohorts in the "
            "current analysis environment, and the possibility that some individual loci have prior prognostic reports "
            "in other cancers or in esophageal cancer. Figures 1–6 and Tables 1–4 should therefore be read as an "
            "internally consistent TCGA case study: our novelty claim is not that any single gene is unprecedented, "
            "but that the integrated design—histology-resolved fast-death epigenetics plus explicit cross-histology "
            "transferability testing within a standardized 60-month framework—fills a practical gap in the prognostic "
            "literature."
        ),
        66: (
            "In TCGA esophageal carcinoma, prognostic models are strongly histology-specific and do not transfer between "
            "EAC and ESCC (Figure 1; Table 1). Fast-death anchored methylation analyses reveal distinct histology-resolved "
            "epigenetic programs (Figures 2–3; Tables 2A–2D) and clinically interpretable 60-month survival separation "
            "(Figure 4; Table 3), while mutation-layer follow-up (Figures 5–6; Table 4) frames mechanistic next steps."
        ),
        47: (
            "Figures. The panels below are referenced sequentially in Results; each subsection interprets the "
            "corresponding figure(s) and table(s) in prose."
        ),
        48: (
            "Figure 1. Cross-apply performance (C-index) using 60-month landmark OS. Histology-specific models "
            "discriminate within histology but do not transfer across histologies. See Results (first subsection) "
            "and Table 1."
        ),
        50: (
            "Figure 2. Fast-death differential methylation by histology (volcano). Left: ESCA/EAC; right: ESCC. "
            "Each point is a locus with Δβ = mean(β_fast)−mean(β_other) and y-axis = −log10(q_diff). "
            "See Results (fast-death methylation subsection) and Tables 2A–2D."
        ),
        52: (
            "Figure 3. Top fast-death hypo- and hypermethylated loci (Δβ) in ESCA/EAC (left) and ESCC (right). "
            "See Results (fast-death methylation subsection) and Tables 2A–2D."
        ),
        54: (
            "Figure 4. Kaplan–Meier curves (60-month landmark OS) stratified by median methylation. Upper row: "
            "ESCA/EAC, three probes with strongest methylation–OS signal (univariate Cox on OS60). Lower row: ESCC, "
            "three probes with strongest fast-death differential methylation (q_diff). See Results (KM subsection) "
            "and Table 3."
        ),
        56: (
            "Figure 5. Exploratory mechanistic follow-up: mutation-associated shifts in expression-derived module "
            "scores (Δ=mut−WT) stratified by histology. See Results (mutation subsection) and Table 4."
        ),
        58: (
            "Figure 6. Mutation×histology interaction screen on 60-month landmark OS (nominal p-values). "
            "See Results (mutation subsection)."
        ),
        74: (
            "All analysis outputs are available in the accompanying folders: result/ (tables) and Plots/ (figures). "
            "Figure 1 corresponds to model_performance_cindex.tsv (Table 1). Figures 2–3 and Tables 2A–2D derive from "
            "methylation_fast_death_EAC_ESCA.tsv and methylation_fast_death_ESCC.tsv. Figure 4 and Table 3 use the same "
            "EAC table’s Cox columns plus ESCC fast-death rankings. Figures 5–6 and Table 4 build on "
            "interaction_mechanism_mut_to_expr_modules.tsv and mutation_x_histology_interactions_OS.tsv."
        ),
    }

    for idx, new_text in updates.items():
        if idx >= len(doc.paragraphs):
            raise IndexError(f"Paragraph index {idx} out of range ({len(doc.paragraphs)})")
        doc.paragraphs[idx].text = new_text

    doc.save(str(path))
    print("Updated narrative:", path)


if __name__ == "__main__":
    main()
