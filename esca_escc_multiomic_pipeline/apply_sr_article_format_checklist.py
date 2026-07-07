"""
Scientific Reports — 'Format of articles' checklist (title page, abstract).

- Replaces abstract with an unstructured, reference-free text of ≤200 words (SR limit).
- Inserts a minimal title page (authors*, affiliation, corresponding contact placeholders)
  after the title if missing, per SR title-page rules.

Run from study root parent folder (same as other manuscript scripts).
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph as DocxParagraph

DOC_NAME = "TCGA_ESCA_ESCC_landmark60_full_scientific_manuscript.docx"

# Official SR cap: 200 words; unstructured; no references.
ABSTRACT_SR = (
    "Esophageal adenocarcinoma (EAC) and esophageal squamous cell carcinoma (ESCC) are often pooled in secondary "
    "analyses despite distinct biology. Using TCGA ESCA (Firehose Legacy), we tested whether multi-omic prognostic "
    "structure transfers across histologies for overall survival with a 60-month landmark and HM450 methylation plus "
    "expression features. Penalized Cox models trained within each histology discriminated strongly when tested in the "
    "same lineage (concordance ≈0.89–0.90) but failed when cross-applied (≈0.45–0.47 versus chance 0.5). A pooled model "
    "with histology-stratified baseline hazards achieved concordance 0.803, indicating recoverable combined-cohort risk "
    "only with explicit histology in the hazard. Among decedents, fast-death versus other deaths uncovered divergent "
    "methylation programs between EAC and ESCC, with representative loci including SIGLEC14, EPHB1, ZNF214, and "
    "RAB11FIP3 in EAC and SHISA3, C7orf13, PRKG1, and SEMA3E in ESCC; selected probes linked to interpretable "
    "60-month survival separation. Exploratory mutation-by-histology screens and expression-module shifts contextualize "
    "somatic variation. EAC and ESCC should be treated as separate prognostic entities for omics-derived scores; "
    "histology-matched external validation is needed. This secondary analysis illustrates how pooling without "
    "transferability testing can mislead interpretation of omics-based risk and supports histology-aware reporting in "
    "public cohort studies. De-identified public-domain TCGA data only; full statistics appear in Methods; mutation "
    "findings are exploratory."
)

TITLE_PAGE_MARKER = "Corresponding author:"


def _insert_paragraph_after(anchor: DocxParagraph, doc: Document) -> DocxParagraph:
    new_el = OxmlElement("w:p")
    anchor._element.addnext(new_el)
    return DocxParagraph(new_el, doc._body)


def _text_before_abstract(doc: Document) -> str:
    parts = []
    for p in doc.paragraphs:
        if p.text.strip() == "Abstract":
            break
        parts.append(p.text)
    return "\n".join(parts)


def _should_add_title_page(doc: Document) -> bool:
    blob = _text_before_abstract(doc)
    if TITLE_PAGE_MARKER in blob:
        return False
    if "[Author One]*" in blob:
        return False
    return True


def _apply_title_page(doc: Document) -> bool:
    if not _should_add_title_page(doc):
        return False
    p0 = doc.paragraphs[0]
    el = p0._element.getnext()
    while el is not None:
        nxt = el.getnext()
        if el.tag.endswith("}sectPr"):
            break
        if el.tag.endswith("}p"):
            t = DocxParagraph(el, doc._body).text.strip()
            if t == "Abstract":
                break
            if t == "":
                el.getparent().remove(el)
        el = nxt

    p0 = doc.paragraphs[0]
    a1 = _insert_paragraph_after(p0, doc)
    a1.add_run(
        "[Author One]*, [Author Two], [Author Three] — replace with final author list; mark corresponding author with *."
    )
    a2 = _insert_paragraph_after(a1, doc)
    a2.add_run(
        "[Affiliation(s): department, institution, city, postcode, country — use superscripts as needed.]"
    )
    a3 = _insert_paragraph_after(a2, doc)
    a3.add_run(
        "*Corresponding author: name.surname@institution.org (add telephone and postal address as required by the journal)."
    )
    _insert_paragraph_after(a3, doc)
    return True


def _apply_abstract(doc: Document) -> None:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() != "Abstract":
            continue
        body = doc.paragraphs[i + 1]
        body.text = ABSTRACT_SR
        return
    raise RuntimeError("Abstract heading not found.")


def main() -> None:
    wc = len(ABSTRACT_SR.split())
    if wc > 200:
        raise RuntimeError(f"Abstract must be ≤200 words for SR (got {wc}).")
    if re.search(r"\[\d+\]", ABSTRACT_SR):
        raise RuntimeError("Abstract must not contain reference-style [n] citations.")
    path = Path(__file__).resolve().parents[1] / DOC_NAME
    doc = Document(str(path))
    _apply_abstract(doc)
    added = _apply_title_page(doc)
    doc.save(str(path))
    print(f"Abstract set to {wc} words (≤200). Title page inserted: {added}. Saved:", path)


if __name__ == "__main__":
    main()
