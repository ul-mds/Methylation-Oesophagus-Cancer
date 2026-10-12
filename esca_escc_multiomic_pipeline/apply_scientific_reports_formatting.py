"""
Apply Cancer Medicine (Wiley) reference styling to the manuscript docx:
- Reference list: journal italic, volume bold, sequential numbering
- Main-text figure captions: bold Fig. n. + short title (full legend at end only)
- End figure legends: bold Figure n. + description
- Table titles: bold Table n. + description
- Table style: Table Grid (editable)
- Refresh display-item guidance paragraph

Idempotent markers:
- Ref formatting: skips if the first numbered reference paragraph already contains an italic run (journal name).
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document


def _clear_runs(paragraph) -> None:
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)


def _add_reference_paragraph(paragraph, num: int, author_title: str, journal: str, volume: str | None, tail: str) -> None:
    _clear_runs(paragraph)
    p = paragraph.add_run(f"{num}. {author_title}")
    p = paragraph.add_run(journal)
    p.italic = True
    if volume is not None:
        paragraph.add_run(" ")
        v = paragraph.add_run(volume)
        v.bold = True
    paragraph.add_run(tail)


# (author + title + trailing space before journal), journal (italic), volume or None, tail after volume/journal
REF_PARTS: list[tuple[str, str, str | None, str]] = [
    (
        "Cancer Genome Atlas Research Network. Integrated genomic characterization of oesophageal carcinoma. ",
        "Nature",
        "541",
        ", 169-175 (2017).",
    ),
    (
        "Li, B. et al. Identification of DNA methylation-driven genes in esophageal squamous cell carcinoma: "
        "a study based on The Cancer Genome Atlas. ",
        "Cancer Cell Int.",
        "19",
        ", 82 (2019).",
    ),
    (
        "Chen, L. et al. Identification of prognostic risk model based on DNA methylation-driven genes in "
        "esophageal adenocarcinoma. ",
        "Biomed. Res. Int.",
        "2021",
        ", 6690910 (2021).",
    ),
    (
        "Cao, W. et al. Multi-faceted epigenetic dysregulation of gene expression promotes esophageal squamous "
        "cell carcinoma. ",
        "Nat. Commun.",
        "11",
        ", 3675 (2020).",
    ),
    (
        "Xi, Y. et al. Multi-omic characterization of genome-wide abnormal DNA methylation reveals diagnostic "
        "and prognostic markers for esophageal squamous-cell carcinoma. ",
        "Sig. Transduct. Target. Ther.",
        "7",
        ", 53 (2022).",
    ),
    (
        "Lyu, Q. et al. Integrated bioinformatics analysis of differences between EAC and ESCC. ",
        "BMC Cancer",
        "25",
        ", 1668 (2025).",
    ),
    (
        "Pinto, R. et al. Molecular stratification of esophageal adenocarcinoma: implications for prognosis and "
        "treatment strategy. ",
        "Oncogene",
        None,
        " (2025); https://doi.org/10.1038/s41388-025-03650-3",
    ),
    (
        "Tang, J. et al. DNA methylation in esophageal cancer: technological advances and early detection "
        "clinical applications. ",
        "Front. Oncol.",
        "15",
        ", 1543190 (2025).",
    ),
    (
        "Wang, G. et al. Analysis of the associations between DNA methylation and clinical features reveals the "
        "potential oncogenic effect of CFAP52 in esophageal squamous cell cancer. ",
        "Cancer Cell Int.",
        "25",
        ", 394 (2025).",
    ),
    (
        "Cui, H. et al. The consensus molecular subtypes of esophageal squamous cell carcinoma. ",
        "Sig. Transduct. Target. Ther.",
        "11",
        ", 65 (2026).",
    ),
    (
        "Zheng, Y. et al. Genome-wide DNA methylation and gene expression profiling characterizes molecular "
        "subtypes of esophagus squamous cell carcinoma for predicting patient survival and immunotherapy efficacy. ",
        "Cancers",
        "14",
        ", 4970 (2022).",
    ),
    (
        "Zhang, X. et al. Multi-omics analysis unveils a four-gene prognostic signature in esophageal squamous "
        "carcinoma and the therapeutic potential of PKP1. ",
        "BMC Cancer",
        "25",
        ", 777 (2025).",
    ),
    (
        "Oesophageal Cancer Clinical and Molecular Stratification (OCCAMS) Consortium et al. Disentangling "
        "oncogenic amplicons in esophageal adenocarcinoma. ",
        "Nat. Commun.",
        "15",
        ", 4074 (2024).",
    ),
    (
        "Chang, Z. et al. Construction of a multimodal esophageal cancer prognostic model with screening of key "
        "genes and investigation of their biological functions. ",
        "Discov. Oncol.",
        "16",
        ", 1560 (2025).",
    ),
    (
        "Tang, P. et al. Integrated machine learning developed a prognosis-related gene signature to predict "
        "prognosis in oesophageal squamous cell carcinoma. ",
        "J. Cell. Mol. Med.",
        "28",
        ", e70171 (2024).",
    ),
    (
        "Tuersun, A. et al. Interpretable machine learning models to predict survival in esophageal cancer: "
        "a study based on the SEER database and external validation in China. ",
        "Front. Physiol.",
        "16",
        ", 1665383 (2025).",
    ),
    (
        "Pu, Y. et al. Estimating the prognosis of esophageal squamous cell carcinoma based on The Cancer Genome "
        "Atlas (TCGA) of m6A methylation-associated genes. ",
        "J. Gastrointest. Oncol.",
        "13",
        ", 1-12 (2022).",
    ),
    (
        "Zhao, D. et al. Advances in multi-omics for esophageal squamous cell carcinoma: diagnostic, prognostic, "
        "and therapeutic perspectives. ",
        "Protein Cell",
        None,
        " (2026); https://doi.org/10.1093/procel/pwag005",
    ),
    (
        "Kunene, V. et al. Prognostic markers in oesophageal and gastric cancer review. Are they ready for clinical "
        "practice? ",
        "ESMO Gastrointest. Oncol.",
        "1",
        ", 100091 (2024).",
    ),
    (
        "Xu, Y. et al. DNA methylation markers in esophageal cancer. ",
        "Front. Genet.",
        "15",
        ", 1354195 (2024).",
    ),
]


def _refs_already_styled(doc: Document) -> bool:
    in_refs = False
    for p in doc.paragraphs:
        t = p.text.strip()
        if t == "References":
            in_refs = True
            continue
        if in_refs:
            if t.startswith("Acknowledgements") or t.startswith("Supplementary Materials"):
                break
            if not t or not re.match(r"^\d+\.", t):
                continue
            if any(r.italic for r in p.runs if (r.text or "").strip()):
                return True
            return False
    return False


def _format_references(doc: Document) -> None:
    if _refs_already_styled(doc):
        print("References already styled (italic journal runs); skipping reference block.")
        return

    in_refs = False
    n = 0
    for p in doc.paragraphs:
        t = p.text.strip()
        if t == "References":
            in_refs = True
            continue
        if in_refs:
            if t.startswith("Acknowledgements") or t.startswith("Supplementary Materials"):
                break
            if not t or not re.match(r"^\d+\.", t):
                continue
            n += 1
            if n > len(REF_PARTS):
                print(f"Warning: reference #{n} not in REF_PARTS; leaving unchanged.")
                continue
            author_title, journal, vol, tail = REF_PARTS[n - 1]
            _add_reference_paragraph(p, n, author_title, journal, vol, tail)
    if n != len(REF_PARTS):
        print(f"Warning: expected {len(REF_PARTS)} numbered references, found {n}.")


def _short_fig_caption(full_text: str) -> tuple[str, str]:
    """Return (bold_label 'Fig. n.', rest = first sentence of description)."""
    m = re.match(r"^Figure (\d+)\.\s*(.*)$", full_text.strip(), re.DOTALL)
    if not m:
        return "", full_text
    num, desc = m.group(1), m.group(2).strip()
    label = f"Fig. {num}."
    if not desc:
        return label, ""
    parts = desc.split(". ")
    first = parts[0].strip()
    if len(parts) > 1:
        first = first + "."
    elif not first.endswith("."):
        first = first + "."
    return label, " " + first


def _format_main_body_figure_captions(doc: Document) -> None:
    fig_leg_i = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "Figure legends":
            fig_leg_i = i
            break
    if fig_leg_i is None:
        print("No 'Figure legends' heading; skipping main-text figure caption trim.")
        return

    for i, p in enumerate(doc.paragraphs):
        if i >= fig_leg_i:
            break
        t = p.text.strip()
        if not re.match(r"^Figure \d+\.", t):
            continue
        label, rest = _short_fig_caption(p.text)
        if not label:
            continue
        _clear_runs(p)
        lb = p.add_run(label)
        lb.bold = True
        if rest:
            p.add_run(rest)


def _format_end_figure_legends(doc: Document) -> None:
    past = False
    for p in doc.paragraphs:
        t = p.text.strip()
        if t == "Figure legends":
            past = True
            continue
        if not past:
            continue
        if t.startswith("Supplementary Materials"):
            break
        if not re.match(r"^Figure \d+\.", t):
            continue
        m = re.match(r"^(Figure \d+\.)\s*(.*)$", p.text.strip(), re.DOTALL)
        if not m:
            continue
        body = m.group(2)
        _clear_runs(p)
        h = p.add_run(m.group(1))
        h.bold = True
        if body:
            p.add_run(" " + body)


def _format_table_caption_paragraphs(doc: Document) -> None:
    for p in doc.paragraphs:
        t = p.text.strip()
        m = re.match(r"^(Table \d+\.)\s*(.*)$", t, re.DOTALL)
        if not m:
            continue
        _clear_runs(p)
        head = p.add_run(m.group(1))
        head.bold = True
        if m.group(2):
            p.add_run(" " + m.group(2))


def _apply_table_grid_style(doc: Document) -> None:
    for tbl in doc.tables:
        try:
            tbl.style = "Table Grid"
        except KeyError:
            try:
                tbl.style = "Light Grid"
            except KeyError:
                pass


def _update_display_item_note(doc: Document) -> None:
    new_txt = (
        "Cancer Medicine author guidelines note: at most eight combined figures and tables in the main article. "
        "This manuscript currently includes six figures and seven tables (thirteen display items). "
        "Before submission, move at least five items to Supplementary Information (for example "
        "Supplementary Tables S1–S5) or merge panels, and cite each in the main text using the "
        "word Supplementary (e.g. Supplementary Table S1; Supplementary Fig. S1)."
    )
    for p in doc.paragraphs:
        if "Cancer Medicine expects at most eight" in p.text or "display items" in p.text and "six figures" in p.text:
            if p.text.strip() == new_txt.strip():
                return
            _clear_runs(p)
            p.add_run(new_txt)
            return


def main() -> None:
    path = Path(__file__).resolve().parents[1] / "TCGA_ESCA_ESCC_landmark60_full_scientific_manuscript.docx"
    doc = Document(str(path))
    _format_references(doc)
    _format_main_body_figure_captions(doc)
    _format_end_figure_legends(doc)
    _format_table_caption_paragraphs(doc)
    _apply_table_grid_style(doc)
    _update_display_item_note(doc)
    doc.save(str(path))
    print("Cancer Medicine formatting applied:", path)


if __name__ == "__main__":
    main()
