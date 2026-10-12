"""
Manuscript First.docx:
  1. Move Fig. 1 (pipeline) to Methods section (after Introduction flow → at Methods start).
  2. Remove Fig. 1 from Results opening; keep Methods citation to Fig. 1.
  3. Renumber references 1–30 in strict order of first appearance (Vancouver).
  4. Max three references per sentence in Introduction rewrites.
  5. Preserve yellow/red Word highlights on non-rewritten paragraphs.
"""
from __future__ import annotations

from esca_escc_multiomic_pipeline.study_paths import study_root

import argparse
import re
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

STUDY_ROOT = study_root()
DEFAULT_DOCX = STUDY_ROOT / "Manuscript First.docx"
W_T = qn("w:t")

# old_ref_number -> new_ref_number (order of first appearance in revised manuscript)
OLD_TO_NEW: dict[int, int] = {
    22: 1,   # Thrift — epidemiology
    21: 2,   # Spechler — Barrett's
    1: 3,    # TCGA ESCA Nature
    6: 4,    # Lyu EAC vs ESCC
    7: 5,    # Pinto EAC stratification
    10: 6,   # Cui ESCC subtypes
    11: 7,   # Zheng methylation subtypes
    23: 8,   # CROSS trial
    24: 9,   # Bremner Barrett origin
    25: 10,  # Tosh phenotype
    26: 11,  # Sarosi bone marrow
    27: 12,  # Hu Barrett stem cell
    19: 13,  # Kunene prognostic review
    28: 14,  # Mandard TRG / predictive
    12: 15,  # Zhang multi-omics ESCC
    13: 16,  # OCCAMS EAC
    14: 17,  # Chang multimodal
    15: 18,  # Tang P ML ESCC
    16: 19,  # Tuersun ML SEER
    17: 20,  # Pu m6A TCGA
    2: 21,   # Li methylation ESCC
    3: 22,   # Chen methylation EAC
    4: 23,   # Cao epigenetic ESCC
    5: 24,   # Xi multi-omic methylation
    18: 25,  # Zhao multi-omics review
    20: 26,  # Xu methylation markers
    8: 27,   # Tang J methylation clinical
    9: 28,   # Wang methylation clinical
    30: 29,  # ABSOLUTE (Limitations — before Methods in SR layout)
    29: 30,  # MSigDB (Methods)
}

INTRO_13 = (
    "Esophageal carcinoma is characterized by two predominant histological subtypes—adenocarcinoma (EAC) "
    "and squamous cell carcinoma (ESCC)—which exhibit divergent epidemiological profiles and molecular "
    "landscapes [1]. The most important risk factors for ESCC are alcohol use and tobacco smoking, whereas "
    "EAC is associated with the precursor lesion Barrett's esophagus, in which the normal squamous lining "
    "is replaced by abnormal columnar mucosa as a result of gastro-oesophageal reflux disease [2]. "
    "Although both entities arise within the oesophagus, their anatomic distribution is not identical: ESCC "
    "is more commonly located in the mid oesophagus, whereas EAC typically occurs in the distal oesophagus "
    "or gastro-oesophageal junction. Genomic characterization through The Cancer Genome Atlas (TCGA) has "
    "established that these histologies represent distinct biological entities rather than interchangeable "
    "variants [3]. Subsequent pan-cancer analyses have further delineated lineage-specific transcriptional "
    "programs and pathway architectures [4,5]. Epigenetic phenotypes and ESCC subtyping are likewise "
    "lineage-linked [6,7]. In several European countries, EAC and ESCC are still managed "
    "within shared neoadjuvant chemoradiotherapy frameworks (e.g., the CROSS regimen) [8]. Despite this "
    "divergence, secondary prognostic modeling efforts frequently aggregate EAC and ESCC into a single "
    "cohort, and many lineage-specific signatures are reported without rigorous assessment of cross-histology "
    "prognostic transferability."
)

INTRO_14 = (
    "For EAC, the cellular origin of Barrett's mucosa remains debated. Competing hypotheses include migration "
    "of columnar gastric (stem) cells, trans-differentiation of native oesophageal squamous cells, and "
    "contribution of bone-marrow–derived progenitors [9–11]. Metaplastic expansion from submucosal gland "
    "ducts provides an additional proposed route [12]. We therefore use cell-of-origin language when "
    "interpreting histology-resolved epigenetic programmes rather than implying a single uncontested "
    "developmental route."
)

INTRO_15 = (
    "Throughout this manuscript, prognostic denotes association with overall survival (or landmark OS) "
    "independent of a specified therapy [13], whereas predictive is reserved for treatment-response contexts "
    "(e.g., pathological regression after chemoradiotherapy) [14]; we do not use the terms interchangeably."
)

INTRO_16 = (
    "This methodological oversight creates a significant gap in the translational utility of public omics "
    "resources. For a prognostic score to be clinically robust as a pan-esophageal tool, it must maintain "
    "its risk-stratification capacity across major histologies. Failure to do so suggests that the underlying "
    "signal is intrinsically tied to the lineage-specific background, meaning that implicit pooling may "
    "obscure rather than clarify biological interpretation. This concern is particularly relevant for recent "
    "multi-omics prognostic constructions validated within isolated cohorts or single lineages [15–17]. "
    "Parallel machine-learning survival frameworks exhibit similar lineage-bound reporting patterns [18–20]."
)

INTRO_17 = (
    "In the present study, we prioritize cross-histology transferability as a primary evaluative endpoint. "
    "Utilizing penalized Cox proportional hazards models within the TCGA-ESCA framework, we quantified the "
    "internal and cross-lineage performance of prognostic scores over a 60-month overall survival horizon. "
    "Beyond prognostic modeling, we shifted our focus toward an outcome-anchored discovery of DNA methylation "
    "programs. Histology-specific methylation-driver panels have been reported for ESCC and EAC [21–23]. "
    "DNA methylation offers a stable and biologically informative substrate for biomarker development, although "
    "many existing panels prioritize prognostic accuracy at the expense of mechanistic interpretability [24]. "
    "By contrasting early attrition cases (rapidly progressing phenotypes) against other decedents within each "
    "lineage, we identified epigenetic programs enriched in aggressive clinical trajectories."
)

INTRO_18 = (
    "Finally, we integrated exploratory mutation-histology interactions and expression module shifts to "
    "provide a somatic context for these lineage-resolved patterns. This approach aligns with the broader "
    "movement toward multi-omics synthesis [25] and refined clinical methylation markers [26] in esophageal "
    "oncology."
)

RESULTS_OPENING = (
    "We interrogated the TCGA ESCA (Firehose Legacy) cohort through a systematic multi-omics pipeline, "
    "structuring our findings to mirror the analysis sequence. We first quantified cross-histology "
    "transferability of prognostic models (Fig. 2; Table 1). We then characterized fast-death methylation "
    "programs within each histology (Fig. 3A–B; Tables 2–3) and linked selected loci to survival separation "
    "(Fig. 4). Finally, we added exploratory somatic context using mutation-associated module shifts and a "
    "mutation×histology interaction screen (Fig. 5A–B; Supplementary Table S7). External exploratory checks "
    "in GEO are summarized in Supplementary Tables S8–S9. An independent TCGA GDC package (Esophageal "
    "Adenocarcinoma (TCGA GDC, 2025)) provided a parallel robustness check for the transferability endpoint."
)


def _para_text(p: Paragraph) -> str:
    return "".join(t.text or "" for t in p._p.iter(W_T))


def _set_para_text(p: Paragraph, text: str) -> None:
    nodes = list(p._p.iter(W_T))
    if nodes:
        nodes[0].text = text
        for n in nodes[1:]:
            n.text = ""
    else:
        p.add_run(text)


def _replace_in_para(p: Paragraph, old: str, new: str) -> bool:
    for t_el in p._p.iter(W_T):
        if t_el.text and old in t_el.text:
            t_el.text = t_el.text.replace(old, new, 1)
            return True
    full = _para_text(p)
    if old not in full:
        return False
    _set_para_text(p, full.replace(old, new, 1))
    return True


def _find_para(doc: Document, pred) -> Paragraph:
    for p in doc.paragraphs:
        if pred(p):
            return p
    raise ValueError("Paragraph not found")


def _find_fig1_blocks(doc: Document) -> tuple[Paragraph | None, Paragraph | None]:
    img_p = cap_p = None
    for i, p in enumerate(doc.paragraphs):
        t = (p.text or "").strip()
        if t.startswith("Fig. 1.") and "Analysis pipeline" in t:
            cap_p = p
            prev = p._p.getprevious()
            if prev is not None and prev.tag.endswith("}p"):
                candidate = Paragraph(prev, p._parent)
                if candidate._p.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
                    img_p = candidate
            break
    return img_p, cap_p


def _move_fig1_to_methods(doc: Document) -> None:
    img_p, cap_p = _find_fig1_blocks(doc)
    if cap_p is None:
        return
    methods = _find_para(doc, lambda p: (_para_text(p).strip() == "Methods"))
    # Insert after Methods heading: caption block then image (image above caption)
    anchor = methods
    if img_p is not None:
        img_el = img_p._p
        img_el.getparent().remove(img_el)
        anchor._p.addnext(img_el)
        anchor = Paragraph(img_el, methods._parent)
    cap_el = cap_p._p
    cap_el.getparent().remove(cap_el)
    anchor._p.addnext(cap_el)


def _extract_references(doc: Document) -> dict[int, str]:
    refs: dict[int, str] = {}
    in_refs = False
    for p in doc.paragraphs:
        t = _para_text(p).strip()
        if t == "References":
            in_refs = True
            continue
        if in_refs:
            if t.startswith("Acknowledgements") or t.startswith("Author contributions"):
                break
            m = re.match(r"^(\d+)\.\s*(.+)", t)
            if m:
                refs[int(m.group(1))] = m.group(2).strip()
    return refs


def _rewrite_reference_list(doc: Document, old_refs: dict[int, str]) -> None:
    in_refs = False
    ref_paras: list[Paragraph] = []
    for p in doc.paragraphs:
        t = _para_text(p).strip()
        if t == "References":
            in_refs = True
            continue
        if in_refs:
            if t.startswith("Acknowledgements") or t.startswith("Author contributions"):
                break
            if re.match(r"^\d+\.", t):
                ref_paras.append(p)
    new_lines: list[str] = []
    for new_num in range(1, 31):
        old_num = next(k for k, v in OLD_TO_NEW.items() if v == new_num)
        body = old_refs.get(old_num, "")
        new_lines.append(f"{new_num}. {body}")
    for p, line in zip(ref_paras, new_lines):
        _set_para_text(p, line)


def _remap_citations_in_text(text: str) -> str:
    """Remap numeric citations only (refs 1–30); skip CI/statistical brackets."""

    def _repl_range(m: re.Match) -> str:
        a, b = int(m.group(1)), int(m.group(2))
        nums = [OLD_TO_NEW.get(n, n) for n in range(a, b + 1)]
        if len(nums) >= 2 and nums == list(range(nums[0], nums[-1] + 1)):
            return f"[{nums[0]}–{nums[-1]}]"
        return "[" + ",".join(str(n) for n in nums) + "]"

    def _repl_list(m: re.Match) -> str:
        parts = [int(x.strip()) for x in m.group(1).split(",")]
        nums = [OLD_TO_NEW.get(n, n) for n in parts]
        return "[" + ",".join(str(n) for n in nums) + "]"

    def _repl_single(m: re.Match) -> str:
        n = int(m.group(1))
        return f"[{OLD_TO_NEW.get(n, n)}]"

    # Only match reference-style brackets (numbers ≤30, no % or letters inside)
    text = re.sub(r"\[(\d{1,2})\s*[–\-—]\s*(\d{1,2})\](?![%\w])", _repl_range, text)
    text = re.sub(r"\[(\d{1,2}(?:\s*,\s*\d{1,2})+)\](?![%\w])", _repl_list, text)
    text = re.sub(r"\[(\d{1,2})\](?![%\w:\d])", _repl_single, text)
    return text


def _apply_citation_remap(doc: Document, skip_indices: set[int]) -> None:
    for i, p in enumerate(doc.paragraphs):
        if i in skip_indices:
            continue
        for t_el in p._p.iter(W_T):
            if t_el.text and "[" in t_el.text:
                new = _remap_citations_in_text(t_el.text)
                if new != t_el.text:
                    t_el.text = new


def _patch_discussion_results_cites(doc: Document) -> None:
    """Insert citations using OLD reference numbers; remapped later."""
    for p in doc.paragraphs:
        t = _para_text(p)
        if t.startswith("Given the transfer failure above, we next tested whether epigenetic"):
            if "[8]" not in t and "[27]" not in t:
                _replace_in_para(
                    p,
                    "are likewise lineage-specific.",
                    "are likewise lineage-specific, extending prior esophageal methylation reviews [8].",
                )
        if t.startswith("Genome-wide contrasts revealed divergent fast-death"):
            if "[9]" not in t and "[28]" not in t:
                _replace_in_para(
                    p,
                    "In ESCC, the companion volcano (Fig. 3A) is dominated",
                    "In ESCC, consistent with methylation–clinical association mapping in squamous disease [9], the companion volcano (Fig. 3A) is dominated",
                )
        if "assay-matched methylation validation remains the translational priority" in t:
            if "[20]" in t or "[8,20]" in t:
                pass  # will be remapped
            elif "[26,27]" not in t:
                _replace_in_para(
                    p,
                    "assay-matched methylation validation remains the translational priority.",
                    "assay-matched methylation validation remains the translational priority [20,8].",
                )


def _audit(doc: Document) -> list[str]:
    issues: list[str] = []
    refs_in_list: set[int] = set()
    in_refs = False
    for p in doc.paragraphs:
        t = _para_text(p).strip()
        if t == "References":
            in_refs = True
            continue
        if in_refs:
            if t.startswith("Acknowledgements"):
                break
            m = re.match(r"^(\d+)\.", t)
            if m:
                refs_in_list.add(int(m.group(1)))

    # Sequential first-appearance check
    first_seen: list[int] = []
    seen: set[int] = set()
    in_refs = False
    for p in doc.paragraphs:
        t = _para_text(p)
        if t.strip() == "References":
            in_refs = True
            continue
        if in_refs or re.match(r"^\d+\.\s", t.strip()):
            continue
        for chunk in re.findall(r"\[\d{1,2}(?:\s*[–\-—]\s*\d{1,2}|\s*,\s*\d{1,2})*\]", t):
            nums: list[int] = []
            rm = re.match(r"\[(\d+)\s*[–\-—]\s*(\d+)\]", chunk)
            if rm:
                nums = list(range(int(rm.group(1)), int(rm.group(2)) + 1))
            elif re.match(r"\[\d+\]", chunk):
                nums = [int(re.search(r"\d+", chunk).group())]
            else:
                nums = [int(x) for x in re.findall(r"\d+", chunk)]
            for n in nums:
                if n not in refs_in_list:
                    issues.append(f"Citation [{n}] not in reference list")
                if n not in seen:
                    seen.add(n)
                    first_seen.append(n)
    for i, n in enumerate(first_seen):
        if n != i + 1:
            issues.append(
                f"First-appearance order break: expected [{i + 1}] at position {i + 1}, got [{n}]"
            )
            break

    # Max 3 refs per sentence (Introduction + body)
    for p in doc.paragraphs:
        t = _para_text(p)
        if not t or t.strip() in ("Introduction", "Methods", "Results", "Discussion"):
            continue
        for sent in re.split(r"(?<=[.!?])\s+", t):
            nums: list[int] = []
            for rm in re.finditer(r"\[(\d{1,2})\s*[–\-—]\s*(\d{1,2})\]", sent):
                nums.extend(range(int(rm.group(1)), int(rm.group(2)) + 1))
            for lm in re.finditer(r"\[(\d{1,2}(?:\s*,\s*\d{1,2})+)\]", sent):
                nums.extend(int(x.strip()) for x in lm.group(1).split(","))
            for sm in re.finditer(r"\[(\d{1,2})\]", sent):
                nums.append(int(sm.group(1)))
            if len(set(nums)) > 3:
                issues.append(f">{3} refs in sentence: {sent[:80]}...")

    # Fig. 1 placement: should be in Methods, not Results opening
    for p in doc.paragraphs:
        if (_para_text(p).startswith("We interrogated") and "Fig. 1" in _para_text(p)):
            issues.append("Results opening still cites Fig. 1")
    methods = False
    fig1_in_methods = False
    fig1_in_results = False
    results = False
    for p in doc.paragraphs:
        t = _para_text(p).strip()
        if t == "Methods":
            methods = True
            results = False
        if t == "Results":
            results = True
            methods = False
        if t.startswith("Fig. 1.") and "Analysis pipeline" in t:
            if methods:
                fig1_in_methods = True
            if results:
                fig1_in_results = True
    if not fig1_in_methods:
        issues.append("Fig. 1 caption not found in Methods section")
    if fig1_in_results:
        issues.append("Fig. 1 still in Results section")

    methods_fig1 = any("Fig. 1" in _para_text(p) for p in doc.paragraphs if _para_text(p).startswith("The end-to-end analytical"))
    if not methods_fig1:
        issues.append("Methods body does not cite Fig. 1")

    return issues


def patch(docx_path: Path) -> None:
    backup = docx_path.with_name(docx_path.stem + ".docx.bak")
    shutil.copy2(docx_path, backup)

    doc = Document(str(docx_path))
    old_refs = _extract_references(doc)

    # Introduction rewrites (sequential refs 1–26; ≤3 refs per sentence)
    intro_map = {
        "Esophageal carcinoma is characterized": INTRO_13,
        "For EAC, the cellular origin": INTRO_14,
        "Throughout this manuscript, prognostic": INTRO_15,
        "This methodological oversight creates": INTRO_16,
        "In the present study, we prioritize": INTRO_17,
        "Finally, we integrated exploratory": INTRO_18,
    }
    intro_indices: set[int] = set()
    for i, p in enumerate(doc.paragraphs):
        t = _para_text(p)
        for prefix, new_text in intro_map.items():
            if t.startswith(prefix):
                _set_para_text(p, new_text)
                intro_indices.add(i)
                break

    # Results opening: remove Fig. 1
    for p in doc.paragraphs:
        if _para_text(p).startswith("We interrogated the TCGA ESCA"):
            _set_para_text(p, RESULTS_OPENING)
            break

    # Move Fig. 1 block to Methods
    _move_fig1_to_methods(doc)

    # Add Results/Discussion citations (old numbering) before global remap
    _patch_discussion_results_cites(doc)

    # Remap remaining old citation numbers → new order (skip rewritten intro)
    _apply_citation_remap(doc, skip_indices=intro_indices)

    # Reorder reference list 1–30
    _rewrite_reference_list(doc, old_refs)

    doc.save(str(docx_path))

    doc = Document(str(docx_path))
    issues = _audit(doc)
    print(f"Saved: {docx_path}")
    print(f"Backup: {backup}")
    if issues:
        print("Audit issues:")
        for x in issues:
            print(" ", x)
    else:
        print("Audit: Fig. 1 in Methods; refs 1–30 sequential; ≤3 refs/sentence.")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--docx", type=Path, default=DEFAULT_DOCX)
    args = ap.parse_args()
    patch(args.docx.resolve())


if __name__ == "__main__":
    main()
