"""Embed the five main figures below their legends without removing manuscript content."""
from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from PIL import Image

ROOT = Path("/Users/senol/Desktop/AML/Esophageal Carcinoma (TCGA, Firehose Legacy)")
PKG = ROOT / "Cancer Medicine submission_package"
FIGURES = PKG / "Figures"
TARGETS = [PKG / "02_Manuscript.docx", PKG / "02_Manuscript_clean.docx"]
MAX_WIDTH_IN = 6.4
MAX_HEIGHT_IN = 6.3


def find_legend(doc: Document, number: int) -> Paragraph:
    heading_seen = False
    prefix = f"Figure {number}."
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == "Figure legends":
            heading_seen = True
            continue
        if heading_seen and text.startswith(prefix):
            return paragraph
    raise ValueError(f"Legend not found: {prefix}")


def paragraph_has_drawing(paragraph: Paragraph) -> bool:
    return bool(paragraph._p.findall(".//" + qn("a:blip")))


def paragraph_after(paragraph: Paragraph) -> Paragraph | None:
    element = paragraph._p.getnext()
    if element is None or element.tag != qn("w:p"):
        return None
    return Paragraph(element, paragraph._parent)


def insert_paragraph_after(paragraph: Paragraph) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    return Paragraph(element, paragraph._parent)


def fitted_size(path: Path) -> tuple[float, float]:
    with Image.open(path) as image:
        width_px, height_px = image.size
    scale = min(MAX_WIDTH_IN / width_px, MAX_HEIGHT_IN / height_px)
    return width_px * scale, height_px * scale


def set_alt_text(paragraph: Paragraph, text: str) -> None:
    doc_properties = paragraph._p.findall(".//" + qn("wp:docPr"))
    if doc_properties:
        doc_properties[0].set("title", text)
        doc_properties[0].set("descr", text)


def embed_figures(path: Path) -> None:
    doc = Document(str(path))
    if doc.inline_shapes:
        raise RuntimeError(
            f"{path.name} already contains {len(doc.inline_shapes)} inline image(s); "
            "stopped to avoid deleting or duplicating existing content."
        )

    for number in range(1, 6):
        image_path = FIGURES / f"Figure_{number}.png"
        if not image_path.is_file():
            raise FileNotFoundError(image_path)

        legend = find_legend(doc, number)
        legend.paragraph_format.page_break_before = True
        legend.paragraph_format.keep_with_next = True

        image_paragraph = paragraph_after(legend)
        if image_paragraph is not None and paragraph_has_drawing(image_paragraph):
            raise RuntimeError(
                f"An image already follows Figure {number}; stopped without replacing it."
            )
        if image_paragraph is None or image_paragraph.text.strip():
            image_paragraph = insert_paragraph_after(legend)

        width, height = fitted_size(image_path)
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.paragraph_format.keep_together = True
        run = image_paragraph.add_run()
        run.add_picture(str(image_path), width=Inches(width), height=Inches(height))
        set_alt_text(
            image_paragraph,
            f"Figure {number}. See the preceding manuscript legend for a complete description.",
        )

    doc.save(str(path))


def main() -> None:
    targets = [target for target in TARGETS if target.is_file()]
    if not targets:
        raise FileNotFoundError("No manuscript file was found.")
    stamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    backup_dir = ROOT / f"manuscript_before_figure_reinsertion_{stamp}"
    backup_dir.mkdir()
    for target in targets:
        shutil.copy2(target, backup_dir / target.name)

    for target in targets:
        embed_figures(target)

    for target in targets:
        check = Document(str(target))
        if len(check.inline_shapes) != 5:
            raise RuntimeError(f"Expected five embedded figures in {target.name}")

    print(f"Backup: {backup_dir}")
    print("Embedded Figures 1–5 below their matching legends in: " + ", ".join(p.name for p in targets))


if __name__ == "__main__":
    main()
