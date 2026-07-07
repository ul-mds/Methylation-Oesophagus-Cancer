"""
Remove supplementary figure blocks that duplicate main-text Figures 1–6 (same PNGs).

Keeps only unique panels (former S7–S8: mediation volcano and top mediation summary),
renumbering them to Supplementary Fig. S1 and S2. Updates the Supplementary Materials intro line.

Safe to re-run: skips removal if no S1–S6 caption blocks remain.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from lxml import etree

DOC_NAME = "TCGA_ESCA_ESCC_landmark60_full_scientific_manuscript.docx"


def _p_has_image(p) -> bool:
    return b'r:embed="' in etree.tostring(p._element)


def _supplementary_fig_heading_index(doc: Document) -> int | None:
    for i, p in enumerate(doc.paragraphs):
        t = (p.text or "").strip()
        if t == "Supplementary Figures" and p.style.name.startswith("Heading"):
            return i
    return None


# Captions that duplicated main-text Figures 1–6 (same PNGs under Plots/).
_DUPLICATE_SOURCE_MARKERS = (
    "Plots/01_cross_apply_cindex",
    "Plots/02_fastdeath_methylation_volcano_panel",
    "Plots/03_fastdeath_top_hypo_hyper_panel",
    "Plots/04_KM_methylation_60mo_panel",
    "Plots/05_mutation_module_heatmap",
    "Plots/06_interaction_screen",
)


def _remove_duplicate_blocks(doc: Document) -> int:
    h_idx = _supplementary_fig_heading_index(doc)
    if h_idx is None:
        raise SystemExit("Supplementary Figures heading not found.")
    removed = 0
    i = h_idx + 1
    while i < len(doc.paragraphs) - 1:
        p_img = doc.paragraphs[i]
        p_cap = doc.paragraphs[i + 1]
        if not _p_has_image(p_img):
            i += 1
            continue
        tc = (p_cap.text or "").strip()
        if not tc.startswith("Supplementary Fig."):
            i += 1
            continue
        if not any(marker in tc for marker in _DUPLICATE_SOURCE_MARKERS):
            i += 1
            continue
        p_cap._element.getparent().remove(p_cap._element)
        p_img._element.getparent().remove(p_img._element)
        removed += 1
    return removed


def _renumber_mediation_figs(doc: Document) -> None:
    for p in doc.paragraphs:
        t = p.text or ""
        if not t.strip().startswith("Supplementary Fig."):
            continue
        t2 = t.replace("Supplementary Fig. S7.", "Supplementary Fig. S1.", 1)
        t2 = t2.replace("Supplementary Fig. S8.", "Supplementary Fig. S2.", 1)
        if t2 != t:
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            p.add_run(t2)


def _update_supp_intro(doc: Document) -> None:
    new_intro = (
        "Raw and processed outputs are stored alongside this manuscript under result/ (tabular data) and "
        "Plots/ (figure images). Supplementary Figures S1–S2 provide mediation-focused panels that are not "
        "reproduced in the main article; main-text Figures 1–6 and Tables 1–7 are cited in the Results."
    )
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith("Raw and processed outputs are stored alongside this manuscript"):
            if "Supplementary Figures S1–S2" in p.text:
                return
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            p.add_run(new_intro)
            return


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    path = root / DOC_NAME
    doc = Document(str(path))
    n = _remove_duplicate_blocks(doc)
    _renumber_mediation_figs(doc)
    _update_supp_intro(doc)
    doc.save(str(path))
    print(f"Removed {n} duplicate supplementary figure pair(s); renumbered mediation to S1–S2. Saved {path}")


if __name__ == "__main__":
    main()
