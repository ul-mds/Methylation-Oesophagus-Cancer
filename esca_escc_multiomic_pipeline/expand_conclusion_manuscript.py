"""Replace the single-sentence Conclusion with two stronger paragraphs (idempotent)."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


def _txt(p) -> str:
    return "".join(t.text or "" for t in p._p.iter(qn("w:t"))).strip()


def main():
    path = Path(__file__).resolve().parents[1] / "TCGA_ESCA_ESCC_landmark60_full_scientific_manuscript.docx"
    doc = Document(str(path))

    for i, p in enumerate(doc.paragraphs):
        if _txt(p) == "Conclusion" and i + 1 < len(doc.paragraphs):
            p73 = doc.paragraphs[i + 1]
            nxt = doc.paragraphs[i + 2] if i + 2 < len(doc.paragraphs) else None
            if _txt(p73).startswith("In TCGA esophageal carcinoma") and nxt and _txt(nxt).startswith(
                "Fast-death anchored methylation"
            ):
                if "That collapse is not a modeling artifact" in _txt(p73):
                    print("Conclusion already expanded:", path)
                    return
            break
    else:
        raise RuntimeError("Could not find Conclusion heading.")

    p73.text = (
        "In TCGA esophageal carcinoma, prognostic models are strongly histology-specific and do not transfer "
        "between EAC and ESCC (Figure 1; Table 1). That collapse is not a modeling artifact: it implies that "
        "“esophageal cancer” cannot be treated as a single prognostic disease for multi-omic risk when coefficients "
        "trained in one lineage are applied to the other. A pooled, histology-aware baseline recovers useful "
        "discrimination in the combined cohort, underscoring that shared platforms are compatible with prognosis "
        "only when histology is respected in the hazard structure—not when signatures are silently assumed to "
        "commute across lineages."
    )

    nxt = doc.paragraphs[i + 2] if i + 2 < len(doc.paragraphs) else None
    second = (
        "Fast-death anchored methylation analyses reveal distinct histology-resolved epigenetic programs "
        "(Figures 2–3; Tables 2–5) and clinically interpretable 60-month survival separation (Figure 4; Table 6), "
        "while mutation-layer follow-up (Figures 5–6; Table 7) frames mechanistic next steps. Together, these layers "
        "argue for reporting esophageal omics findings with explicit histology context, for validating methylation "
        "and somatic hypotheses in external cohorts matched on lineage and assay, and for interpreting medium-term "
        "risk within the same landmark horizon used for discovery—so that aggressive epigenetic biology, "
        "transferable model failure, and exploratory genomics are read as a single, coherent case study rather than "
        "disconnected signatures."
    )
    if nxt and _txt(nxt).startswith("Fast-death anchored"):
        nxt.text = second
    else:
        new_p = OxmlElement("w:p")
        p73._element.addnext(new_p)
        Paragraph(new_p, p73._parent).text = second

    doc.save(str(path))
    print("Conclusion expanded:", path)


if __name__ == "__main__":
    main()
