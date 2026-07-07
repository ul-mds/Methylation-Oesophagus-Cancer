"""
Append peer-reviewed literature (2022–2026) to References and weave bracketed
citations [n] into Introduction, Methods, and Discussion. Idempotent on reference
lines: skips if 'Lyu Q' already present in References.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def _txt(p) -> str:
    return "".join(t.text or "" for t in p._p.iter(qn("w:t"))).strip()


def _find_paragraph(doc: Document, startswith: str):
    for p in doc.paragraphs:
        if _txt(p).startswith(startswith):
            return p
    raise KeyError(startswith)


def _refs_already_extended(doc: Document) -> bool:
    for p in doc.paragraphs:
        t = _txt(p)
        if "Lyu Q" in t and "12885-025-15090" in t:
            return True
    return False


def main():
    path = Path(__file__).resolve().parents[1] / "TCGA_ESCA_ESCC_landmark60_full_scientific_manuscript.docx"
    doc = Document(str(path))

    for p in doc.paragraphs:
        t = _txt(p)
        if t.startswith("7. Pinto R, Sjurgard"):
            p.text = (
                "7. Pinto R, et al. Molecular stratification of esophageal adenocarcinoma: implications for "
                "prognosis and treatment strategy. Oncogene. 2025. doi:10.1038/s41388-025-03650-3"
            )
            break

    replacements: list[tuple[str, str]] = [
        (
            "The TCGA esophageal carcinoma cohort combines esophageal adenocarcinoma (EAC) and esophageal squamous cell carcinoma (ESCC), yet integrated genomics has established that these histologies represent distinct molecular diseases rather than interchangeable subtypes of a single entity. A practical consequence—often unstated in prognostic omics studies—is that risk models optimized in one histology may lack external validity when applied to the other, even within the same data resource. We therefore structured the Results around a single logical arc: (i) quantify cross-histology transferability of multi-feature Cox scores under a 60-month landmark overall survival endpoint (Figure 1; Table 1); (ii) ask whether clinically extreme courses (early death among decedents) exhibit histology-specific DNA methylation programs (Figures 2–3; Tables 2–5); (iii) relate prioritized probes to interpretable survival separation at the same clinical horizon (Figure 4; Table 6); and (iv) situate these epigenetic patterns alongside exploratory somatic analyses in which histology may modify mutation–outcome associations and downstream expression modules (Figures 5–6; Table 7). For readability, the empirical blocks below follow this same order under four subsection headings, with each table and figure placed in the substudy to which it primarily belongs.",
            "The TCGA esophageal carcinoma cohort combines esophageal adenocarcinoma (EAC) and esophageal squamous cell carcinoma (ESCC), yet integrated genomics has established that these histologies represent distinct molecular diseases rather than interchangeable subtypes of a single entity [1,6,7]. A practical consequence—often unstated in prognostic omics studies—is that risk models optimized in one histology may lack external validity when applied to the other, even within the same data resource. We therefore structured the Results around a single logical arc: (i) quantify cross-histology transferability of multi-feature Cox scores under a 60-month landmark overall survival endpoint (Figure 1; Table 1); (ii) ask whether clinically extreme courses (early death among decedents) exhibit histology-specific DNA methylation programs (Figures 2–3; Tables 2–5); (iii) relate prioritized probes to interpretable survival separation at the same clinical horizon (Figure 4; Table 6); and (iv) situate these epigenetic patterns alongside exploratory somatic analyses in which histology may modify mutation–outcome associations and downstream expression modules (Figures 5–6; Table 7). For readability, the empirical blocks below follow this same order under four subsection headings, with each table and figure placed in the substudy to which it primarily belongs. Parallel recent work continues to refine methylation-based risk models, liquid-biopsy-oriented profiling narratives, and consensus ESCC subtypes that jointly motivate histology-resolved reporting [5,8–10].",
        ),
        (
            "Esophageal carcinoma is dominated by two major histologies—adenocarcinoma (EAC) and squamous cell carcinoma (ESCC)—that differ in epidemiology, anatomic distribution, and therapeutic context. TCGA’s integrated analysis emphasized that these histologies are not merely variants of one disease but represent distinct molecular entities. Yet in downstream prognostic modeling, it remains common to either pool them or to develop histology-specific signatures without asking a basic translational question: if a model captures general “esophageal cancer” prognosis, should it not generalize across histologies?",
            "Esophageal carcinoma is dominated by two major histologies—adenocarcinoma (EAC) and squamous cell carcinoma (ESCC)—that differ in epidemiology, anatomic distribution, and therapeutic context. TCGA’s integrated analysis emphasized that these histologies are not merely variants of one disease but represent distinct molecular entities [1]. Subsequent integrative analyses continue to resolve transcriptomic and pathway divergence between EAC and ESCC in public cohorts [6], while recent consensus subtyping in ESCC consolidates prior taxonomies and ties subtype structure to multi-omic traits including methylation and outcome [10]. EAC-focused stratification reviews likewise frame prognosis and treatment decisions in histology-specific molecular terms [7]. Yet in downstream prognostic modeling, it remains common to either pool them or to develop histology-specific signatures without asking a basic translational question: if a model captures general “esophageal cancer” prognosis, should it not generalize across histologies?",
        ),
        (
            "The absence of an explicit transferability test creates an interpretability gap. A prognostic model that performs well within one histology may simply be learning biology that is specific to that histology, making cross-histology deployment inappropriate. Conversely, a model that truly reflects shared aggressiveness biology might retain performance when applied across histologies. We therefore foregrounded transferability as a primary endpoint: we trained EAC-only and ESCC-only prognostic models and measured their performance both within histology and when applied to the other histology, with concordance summarized in Figure 1 and Table 1.",
            "The absence of an explicit transferability test creates an interpretability gap. A prognostic model that performs well within one histology may simply be learning biology that is specific to that histology, making cross-histology deployment inappropriate. Conversely, a model that truly reflects shared aggressiveness biology might retain performance when applied across histologies. We therefore foregrounded transferability as a primary endpoint: we trained EAC-only and ESCC-only prognostic models and measured their performance both within histology and when applied to the other histology, with concordance summarized in Figure 1 and Table 1. This emphasis complements histology-resolved omics contrasts that document how differently EAC and ESCC organize molecular risk [6,7].",
        ),
        (
            "In parallel, DNA methylation is attractive for prognosis research because it is stable, captures both cell-of-origin and tumor state, and can be measured robustly. Numerous methylation-based prognostic signatures have been reported for esophageal cancer, particularly ESCC. Rather than pursuing another signature for its own sake, we anchored methylation discovery to an intuitive clinical phenotype—fast death among decedents—and assessed methylation programs separately in EAC and ESCC within a 60-month landmark OS window, visualized as paired histology volcanoes and bar panels (Figures 2–3) with full locus-level ranks in Tables 2–5. We then linked prioritized loci to 60-month Kaplan–Meier separation (Figure 4) and, in EAC, to univariate Cox rankings (Table 6). This framing aims to produce results that are both biologically interpretable and clinically communicable.",
            "In parallel, DNA methylation is attractive for prognosis research because it is stable, captures both cell-of-origin and tumor state, and can be measured robustly. Numerous methylation-based prognostic signatures have been reported for esophageal cancer, particularly ESCC, including TCGA-grounded multi-omic methylation surveys and methylation–clinical association studies [2,4,5,9]. Reviews summarize how sequencing platforms now enable genome-wide methylation profiling for detection and risk stratification in esophageal cancer [8]. Rather than pursuing another signature for its own sake, we anchored methylation discovery to an intuitive clinical phenotype—fast death among decedents—and assessed methylation programs separately in EAC and ESCC within a 60-month landmark OS window, visualized as paired histology volcanoes and bar panels (Figures 2–3) with full locus-level ranks in Tables 2–5. We then linked prioritized loci to 60-month Kaplan–Meier separation (Figure 4) and, in EAC, to univariate Cox rankings (Table 6). This framing aims to produce results that are both biologically interpretable and clinically communicable.",
        ),
        (
            "We used the cBioPortal-format “Esophageal Carcinoma (TCGA, Firehose Legacy)” study package. Histology was determined from the sample-level clinical file (ONCOTREE_CODE: ESCA vs ESCC). Overall survival time (OS_MONTHS) and status (OS_STATUS) were obtained from the patient-level clinical file and mapped to samples. Analyses were restricted to primary tumors and to samples with aligned clinical and omics availability.",
            "We used the cBioPortal-format “Esophageal Carcinoma (TCGA, Firehose Legacy)” study package [1]. Histology was determined from the sample-level clinical file (ONCOTREE_CODE: ESCA vs ESCC). Overall survival time (OS_MONTHS) and status (OS_STATUS) were obtained from the patient-level clinical file and mapped to samples. Analyses were restricted to primary tumors and to samples with aligned clinical and omics availability. Our use of this harmonized repository parallels other recent TCGA-based methylation integrative analyses in ESCC [5,9].",
        ),
        (
            "We trained Cox proportional hazards models separately in EAC and ESCC (histology-specific models) and evaluated them both within-histology and by cross-applying them to the opposite histology. Performance was quantified by the concordance index on OS60. We also trained a pooled Cox model with a histology-stratified baseline hazard to provide a combined, histology-aware comparator.",
            "We trained Cox proportional hazards models separately in EAC and ESCC (histology-specific models) and evaluated them both within-histology and by cross-applying them to the opposite histology. Performance was quantified by the concordance index on OS60. We also trained a pooled Cox model with a histology-stratified baseline hazard to provide a combined, histology-aware comparator. The explicit cross-apply contrast extends histology-first prognostic perspectives emphasized in recent differential and stratification studies of esophageal carcinoma [6,7].",
        ),
        (
            "Within each histology, we defined a fast-death phenotype among decedents as those with the shortest OS (top 20% of deaths; minimum 10). We compared HM450 beta values between fast-death samples and all other samples in the same histology, reporting Δβ=mean(β_fast)−mean(β_other). Multiple testing was controlled using Benjamini–Hochberg FDR (q_diff).",
            "Within each histology, we defined a fast-death phenotype among decedents as those with the shortest OS (top 20% of deaths; minimum 10). We compared HM450 beta values between fast-death samples and all other samples in the same histology, reporting Δβ=mean(β_fast)−mean(β_other). Multiple testing was controlled using Benjamini–Hochberg FDR (q_diff). Conceptually, this contrasts aggressive clinical trajectories within histology rather than relying solely on pan-tumor versus normal methylation contrasts [4,5,9].",
        ),
        (
            "The primary finding is straightforward but often left implicit: in TCGA esophageal carcinoma, prognostic signal is not shared uniformly across histologies. Models trained within EAC and ESCC each performed well within their own histology, yet cross-application failed—Figure 1 and Table 1 make this collapse explicit for readers who prefer graphical or tabular summaries. This transferability failure is not a technical nuisance; it is a biological result that argues against treating EAC and ESCC as a single prognostic entity and helps explain why many published signatures do not replicate when moved across clinically heterogeneous cohorts.",
            "The primary finding is straightforward but often left implicit: in TCGA esophageal carcinoma, prognostic signal is not shared uniformly across histologies. Models trained within EAC and ESCC each performed well within their own histology, yet cross-application failed—Figure 1 and Table 1 make this collapse explicit for readers who prefer graphical or tabular summaries. This transferability failure is not a technical nuisance; it is a biological result that argues against treating EAC and ESCC as a single prognostic entity and helps explain why many published signatures do not replicate when moved across clinically heterogeneous cohorts [6,7].",
        ),
        (
            "A second contribution is the clinically anchored epigenetic analysis. Rather than optimizing a high-dimensional signature, we asked what methylation programs distinguish the most aggressive clinical courses—fast death among decedents—within each histology. Figures 2–3, together with Tables 2–5, show that the leading hypo- and hypermethylation candidates differ sharply between histologies, consistent with distinct cell-of-origin and pathway wiring. Figure 4 then demonstrates that carefully chosen probes can still stratify 60-month OS within each histology, with EAC priorities guided by Table 6. These programs can be communicated as testable hypotheses (e.g., SHISA3-associated epigenetic state in aggressive ESCC) and evaluated in external cohorts with methylation assays.",
            "A second contribution is the clinically anchored epigenetic analysis. Rather than optimizing a high-dimensional signature, we asked what methylation programs distinguish the most aggressive clinical courses—fast death among decedents—within each histology. Figures 2–3, together with Tables 2–5, show that the leading hypo- and hypermethylation candidates differ sharply between histologies, consistent with distinct cell-of-origin and pathway wiring, in line with multi-omic methylation landscapes and subtype-resolved epigenetic heterogeneity described in recent ESCC work [5,9,10]. Figure 4 then demonstrates that carefully chosen probes can still stratify 60-month OS within each histology, with EAC priorities guided by Table 6. These programs can be communicated as testable hypotheses (e.g., SHISA3-associated epigenetic state in aggressive ESCC) and evaluated in external cohorts with methylation assays [8].",
        ),
        (
            "Our use of a 60-month landmark OS window serves two purposes: it matches common clinical decision horizons and reduces sensitivity to late follow-up differences. That said, landmarking can reduce power for late events, and ESCC contributes fewer OS events within the window than EAC, which limits precision for Kaplan–Meier displays such as Figure 4 (lower row) even when fast-death methylation contrasts are strong (Figures 2–3; Tables 2–3). Table 6 summarizes EAC-only methylation–OS Cox screening on OS60. The mutation×histology interaction results should be interpreted cautiously; interaction testing is power-intensive, and we deliberately positioned Figure 6 and Table 7 as exploratory. Their value lies in connecting candidate alterations to downstream pathway context via Figure 5, which helps prioritize targets for future validation rather than establishing definitive effect modification.",
            "Our use of a 60-month landmark OS window serves two purposes: it matches common clinical decision horizons and reduces sensitivity to late follow-up differences. That said, landmarking can reduce power for late events, and ESCC contributes fewer OS events within the window than EAC, which limits precision for Kaplan–Meier displays such as Figure 4 (lower row) even when fast-death methylation contrasts are strong (Figures 2–3; Tables 2–3). Table 6 summarizes EAC-only methylation–OS Cox screening on OS60. The mutation×histology interaction results should be interpreted cautiously; interaction testing is power-intensive, and we deliberately positioned Figure 6 and Table 7 as exploratory. Their value lies in connecting candidate alterations to downstream pathway context via Figure 5, which helps prioritize targets for future validation rather than establishing definitive effect modification. Independently, large ESCC consortia now relate consensus molecular subtypes to prognosis and treatment response, including epigenetic readouts at the subtype level [10].",
        ),
        (
            "Limitations include reliance on a cBioPortal study package without a full HM450 probe annotation table within the package (gene tokens reflect probe naming), the lack of external validation cohorts in the current analysis environment, and the possibility that some individual loci have prior prognostic reports in other cancers or in esophageal cancer. Figures 1–6 and Tables 1–7 should therefore be read as an internally consistent TCGA case study: our novelty claim is not that any single gene is unprecedented, but that the integrated design—histology-resolved fast-death epigenetics plus explicit cross-histology transferability testing within a standardized 60-month framework—fills a practical gap in the prognostic literature.",
            "Limitations include reliance on a cBioPortal study package without a full HM450 probe annotation table within the package (gene tokens reflect probe naming), the lack of external validation cohorts in the current analysis environment, and the possibility that some individual loci have prior prognostic reports in other cancers or in esophageal cancer. Figures 1–6 and Tables 1–7 should therefore be read as an internally consistent TCGA case study: our novelty claim is not that any single gene is unprecedented, but that the integrated design—histology-resolved fast-death epigenetics plus explicit cross-histology transferability testing within a standardized 60-month framework—fills a practical gap emphasized by recent histology-resolved omics syntheses and methylation-focused reviews [6–10].",
        ),
    ]

    for old, new in replacements:
        found = False
        for p in doc.paragraphs:
            if _txt(p) == old:
                p.text = new
                found = True
                break
        if not found:
            # allow re-run after partial edit: match if new already applied
            for p in doc.paragraphs:
                if _txt(p) == new:
                    found = True
                    break
            if not found:
                raise RuntimeError("Paragraph text mismatch; manuscript may have been edited manually.")

    if not _refs_already_extended(doc):
        new_refs = [
            "6. Lyu Q, Chai Y, Chen W, Chen Y, Li Y. Integrated bioinformatics analysis of differences between EAC and ESCC. BMC Cancer. 2025;25:1668. doi:10.1186/s12885-025-15090-z",
            "7. Pinto R, et al. Molecular stratification of esophageal adenocarcinoma: implications for prognosis and treatment strategy. Oncogene. 2025. doi:10.1038/s41388-025-03650-3",
            "8. Tang J, Zhang W, Dai L, Wu D, Qiu J, Liu J. DNA methylation in esophageal cancer: technological advances and early detection clinical applications. Front Oncol. 2025;15:1543190. doi:10.3389/fonc.2025.1543190",
            "9. Wang G, Lv M, Wang Y. Analysis of the associations between DNA methylation and clinical features reveals the potential oncogenic effect of CFAP52 in esophageal squamous cell cancer. Cancer Cell Int. 2025;25:394. doi:10.1186/s12935-025-03934-4",
            "10. Cui H, Zhu Z, Xu E, et al. The consensus molecular subtypes of esophageal squamous cell carcinoma. Signal Transduct Target Ther. 2026;11:65. doi:10.1038/s41392-026-02577-9",
        ]
        # Fix ref 5 to full citation (idempotent if already Xi et al.)
        xi_line = (
            "5. Xi Y, Lin Y, Guo W, et al. Multi-omic characterization of genome-wide abnormal DNA methylation "
            "reveals diagnostic and prognostic markers for esophageal squamous-cell carcinoma. "
            "Signal Transduct Target Ther. 2022;7:53. doi:10.1038/s41392-022-00873-8"
        )
        for p in doc.paragraphs:
            t = _txt(p)
            if t.startswith("5. Multi-omic characterization") or t.startswith("5. Xi Y"):
                p.text = xi_line
                break
        else:
            raise RuntimeError("Could not find reference 5 to replace.")

        from docx.oxml import OxmlElement
        from docx.text.paragraph import Paragraph as DocxParagraph

        insert_after = None
        for p in doc.paragraphs:
            if _txt(p).startswith("5. Xi Y"):
                insert_after = p
                break
        if insert_after is None:
            raise RuntimeError("Missing updated ref 5.")

        prev_el = insert_after._element
        for line in new_refs:
            new_p = OxmlElement("w:p")
            prev_el.addnext(new_p)
            DocxParagraph(new_p, insert_after._parent).add_run(line)
            prev_el = new_p

    doc.save(str(path))
    print("Literature (2022–2026) merged:", path)


if __name__ == "__main__":
    main()
