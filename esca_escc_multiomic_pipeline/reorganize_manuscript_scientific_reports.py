"""
Scientific Reports–oriented manuscript cleanup:

- Introduction: remove figure/table citations; keep literature-driven framing.
- Reorder body: Introduction → Methods → Results → Discussion → References → end matter.
- Discussion: remove display-item citations (Results hold primary citations); drop redundant summary paras.
- Remove empty Heading 2 and duplicate conclusion blocks inside the old Discussion pocket.
- Strip "(Main-text Fig. …)" cross-references from Supplementary Figure captions.

Run from study root; rewrites TCGA_ESCA_ESCC_landmark60_full_scientific_manuscript.docx in place.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

DOC_NAME = "TCGA_ESCA_ESCC_landmark60_full_scientific_manuscript.docx"


def _para_text(el) -> str:
    parts: list[str] = []
    for t in el.iter(qn("w:t")):
        if t.text:
            parts.append(t.text)
    return "".join(parts)


def _elements_between_paragraphs(doc: Document, p_start: int, p_end: int) -> list:
    """Inclusive paragraph index range; includes any w:tbl between those paragraphs."""
    body = doc.element.body
    start_el = None
    end_el = None
    pi = -1
    for el in body:
        if el.tag == qn("w:p"):
            pi += 1
            if pi == p_start:
                start_el = el
            if pi == p_end:
                end_el = el
                break
    if start_el is None or end_el is None:
        raise ValueError(f"Paragraph range {p_start}-{p_end} not found")
    out: list = []
    capturing = False
    for el in body:
        if el is start_el:
            capturing = True
        if capturing:
            out.append(el)
        if el is end_el:
            break
    return out


def _elements_from_paragraph_to_end(doc: Document, p_start: int) -> list:
    body = doc.element.body
    start_el = None
    pi = -1
    for el in body:
        if el.tag == qn("w:p"):
            pi += 1
            if pi == p_start:
                start_el = el
                break
    if start_el is None:
        raise ValueError(f"Paragraph {p_start} not found")
    out: list = []
    capturing = False
    for el in body:
        if el is start_el:
            capturing = True
        if el.tag == qn("w:sectPr"):
            break
        if capturing:
            out.append(el)
    return out


def _rebuild_body(doc: Document, ordered_blocks: list[list]) -> None:
    body = doc.element.body
    sect_pr = None
    for el in list(body):
        if el.tag == qn("w:sectPr"):
            sect_pr = el
        body.remove(el)
    for block in ordered_blocks:
        for el in block:
            body.append(el)
    if sect_pr is not None:
        body.append(sect_pr)


def _strip_maintext_refs_from_supplementary(doc: Document) -> None:
    pat = re.compile(r"\s*\(?Main-text\s+[Ff]ig\.\s*[^)]+\)\.?", re.I)
    for p in doc.paragraphs:
        t = p.text or ""
        if not t.strip().startswith("Supplementary Fig."):
            continue
        nt = pat.sub("", t)
        nt = re.sub(r"\s{2,}", " ", nt).strip()
        if nt != t:
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            p.add_run(nt)


def _edit_introduction_and_discussion(doc: Document) -> None:
    intro_p10 = (
        "The absence of an explicit transferability test creates an interpretability gap. A prognostic model that "
        "performs well within one histology may simply be learning biology that is specific to that histology, making "
        "cross-histology deployment inappropriate. Conversely, a model that truly reflects shared aggressiveness biology "
        "might retain performance when applied across histologies. We therefore foregrounded transferability as a "
        "primary endpoint: we trained EAC-only and ESCC-only prognostic models and measured their performance both "
        "within histology and when applied to the other histology, using concordance indices that summarise agreement "
        "between predicted risk ordering and observed outcomes. This emphasis complements histology-resolved omics "
        "contrasts that document how differently EAC and ESCC organize molecular risk [6,7]."
    )
    intro_p11 = (
        "In parallel, DNA methylation is attractive for prognosis research because it is stable, captures both "
        "cell-of-origin and tumor state, and can be measured robustly. Numerous methylation-based prognostic signatures "
        "have been reported for esophageal cancer, particularly ESCC, including TCGA-grounded multi-omic methylation "
        "surveys and methylation–clinical association studies [2,4,5,9,11,17,20]. Reviews summarize how sequencing "
        "platforms now enable genome-wide methylation profiling for detection and risk stratification in esophageal "
        "cancer [8]. Rather than pursuing another signature for its own sake, we anchored methylation discovery to an "
        "intuitive clinical phenotype—fast death among decedents—and assessed methylation programs separately in EAC and "
        "ESCC within a pre-specified follow-up window. Survival contrasts, Cox screening, and mechanistic follow-up are "
        "presented in the Results. This framing aims to produce findings that are both biologically interpretable and "
        "clinically communicable."
    )
    disc_p69 = (
        "The primary finding is straightforward but often left implicit: in TCGA esophageal carcinoma, prognostic "
        "signal is not shared uniformly across histologies. Models trained within EAC and ESCC each performed well "
        "within their own histology, yet cross-application failed, as quantified in the Results. This transferability "
        "failure is not a technical nuisance; it is a biological result that argues against treating EAC and ESCC as a "
        "single prognostic entity and helps explain why many published signatures do not replicate when moved across "
        "clinically heterogeneous cohorts [6,7,15,16]."
    )
    disc_p70 = (
        "A second contribution is the clinically anchored epigenetic analysis. Rather than optimizing a "
        "high-dimensional signature, we asked what methylation programs distinguish the most aggressive clinical "
        "courses—fast death among decedents—within each histology. As shown in the Results, the leading hypo- and "
        "hypermethylation candidates differ sharply between histologies, consistent with distinct cell-of-origin and "
        "pathway wiring, in line with multi-omic methylation landscapes and subtype-resolved epigenetic heterogeneity "
        "described in recent ESCC work [5,9,10]. Probe-level survival contrasts and EAC Cox screening support "
        "interpretable stratification within each lineage. These programs can be communicated as testable hypotheses "
        "(e.g., SHISA3-associated epigenetic state in aggressive ESCC) and evaluated in external cohorts with methylation "
        "assays [8]."
    )
    disc_p71 = (
        "Our use of a 60-month OS window serves two purposes: it matches common clinical decision horizons and reduces "
        "sensitivity to late follow-up differences. That said, capping follow-up can reduce power for late events, and "
        "ESCC contributes fewer OS events within the window than EAC, which limits precision for some survival "
        "displays even when fast-death methylation contrasts are strong. Methylation–OS screening and mutation×histology "
        "interaction testing are therefore interpreted as exploratory and event-limited; their value lies in connecting "
        "candidate alterations to pathway context for prioritisation in external validation rather than in definitive "
        "effect modification. Independently, large ESCC consortia now relate consensus molecular subtypes to prognosis "
        "and treatment response, including epigenetic readouts at the subtype level [10], while EAC-focused genomic "
        "architecture studies motivate separate mechanistic priors for amplification-driven disease [13]."
    )
    disc_p72 = (
        "Limitations include reliance on a cBioPortal study package without a full HM450 probe annotation table within "
        "the package (gene tokens reflect probe naming), the lack of external validation cohorts in the current "
        "analysis environment, and the possibility that some individual loci have prior prognostic reports in other "
        "cancers or in esophageal cancer. The reported outputs should therefore be read as an internally consistent "
        "TCGA case study: the novelty claim is not that any single gene is unprecedented, but that the integrated "
        "design—histology-resolved fast-death epigenetics plus explicit cross-histology transferability testing within "
        "a standardized follow-up framework—fills a practical gap emphasized by recent histology-resolved omics "
        "syntheses and methylation-focused reviews [6–20]."
    )
    methods_p77_prefix = (
        "We used the cBioPortal-format “Esophageal Carcinoma (TCGA, Firehose Legacy)” study package [1]. "
        "For each sample we retained Illumina HM450 DNA methylation beta values, RNA-seq–derived expression features "
        "(z-scores as distributed in the study matrices), and binary somatic mutation calls where available, alongside "
        "clinical fields required for histology assignment and survival. "
    )

    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith("The absence of an explicit transferability test"):
            _set_paragraph_plain(p, intro_p10)
        elif t.startswith("In parallel, DNA methylation is attractive") and "Fig. 2a" in (p.text or ""):
            _set_paragraph_plain(p, intro_p11)
        elif t.startswith("The primary finding is straightforward") and "Fig. 1a" in (p.text or ""):
            _set_paragraph_plain(p, disc_p69)
        elif t.startswith("A second contribution is the clinically anchored") and "Fig. 2a" in (p.text or ""):
            _set_paragraph_plain(p, disc_p70)
        elif t.startswith("Our use of a 60-month OS window") and "Fig. 4d" in (p.text or ""):
            _set_paragraph_plain(p, disc_p71)
        elif t.startswith("Limitations include reliance on a cBioPortal") and "Figures 1–6" in (p.text or ""):
            _set_paragraph_plain(p, disc_p72)
        elif t.startswith("We used the cBioPortal-format") and "Histology was determined" in (p.text or ""):
            rest = (p.text or "").split("Histology was determined", 1)[1]
            _set_paragraph_plain(p, methods_p77_prefix + "Histology was determined" + rest)


def _set_paragraph_plain(p: Paragraph, text: str) -> None:
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    p.add_run(text)


def _filter_results_block(doc: Document, els: list) -> list:
    """Drop empty Heading 2 paragraphs inside Results."""
    out: list = []
    for el in els:
        if el.tag != qn("w:p"):
            out.append(el)
            continue
        para = Paragraph(el, doc._body)
        if not (para.text or "").strip() and para.style.name.startswith("Heading 2"):
            continue
        out.append(el)
    return out


def _filter_discussion_block(doc: Document, els: list) -> list:
    """Remove redundant summary paragraphs (old 73–74) and set Discussion title to Heading 1."""
    skip_prefixes = (
        "In TCGA esophageal carcinoma, prognostic models are strongly histology-specific and do not transfer between",
        "Fast-death anchored methylation analyses reveal distinct histology-resolved epigenetic programs",
    )
    out: list = []
    for el in els:
        if el.tag != qn("w:p"):
            out.append(el)
            continue
        txt = _para_text(el).strip()
        if any(txt.startswith(s) for s in skip_prefixes):
            continue
        para = Paragraph(el, doc._body)
        if txt == "Discussion":
            para.style = doc.styles["Heading 1"]
        out.append(el)
    return out


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    path = root / DOC_NAME
    doc = Document(str(path))

    # Paragraph indices verified against pre-reorder manuscript (Introduction → Results → Discussion → Methods → Refs).
    intro = _elements_between_paragraphs(doc, 0, 12)
    # Include trailing blank paragraphs before Discussion (former indices 66–67).
    results = _filter_results_block(doc, _elements_between_paragraphs(doc, 13, 67))
    discussion_raw = _elements_between_paragraphs(doc, 68, 74)
    methods = _elements_between_paragraphs(doc, 75, 87)
    refs_onward = _elements_from_paragraph_to_end(doc, 88)

    _edit_introduction_and_discussion(doc)
    discussion = _filter_discussion_block(doc, discussion_raw)

    # Detach selected elements from body (preserve sectPr handled in rebuild).
    seen = set()
    for block in (intro, methods, results, discussion, refs_onward):
        for el in block:
            seen.add(id(el))

    body = doc.element.body
    for el in list(body):
        if el.tag == qn("w:sectPr"):
            continue
        if id(el) in seen:
            el.getparent().remove(el)

    _rebuild_body(doc, [intro, methods, results, discussion, refs_onward])
    _strip_maintext_refs_from_supplementary(doc)

    doc.save(str(path))
    print("Reorganized manuscript:", path)


if __name__ == "__main__":
    main()
