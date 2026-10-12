"""
Re-embed main-text figures in Manuscript First.docx (image paragraph immediately above caption).

Layout matches this manuscript: [PNG] then [Fig. N. caption ...]
Only inline Results/Introduction captions (Fig. N.) — not the trailing Figure legends block.
"""
from __future__ import annotations

from esca_escc_multiomic_pipeline.study_paths import study_root

import argparse
import re
import shutil
import zipfile
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph

STUDY_ROOT = study_root()
DEFAULT_DOCX = STUDY_ROOT / "Manuscript First.docx"
PLOT_DIR = STUDY_ROOT / "Plots"

REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
IMAGE_REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"

# prefix, png filename, width in inches
FIGURE_SPECS: list[tuple[str, str, float]] = [
    ("Fig. 1.", "00_analysis_pipeline_flowchart.png", 6.4),
    ("Fig. 2.", "01_cross_apply_cindex.png", 6.2),
    ("Fig. 3.", "02_fastdeath_methylation_volcano_panel.png", 6.6),
    ("Fig. 4.", "04_KM_methylation_60mo_panel.png", 6.6),
    ("Fig. 5.", "05_somatic_combined_A_heatmap_B_interaction.png", 6.2),
]


def _insert_paragraph_before(paragraph: Paragraph, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    return new_para


def _p_has_image(p: Paragraph) -> bool:
    return bool(p._p.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"))


def _ensure_image_before_caption(caption_para: Paragraph, png: Path, width: Inches) -> None:
    cap_el = caption_para._p
    prev = cap_el.getprevious()
    if prev is not None and prev.tag.endswith("}p"):
        img_para = Paragraph(prev, caption_para._parent)
        if _p_has_image(img_para):
            for r in list(img_para.runs):
                r._r.getparent().remove(r._r)
            img_para.add_run().add_picture(str(png), width=width)
            return
        if not (img_para.text or "").strip():
            for r in list(img_para.runs):
                r._r.getparent().remove(r._r)
            img_para.add_run().add_picture(str(png), width=width)
            return
    pic_p = _insert_paragraph_before(caption_para, "")
    pic_p.add_run().add_picture(str(png), width=width)


def _find_inline_caption(doc: Document, prefix: str) -> Paragraph:
    in_legends = False
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t == "Figure legends":
            in_legends = True
            continue
        if in_legends:
            continue
        if t.startswith(prefix):
            return p
    raise ValueError(f"Inline caption not found: {prefix!r}")


def _prune_orphan_image_parts(docx_path: Path) -> None:
    with zipfile.ZipFile(docx_path, "r") as zin:
        doc_xml = zin.read("word/document.xml").decode("utf-8")
        used_rids = set(re.findall(r'r:embed="(rId\d+)"', doc_xml))
        rels_data = zin.read("word/_rels/document.xml.rels")

    root = ET.fromstring(rels_data)
    remove_rels = []
    keep_media: set[str] = set()
    for rel in root.findall(f"{{{REL_NS}}}Relationship"):
        if rel.get("Type") != IMAGE_REL_TYPE:
            continue
        rid = rel.get("Id")
        tgt = rel.get("Target", "").replace("\\", "/")
        if rid in used_rids:
            keep_media.add(tgt)
        else:
            remove_rels.append(rel)
    for rel in remove_rels:
        root.remove(rel)
    new_rels = ET.tostring(root, encoding="utf-8", xml_declaration=True)

    buf = BytesIO()
    with zipfile.ZipFile(docx_path, "r") as zin, zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            name = item.filename
            if name == "word/_rels/document.xml.rels":
                zout.writestr(item, new_rels)
                continue
            if name.startswith("word/media/") and name.endswith(".png"):
                rel_path = name.removeprefix("word/")
                if rel_path not in keep_media:
                    continue
            zout.writestr(item, zin.read(name))
    docx_path.write_bytes(buf.getvalue())


def reembed(docx_path: Path) -> None:
    for _, stem, _ in FIGURE_SPECS:
        png = PLOT_DIR / stem
        if not png.is_file():
            raise FileNotFoundError(f"Missing plot: {png}")

    backup = docx_path.with_suffix(docx_path.suffix + ".pre_reembed_figs.bak")
    shutil.copy2(docx_path, backup)

    doc = Document(str(docx_path))
    for prefix, stem, w in FIGURE_SPECS:
        cap = _find_inline_caption(doc, prefix)
        _ensure_image_before_caption(cap, PLOT_DIR / stem, Inches(w))

    doc.save(str(docx_path))
    _prune_orphan_image_parts(docx_path)

    doc2 = Document(str(docx_path))
    n_img = sum(1 for r in doc2.part.rels.values() if "image" in r.reltype)
    size_mb = docx_path.stat().st_size / 1024 / 1024
    print(f"Backup: {backup}")
    print(f"Re-embedded {len(FIGURE_SPECS)} figures in: {docx_path}")
    print(f"Embedded images: {n_img}; file size: {size_mb:.2f} MB")


def main() -> None:
    ap = argparse.ArgumentParser(description="Re-embed main figures in Manuscript First.docx")
    ap.add_argument("--docx", type=Path, default=DEFAULT_DOCX)
    args = ap.parse_args()
    reembed(args.docx.resolve())


if __name__ == "__main__":
    main()
