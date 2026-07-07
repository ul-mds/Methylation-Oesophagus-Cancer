"""
Replace left/right figure language with panel letters and standardise in-text citations
as Fig. 1a, Fig. 2a,b, Fig. 4a–f, etc. Skips long Figure-legend paragraphs (>350 chars).
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document

DOC = "TCGA_ESCA_ESCC_landmark60_full_scientific_manuscript.docx"


def _transform(text: str) -> str:
    if len(text) > 350 and text.strip().startswith("Figure ") and "Genome-wide" in text:
        return text
    if len(text) > 350 and text.strip().startswith("Figure ") and "Histology-specific versus cross-lineage" in text:
        return text
    if len(text) > 350 and text.strip().startswith("Figure ") and "Kaplan–Meier estimates" in text:
        return text
    if len(text) > 350 and text.strip().startswith("Figure ") and "Exploratory heatmap" in text:
        return text
    if len(text) > 350 and text.strip().startswith("Figure ") and "Screen of statistical interaction" in text:
        return text
    if len(text) > 350 and text.strip().startswith("Figure ") and "Ranked display" in text:
        return text

    pairs = [
        ("Figure 2, left", "Fig. 2a"),
        ("Figure 2, right", "Fig. 2b"),
        ("Figure 4, upper row", "Fig. 4a–c"),
        ("lower row of Figure 4", "Fig. 4d–f"),
        ("Figure 4 (lower row)", "Fig. 4d–f"),
        ("Figures 2–3", "Fig. 2a,b and Fig. 3a,b"),
        ("Figures 5–6", "Fig. 5a and Fig. 6a"),
        ("Figure 5 summarizes", "Fig. 5a summarizes"),
        ("Figure 6 then presents", "Fig. 6a then presents"),
        ("Figure 5 and Table 7", "Fig. 5a and Table 7"),
        ("Figure 6 and Table 7", "Fig. 6a and Table 7"),
        ("positioned Figure 6 and Table 7", "positioned Fig. 6a and Table 7"),
        ("via Figure 5,", "via Fig. 5a,"),
        ("Read sequentially, Figure 5 and Table 7", "Read sequentially, Fig. 5a and Table 7"),
        ("while Figure 6 prioritizes", "while Fig. 6a prioritizes"),
        ("while Figure 6 prioritises", "while Fig. 6a prioritises"),
        ("Figure 1 visualizes", "Fig. 1a visualizes"),
        ("Figure 1 and Table 1", "Fig. 1a and Table 1"),
        ("Figure 1;", "Fig. 1a;"),
        ("(Figure 1;", "(Fig. 1a;"),
        ("Figure 3 compresses", "Fig. 3a,b compresses"),
        ("in Figure 3.", "in Fig. 3a,b."),
        ("Figure 4 then", "Fig. 4a–f then"),
        ("(Figure 4;", "(Fig. 4a–f;"),
        (" (Figure 4)", " (Fig. 4a–f)"),
        ("(Figure 4; Table 6)", "(Fig. 4a–f; Table 6)"),
        ("(Figures 5–6; Table 7)", "(Fig. 5a and Fig. 6a; Table 7)"),
        ("main-text Figure 6)", "main-text Fig. 6a)"),
        ("main-text Figure 1.)", "main-text Fig. 1a.)"),
        ("main-text Figure 2.)", "main-text Fig. 2a,b.)"),
        ("main-text Figure 3.)", "main-text Fig. 3a,b.)"),
        ("main-text Figure 4.)", "main-text Fig. 4a–f.)"),
        ("main-text Figure 5.)", "main-text Fig. 5a.)"),
        ("main-text Figure 6.)", "main-text Fig. 6a.)"),
    ]
    for old, new in pairs:
        text = text.replace(old, new)

    text = text.replace("addition1al", "additional")
    text = text.replace(
        "so the Fig. 4d–f extends",
        "so Fig. 4d–f extend",
    )
    text = text.replace(
        "Together, Fig. 2a,b and Fig. 3a,b and Tables",
        "Together, Fig. 2a,b, Fig. 3a,b, and Tables",
    )
    text = text.replace(
        "Fig. 3a,b compresses these landscapes",
        "Panels a and b of Fig. 3 compress these landscapes",
    )

    supp_fixes = [
        (
            "Left: ESCA/EAC; right: ESCC.",
            "Panel a, ESCA/EAC; panel b, ESCC.",
        ),
        (
            "in ESCA/EAC (left) and ESCC (right). (Main-text Figure 3.)",
            "in ESCA/EAC (panel a) and ESCC (panel b). (Main-text Fig. 3a,b.)",
        ),
        ("(Main-text Figure 1.)", "(Main-text Fig. 1a.)"),
        ("(Main-text Figure 2.)", "(Main-text Fig. 2a,b.)"),
        ("(Main-text Figure 4.)", "(Main-text Fig. 4a–f.)"),
        ("(Main-text Figure 5.)", "(Main-text Fig. 5a.)"),
        ("(Main-text Figure 6.)", "(Main-text Fig. 6a.)"),
    ]
    for old, new in supp_fixes:
        text = text.replace(old, new)

    return text


def _short_captions(text: str) -> str:
    m = {
        "Fig. 1. Cross-apply performance (C-index) using overall survival (OS).": (
            "Fig. 1a. Cross-apply performance (C-index) using overall survival (OS) (single summary panel)."
        ),
        "Fig. 2. Fast-death differential methylation by histology (volcano).": (
            "Fig. 2a,b. Fast-death differential methylation by histology (volcano; panel a, ESCA/EAC; panel b, ESCC)."
        ),
        "Fig. 3. Top fast-death hypo- and hypermethylated loci (Δβ) in ESCA/EAC (left) and ESCC (right).": (
            "Fig. 3a,b. Top fast-death hypo- and hypermethylated loci (Δβ; panel a, ESCA/EAC; panel b, ESCC)."
        ),
        "Fig. 4. Kaplan–Meier curves for overall survival stratified by median methylation.": (
            "Fig. 4a–f. Kaplan–Meier curves for overall survival stratified by median methylation "
            "(panels a–c, ESCA/EAC; panels d–f, ESCC)."
        ),
        "Fig. 5. Exploratory mechanistic follow-up: mutation-associated shifts in expression-derived module scores (Δ=mut−WT) stratified by histology.": (
            "Fig. 5a. Exploratory mechanistic follow-up: mutation-associated shifts in expression-derived module scores "
            "(Δ=mut−WT) stratified by histology (single heatmap panel)."
        ),
        "Fig. 6. Mutation×histology interaction screen on overall survival (nominal p-values).": (
            "Fig. 6a. Mutation×histology interaction screen on overall survival (nominal p-values; single scatter panel)."
        ),
    }
    return m.get(text.strip(), text)


def main() -> None:
    path = Path(__file__).resolve().parents[1] / DOC
    doc = Document(str(path))
    for p in doc.paragraphs:
        t = p.text
        if not t:
            continue
        nt = _short_captions(t)
        nt = _transform(nt)
        if nt != t:
            p.text = nt

    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    t = p.text
                    if not t:
                        continue
                    nt = _transform(t)
                    if nt != t:
                        p.text = nt

    doc.save(str(path))
    print("Figure panel citations normalised:", path)


if __name__ == "__main__":
    main()
