"""Resolve final scientific and language issues found after package re-audit."""
from __future__ import annotations

import math
import re
import shutil
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lifelines import CoxPHFitter
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .survival_60mo import truncate_os_at_months

ROOT = Path("/Users/senol/Desktop/AML/Esophageal Carcinoma (TCGA, Firehose Legacy)")
PKG = ROOT / "Cancer Medicine submission_package"
MANUSCRIPT = PKG / "02_Manuscript.docx"
RESULT = ROOT / "result"
SUP = PKG / "Supplementary"
YELLOW = WD_COLOR_INDEX.YELLOW
YELLOW_FILL = PatternFill("solid", fgColor="FFF59D")


def clear_runs(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag in {qn("w:r"), qn("w:hyperlink")}:
            paragraph._p.remove(child)


def set_highlighted(paragraph: Paragraph, text: str) -> None:
    clear_runs(paragraph)
    run = paragraph.add_run(text)
    run.font.highlight_color = YELLOW


def find_starts(doc: Document, prefix: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(prefix)


def bh_fdr(pvalues: pd.Series) -> pd.Series:
    values = pvalues.to_numpy(dtype=float)
    out = np.full(values.shape, np.nan)
    valid = np.isfinite(values)
    p = values[valid]
    if not len(p):
        return pd.Series(out, index=pvalues.index)
    order = np.argsort(p)
    ranked = p[order]
    q = ranked * len(ranked) / np.arange(1, len(ranked) + 1)
    q = np.minimum.accumulate(q[::-1])[::-1]
    q = np.clip(q, 0, 1)
    restored = np.empty_like(q)
    restored[order] = q
    out[valid] = restored
    return pd.Series(out, index=pvalues.index)


def recompute_eac_os60_cox() -> pd.DataFrame:
    path = RESULT / "methylation_fast_death_EAC_ESCA.tsv"
    result = pd.read_csv(path, sep="\t")
    selected = result.index[result["cox_p_meth"].notna()].tolist()

    master = pd.read_csv(RESULT / "master_clinical_samples.tsv", sep="\t", low_memory=False)
    master = master[master["HISTOLOGY"] == "EAC/ESCA"].copy()
    master = master.dropna(subset=["OS_MONTHS", "OS_EVENT"])
    master["SAMPLE_ID"] = master["SAMPLE_ID"].astype(str)
    master["OS_EVENT"] = master["OS_EVENT"].astype(int)
    t60, e60 = truncate_os_at_months(master["OS_MONTHS"], master["OS_EVENT"], months=60.0)
    clinical = pd.DataFrame(
        {
            "sample": master["SAMPLE_ID"].to_numpy(),
            "time": t60.to_numpy(dtype=float),
            "event": e60.to_numpy(dtype=int),
        }
    ).set_index("sample")

    meth = pd.read_csv(ROOT / "data_methylation_hm450.txt", sep="\t", low_memory=False)
    probe_col = meth.columns[0]
    meth = meth.set_index(probe_col)
    meth.index = meth.index.astype(str)
    meth.columns = meth.columns.astype(str)
    shared = [sample for sample in clinical.index if sample in meth.columns]
    clinical = clinical.loc[shared]

    betas: list[float] = []
    pvals: list[float] = []
    ns: list[int] = []
    events: list[int] = []
    for number, idx in enumerate(selected, 1):
        probe = str(result.at[idx, "probe"])
        if probe not in meth.index:
            betas.append(np.nan)
            pvals.append(np.nan)
            ns.append(0)
            events.append(0)
            continue
        beta = pd.to_numeric(meth.loc[probe, shared], errors="coerce")
        frame = clinical.join(beta.rename("beta"), how="inner").dropna()
        if len(frame) < 20 or frame["event"].sum() < 5 or frame["beta"].nunique() < 3:
            betas.append(np.nan)
            pvals.append(np.nan)
            ns.append(len(frame))
            events.append(int(frame["event"].sum()))
            continue
        try:
            cph = CoxPHFitter()
            cph.fit(frame[["time", "event", "beta"]], duration_col="time", event_col="event")
            betas.append(float(cph.params_["beta"]))
            pvals.append(float(cph.summary.loc["beta", "p"]))
        except Exception:
            betas.append(np.nan)
            pvals.append(np.nan)
        ns.append(len(frame))
        events.append(int(frame["event"].sum()))
        if number % 250 == 0:
            print(f"OS60 Cox {number}/{len(selected)}")

    result.loc[selected, "cox_beta_meth"] = betas
    result.loc[selected, "cox_p_meth"] = pvals
    result.loc[selected, "cox_q_meth"] = bh_fdr(pd.Series(pvals)).to_numpy()
    result.loc[selected, "cox_n"] = ns
    result.loc[selected, "cox_events"] = events
    result.to_csv(path, sep="\t", index=False)
    return result


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    clear_runs(paragraph)
    run = paragraph.add_run(text)
    run.font.highlight_color = YELLOW


def resize_table(table, n_body_rows: int) -> None:
    target = n_body_rows + 1
    while len(table.rows) > target:
        table._tbl.remove(table.rows[-1]._tr)
    while len(table.rows) < target:
        table.add_row()


def revise_word_tables_and_language(eac: pd.DataFrame) -> None:
    doc = Document(str(MANUSCRIPT))
    escc = pd.read_csv(RESULT / "methylation_fast_death_ESCC.tsv", sep="\t")

    # Safer title and complete terminology.
    set_highlighted(
        doc.paragraphs[0],
        "Histology-dependent prognostic model transportability and candidate methylation patterns "
        "in esophageal carcinoma",
    )
    set_highlighted(find_starts(doc, "Acknowledgements"), "Acknowledgments")
    set_highlighted(
        find_starts(doc, "This work was supported by"),
        "This work was supported by the Saxon Ministry of Science, Research and the Arts through the "
        "Tag-White project (grant numbers 100602109 [Mittweida] and 100601615 [Leipzig]) and by the "
        "Leipzig University Open Access Publishing Fund.",
    )
    set_highlighted(
        find_starts(doc, "The authors thank the TCGA Research Network"),
        "The authors thank the TCGA Research Network and the cBioPortal team for providing curated "
        "public data. Figures were generated with Python using matplotlib and seaborn. Analysis scripts "
        "and package versions are available in the repository described in the Data availability statement.",
    )

    abstract = find_starts(doc, "Background: Esophageal adenocarcinoma")
    set_highlighted(
        abstract,
        abstract.text.replace(
            "Using TCGA ESCA Firehose Legacy data",
            "Using The Cancer Genome Atlas Esophageal Carcinoma (TCGA-ESCA) Firehose Legacy data",
        ).replace(
            "External GEO analyses",
            "External Gene Expression Omnibus (GEO) analyses",
        ).replace(
            "different short-survival-associated patterns",
            "different candidate short-survival-associated patterns",
        ),
    )

    set_highlighted(
        find_starts(doc, "Esophageal carcinoma is characterized by"),
        "Esophageal carcinoma comprises two predominant histological subtypes: esophageal adenocarcinoma "
        "(EAC) and esophageal squamous cell carcinoma (ESCC), which have divergent epidemiological and "
        "molecular profiles [1]. Major risk factors for ESCC include alcohol use and tobacco smoking, "
        "whereas EAC is associated with Barrett’s esophagus, a precursor condition arising in the setting "
        "of gastroesophageal reflux disease [2]. ESCC more commonly occurs in the mid-esophagus, whereas "
        "EAC usually arises in the distal esophagus or at the gastroesophageal junction. The Cancer Genome "
        "Atlas established that EAC and ESCC are molecularly distinct [3]. Subsequent comparative analyses "
        "[4], EAC molecular studies [5,16], and ESCC subtype analyses [6,7] further demonstrated "
        "histology-associated genomic, transcriptional, and epigenetic features. The CROSS trial included "
        "both histologies within a shared neoadjuvant chemoradiotherapy framework [8]. Nevertheless, "
        "secondary prognostic analyses frequently pool EAC and ESCC or evaluate only one histology without "
        "testing cross-histology transportability.",
    )
    set_highlighted(
        find_starts(doc, "Throughout this manuscript, prognostic denotes"),
        "Throughout this manuscript, prognostic denotes association with overall survival truncated at "
        "60 months (OS60), independent of a specified therapy [9]. Predictive is reserved for treatment-"
        "response settings, such as pathological regression after chemoradiotherapy [10]. Because the "
        "cellular origins of Barrett’s metaplasia and EAC remain incompletely resolved [11–14], we "
        "interpret histology-associated epigenetic patterns as lineage-linked observations rather than "
        "evidence of a single developmental origin.",
    )
    set_highlighted(
        find_starts(doc, "A prognostic score intended for esophageal carcinoma"),
        find_starts(doc, "A prognostic score intended for esophageal carcinoma").text.replace(
            "[15–20]", "[15,17–20]"
        ),
    )
    set_highlighted(
        find_starts(doc, "Finally, we added exploratory"),
        "Finally, we added exploratory mutation–histology screens and expression-module contrasts to "
        "provide somatic context for the methylation findings [25,26].",
    )

    # Clarify the two analysis populations.
    cohort_p = find_starts(doc, "The analysis sequence is shown in Fig. 1.")
    set_highlighted(
        cohort_p,
        cohort_p.text.replace(
            "The Firehose Legacy dataset was used as the fixed discovery cohort.",
            "The complete paired multi-omic modeling cohort included 89 EAC and 95 ESCC samples. "
            "The methylation-only short-survival analysis included one additional ESCC sample (n=96) "
            "that lacked the complete paired feature set required for prognostic modeling. The Firehose "
            "Legacy dataset was used as the fixed discovery cohort.",
        ),
    )
    set_highlighted(
        find_starts(doc, "We examined the shortest-survival tail"),
        find_starts(doc, "We examined the shortest-survival tail").text.replace(
            "fast-death group", "short-survival group"
        ).replace(
            "fast-death samples", "short-survival samples"
        ).replace(
            "β_fast-death", "β_short-survival"
        ),
    )
    set_highlighted(
        find_starts(doc, "The GDC reprocessing of overlapping"),
        find_starts(doc, "The GDC reprocessing of overlapping").text.replace(
            "This improvement indicates that histology-related baseline hazard differences must be modeled "
            "explicitly; it does not demonstrate",
            "This result suggests that accounting for histology-related baseline-hazard differences improved "
            "apparent pooled-model discrimination; it does not demonstrate",
        ),
    )
    set_highlighted(
        find_starts(doc, "We next examined whether methylation correlates"),
        find_starts(doc, "We next examined whether methylation correlates").text.replace(
            "We next examined whether methylation correlates of short survival also differed between histologies",
            "We next characterized methylation correlates of short survival separately within each histology",
        ),
    )
    set_highlighted(
        find_starts(doc, "As exploratory follow-up"),
        find_starts(doc, "As exploratory follow-up").text.replace(
            "These are parallel observations to the transport failure",
            "These observations complement the transportability analysis",
        ),
    )

    # Main Table 2: 10 ESCC hypo + 10 ESCC hyper, no EAC rows.
    hypo = escc[escc["delta_fast_minus_other"] < 0].sort_values("q_diff").head(10).copy()
    hyper = escc[escc["delta_fast_minus_other"] > 0].sort_values("q_diff").head(10).copy()
    t2 = pd.concat([hypo.assign(Direction="Hypomethylated"), hyper.assign(Direction="Hypermethylated")])
    table2 = doc.tables[1]
    resize_table(table2, len(t2))
    for row, (_, item) in zip(table2.rows[1:], t2.iterrows()):
        values = [
            str(item["probe"]),
            str(item["Direction"]),
            f"{item['mean_beta_fast']:.3f}",
            f"{item['mean_beta_other']:.3f}",
            f"{item['delta_fast_minus_other']:.3f}",
            f"{item['q_diff']:.3e}" if item["q_diff"] < 0.001 else f"{item['q_diff']:.3f}",
        ]
        for cell, value in zip(row.cells, values):
            set_cell_text(cell, value)

    # Main Table 3 rebuilt from the corrected OS60 Cox screen.
    top = eac.dropna(subset=["cox_q_meth"]).sort_values("cox_q_meth").head(5)
    table3 = doc.tables[2]
    resize_table(table3, len(top))
    for row, (_, item) in zip(table3.rows[1:], top.iterrows()):
        values = [
            str(item["probe"]),
            f"{item['cox_beta_meth']:.3f}",
            f"{item['cox_q_meth']:.3f}",
            str(int(item["cox_n"])),
            str(int(item["cox_events"])),
        ]
        for cell, value in zip(row.cells, values):
            set_cell_text(cell, value)

    table3_note = find_starts(doc, "Univariate Cox models of continuous HM450")
    set_highlighted(
        table3_note,
        "Univariate Cox models of continuous HM450 β versus OS60 were fitted separately for 2,000 "
        "prespecified EAC probes. Cox β is the log-hazard increment per unit increase in β, and q is "
        "the Benjamini–Hochberg FDR across the screened probes. Each fit included 89 patients and 44 "
        "OS60 events. This exploratory screen does not define a multivariable signature.",
    )
    set_highlighted(
        find_starts(doc, "†Cox β"),
        "†Cox β, log-hazard per unit increase in methylation beta. ‡q, Benjamini–Hochberg FDR. "
        "§n, sample size. ¶Events, number of OS60 events.",
    )

    # Improve table notes and final terminology.
    for prefix in [
        "The histology-stratified fast-death analysis",
        "We repeated patient classification using",
        "The genome-wide contrasts showed different fast-death",
        "Table 1 lists the leading EAC/ESCA",
        "Table 2 lists the leading ESCC",
        "Fast-death denotes the shortest OS60 among EAC/ESCA",
        "Fast-death denotes the shortest OS60 among ESCC",
    ]:
        try:
            paragraph = find_starts(doc, prefix)
        except ValueError:
            continue
        text = paragraph.text
        text = text.replace("fast-death", "short-survival")
        text = text.replace("EAC/ESCA", "EAC subset of TCGA-ESCA")
        set_highlighted(paragraph, text)

    heading = find_starts(doc, 'The "Fast-Death" Phenotype')
    set_highlighted(heading, "Identification of Histology-Specific Short-Survival Methylation Patterns")
    for paragraph in doc.paragraphs:
        text = paragraph.text
        if "fast-death" in text:
            set_highlighted(paragraph, text.replace("fast-death", "short-survival"))

    # Correct Figure 5 multiplicity wording and S8 punctuation.
    fig5 = find_starts(doc, "Figure 5. Exploratory somatic context")
    set_highlighted(
        fig5,
        fig5.text.replace(
            "genome-wide multiplicity correction",
            "multiplicity correction across the 80 screened genes",
        ),
    )
    s8 = find_starts(doc, "Supplementary Table S8.")
    set_highlighted(s8, s8.text.replace("genes,110", "genes; 110").replace("top 1200", "top 1,200"))

    # Conclusions: different candidate patterns, not direct proof of distinctness.
    conclusion = find_starts(doc, "Prognostic models intended for use")
    set_highlighted(
        conclusion,
        conclusion.text.replace(
            "identified distinct short-survival-associated patterns",
            "identified different candidate short-survival-associated patterns",
        ),
    )

    # Standardize reference formatting for 11 and 16.
    set_highlighted(
        find_starts(doc, "11. Hassall"),
        "11. Hassall, E. Barrett’s esophagus: congenital or acquired? Am. J. Gastroenterol. "
        "1993;88:819–824. PMID: 8503373.",
    )
    ref16 = find_starts(doc, "16. Ng")
    set_highlighted(
        ref16,
        "16. Ng, A. W. T. et al. Disentangling oncogenic amplicons in esophageal adenocarcinoma. "
        "Nat. Commun. 2024;15:4074. https://doi.org/10.1038/s41467-024-47619-4.",
    )

    doc.save(str(MANUSCRIPT))
    clean = Document(str(MANUSCRIPT))
    for element in clean._element.findall(".//" + qn("w:highlight")):
        element.getparent().remove(element)
    clean.save(str(PKG / "02_Manuscript_clean.docx"))

    write_main_table_workbooks(t2, top)


def write_main_table_workbooks(t2: pd.DataFrame, top: pd.DataFrame) -> None:
    tables = PKG / "Tables"

    wb = Workbook()
    ws = wb.active
    ws.title = "Table_2"
    ws.append(["Table 2. ESCC: leading short-survival-associated hypo- and hypermethylated HM450 loci."])
    ws["A1"].fill = YELLOW_FILL
    ws["A1"].font = Font(bold=True)
    ws.append([])
    ws.append(["probe", "direction", "mean_beta_short_survival", "mean_beta_other", "delta", "q_diff"])
    for cell in ws[3]:
        cell.fill = YELLOW_FILL
        cell.font = Font(bold=True)
    for _, item in t2.iterrows():
        ws.append(
            [
                item["probe"], item["Direction"], item["mean_beta_fast"], item["mean_beta_other"],
                item["delta_fast_minus_other"], item["q_diff"],
            ]
        )
    ws.append([])
    ws.append(["REVISED FOR REVIEW: all rows are from the ESCC source table."])
    ws.cell(ws.max_row, 1).fill = YELLOW_FILL
    wb.save(tables / "Table_2.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Table_3"
    ws.append(["Table 3. EAC: univariate Cox associations between continuous probe methylation and OS60."])
    ws["A1"].fill = YELLOW_FILL
    ws["A1"].font = Font(bold=True)
    ws.append([])
    ws.append(["probe", "cox_beta_meth", "cox_q_meth", "cox_n", "cox_events"])
    for cell in ws[3]:
        cell.fill = YELLOW_FILL
        cell.font = Font(bold=True)
    for _, item in top.iterrows():
        ws.append(
            [
                item["probe"], item["cox_beta_meth"], item["cox_q_meth"],
                int(item["cox_n"]), int(item["cox_events"]),
            ]
        )
    ws.append([])
    ws.append(["REVISED FOR REVIEW: Cox models use OS60; 44 events."])
    ws.cell(ws.max_row, 1).fill = YELLOW_FILL
    wb.save(tables / "Table_3.xlsx")


def clean_supplementary_metadata() -> None:
    xlsx = SUP / "02_Supplementary_Tables.xlsx"
    wb = load_workbook(xlsx)
    for name in list(wb.sheetnames):
        if name.startswith("File_info") and name != "File_info":
            del wb[name]
    if "File_info" in wb.sheetnames:
        info = wb["File_info"]
        info["A4"] = (
            "Corrected machine-readable Tables S1–S12. S2–S5 are direction-specific methylation "
            "tables; S6 is the EAC OS60 Cox screen."
        )
        info["A4"].alignment = Alignment(wrap_text=True)
        for row in info.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "fast-death" in cell.value.lower():
                    cell.value = re.sub(
                        "fast-death", "short-survival", cell.value, flags=re.IGNORECASE
                    )
                    cell.fill = YELLOW_FILL
    if "S12" in wb.sheetnames:
        ws = wb["S12"]
        for row in ws.iter_rows(min_row=2):
            item = str(row[0].value)
            if item == "13":
                row[3].value = "Methods; analysis-repository model coefficient files"
                row[3].fill = YELLOW_FILL
            if item == "20":
                row[3].value = "Analysis-repository coefficient files; no nomogram"
                row[3].fill = YELLOW_FILL
    wb.save(xlsx)
    shutil.copy2(xlsx, SUP / "Supporting_Information_tables.xlsx")
    review_dir = SUP / "Review_highlighted"
    review_dir.mkdir(exist_ok=True)
    shutil.copy2(xlsx, review_dir / "02_Supplementary_Tables_review.xlsx")


def patch_endpoint_tsv() -> None:
    path = PKG / "Supplementary" / "Tables" / "S01_model_performance_cindex.tsv"
    frame = pd.read_csv(path, sep="\t")
    frame["endpoint"] = "OS60_truncated"
    frame.to_csv(path, sep="\t", index=False)


def sync_regenerated_figures() -> None:
    plots = ROOT / "Plots"
    figures = PKG / "Figures"
    copies = {
        plots / "00_analysis_pipeline_flowchart.png": figures / "Figure_1.png",
        plots / "00_analysis_pipeline_flowchart.pdf": figures / "Figure_1.pdf",
        plots / "03_fastdeath_methylation_combined_A_volcano_B_ranked.png": figures / "Figure_3.png",
        plots / "04_KM_methylation_60mo_panel.png": figures / "Figure_4.png",
    }
    for source, destination in copies.items():
        if not source.is_file():
            raise FileNotFoundError(source)
        shutil.copy2(source, destination)


def normalize_remaining_manuscript_terminology() -> None:
    doc = Document(str(MANUSCRIPT))
    for paragraph in doc.paragraphs:
        text = paragraph.text
        revised = re.sub(r"fast-death", "short-survival", text, flags=re.IGNORECASE)
        revised = revised.replace(
            "Short-survival death-tail sensitivity (10%, 20%, 30%) on the OS60 cohort. "
            "Short-survival = shortest-survival tail among OS60 decedents (minimum n=10).",
            "Short-survival threshold sensitivity (10%, 20%, and 30%) in the OS60 cohort. "
            "The short-survival group comprised the shortest-survival tail among OS60 decedents "
            "(minimum n = 10).",
        )
        if "Full data: Supplementary/Tables/S10_" in revised:
            revised = revised.split("Full data:", 1)[0].rstrip() + (
                " Complete results are provided in Supporting Information file 2 (Table S10)."
            )
        if revised != text:
            set_highlighted(paragraph, revised)

    targeted_revisions = {
        "short-survival methylation programs by histology": (
            "Short-Survival Methylation Patterns by Histology"
        ),
        "Supplementary Table S6.": (
            "Supplementary Table S6. EAC: complete univariate Cox screen of continuous HM450 probe "
            "methylation versus OS60 (2,000 prespecified probes; Benjamini–Hochberg FDR)."
        ),
        "Supplementary Table S8.": (
            "Supplementary Table S8. GSE161533 (GEO GPL570): overlap between a prespecified union of "
            "110 TCGA ESCC genes and gene-level summaries of the top 1,200 probes ranked by the absolute "
            "mean paired tumor–normal log2 fold change (n = 28 patients)."
        ),
        "Supplementary Table S10.": (
            "Supplementary Table S10. Sensitivity of the short-survival definition at 10%, 20%, and "
            "30% in the OS60 cohort. The short-survival group comprised the shortest-survival tail among "
            "OS60 decedents (minimum n = 10). Methylation contrasts compared the short-survival group "
            "with all other samples within each histology using Welch tests and Benjamini–Hochberg FDR. "
            "The table reports patient-membership and top-200-probe Jaccard indices relative to the 20% "
            "primary definition."
        ),
        "Supplementary Table S11.": (
            "Supplementary Table S11. Exploratory mutation-by-histology interaction screen on OS60 "
            "across the 80 most frequently mutated genes (nominal Wald P values; hypothesis-generating)."
        ),
    }
    for prefix, replacement in targeted_revisions.items():
        set_highlighted(find_starts(doc, prefix), replacement)

    for prefix in [
        "Table 1 lists the leading EAC subset",
        "short-survival denotes the shortest OS60 among EAC",
        "short-survival denotes the shortest OS60 among ESCC",
        "Figure 3. short-survival",
    ]:
        paragraph = find_starts(doc, prefix)
        text = paragraph.text
        if prefix.startswith("Table 1"):
            text = text.replace(". short-survival was", ". Short-survival was")
        elif prefix.startswith("Figure 3"):
            text = text.replace("Figure 3. short-survival", "Figure 3. Short-survival", 1)
        else:
            text = "Short-survival" + text[len("short-survival"):]
        set_highlighted(paragraph, text)

    for prefix in [
        "Supplementary Table S2. EAC/ESCA:",
        "Supplementary Table S4. EAC/ESCA:",
    ]:
        paragraph = find_starts(doc, prefix)
        set_highlighted(paragraph, paragraph.text.replace("EAC/ESCA:", "EAC:"))

    for table in doc.tables[:2]:
        for cell in table.rows[0].cells:
            if re.search(r"fast-death", cell.text, flags=re.IGNORECASE):
                set_cell_text(
                    cell,
                    re.sub(r"fast-death", "short-survival", cell.text, flags=re.IGNORECASE),
                )

    doc.save(str(MANUSCRIPT))
    clean = Document(str(MANUSCRIPT))
    for element in clean._element.findall(".//" + qn("w:highlight")):
        element.getparent().remove(element)
    clean.save(str(PKG / "02_Manuscript_clean.docx"))


def main() -> None:
    eac = recompute_eac_os60_cox()
    revise_word_tables_and_language(eac)
    clean_supplementary_metadata()
    patch_endpoint_tsv()
    print("Final scientific and language revisions applied")


if __name__ == "__main__":
    main()
