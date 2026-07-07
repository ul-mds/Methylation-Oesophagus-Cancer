"""
Insert each supplementary figure PNG immediately *above* its Supplementary Fig. S* caption.

Only unique (non–main-text) panels: mediation summaries (former 07/08), numbered S1–S2.

If an image paragraph already sits above the caption, its picture is replaced so regenerated PNGs refresh the Word file.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from lxml import etree

DOC_NAME = "Supplementary/Supplementary_Figures.docx"

SPECS: list[tuple[str, str, float]] = [
    ("Supplementary Fig. S1.", "S1_mediation_delta_volcano.png", 6.2),
    ("Supplementary Fig. S2.", "S2_top_histology_specific_mediation.png", 6.2),
]


def _p_has_embedded_image(p_el) -> bool:
    return b'r:embed="' in etree.tostring(p_el)


def _ensure_image_above_caption(caption_para: Paragraph, png: Path, width_in: float) -> None:
    if not png.is_file():
        print(f"Missing plot, skip: {png}")
        return
    cap_el = caption_para._p
    prev = cap_el.getprevious()
    if prev is not None and prev.tag.endswith("}p") and _p_has_embedded_image(prev):
        pic_p = Paragraph(prev, caption_para._parent)
        for r in list(pic_p.runs):
            r._r.getparent().remove(r._r)
        pic_p.add_run().add_picture(str(png), width=Inches(width_in))
        print(f"Refreshed above caption: {png.name}")
        return

    new_el = OxmlElement("w:p")
    cap_el.addprevious(new_el)
    pic_p = Paragraph(new_el, caption_para._parent)
    pic_p.add_run().add_picture(str(png), width=Inches(width_in))
    print(f"Embedded above caption: {png.name}")


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    plots = root / "Supplementary" / "Figures"
    path = root / DOC_NAME
    doc = Document(str(path))

    for prefix, fname, w in SPECS:
        png = plots / fname
        found = False
        for p in doc.paragraphs:
            if (p.text or "").strip().startswith(prefix):
                _ensure_image_above_caption(p, png, w)
                found = True
                break
        if not found:
            print(f"Caption not found for {prefix}")

    doc.save(str(path))
    print("Saved:", path)


if __name__ == "__main__":
    main()
