"""
Insert thematic subsection titles (Heading 3) into the Results section, fix
mis-styled narrative paragraphs, and tighten the opening roadmap so the section
reads as one continuous story organized into labeled substudies.

Idempotent: skips a title if that exact paragraph text already exists.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


def _insert_paragraph_before(paragraph: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._element.addprevious(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    if style is not None:
        new_para.style = style
    return new_para


def _has_exact_title(doc: Document, title: str) -> bool:
    return any((p.text or "").strip() == title for p in doc.paragraphs)


def _find_startswith(doc: Document, prefix: str) -> Paragraph:
    for p in doc.paragraphs:
        if (p.text or "").strip().startswith(prefix):
            return p
    raise ValueError(f"No paragraph starting with: {prefix!r}")


def main():
    path = Path(__file__).resolve().parents[1] / "TCGA_ESCA_ESCC_landmark60_full_scientific_manuscript.docx"
    doc = Document(str(path))

    # Opening paragraph: keep one narrative block; note substudy headers (English manuscript).
    p_results = _find_startswith(doc, "The TCGA esophageal carcinoma cohort combines")
    roadmap = (
        "The TCGA esophageal carcinoma cohort combines esophageal adenocarcinoma (EAC) and esophageal squamous "
        "cell carcinoma (ESCC), yet integrated genomics has established that these histologies represent distinct "
        "molecular diseases rather than interchangeable subtypes of a single entity. A practical consequence—often "
        "unstated in prognostic omics studies—is that risk models optimized in one histology may lack external "
        "validity when applied to the other, even within the same data resource. We therefore structured the "
        "Results around a single logical arc: (i) quantify cross-histology transferability of multi-feature Cox "
        "scores under a 60-month landmark overall survival endpoint (Figure 1; Table 1); (ii) ask whether clinically "
        "extreme courses (early death among decedents) exhibit histology-specific DNA methylation programs "
        "(Figures 2–3; Tables 2–5); (iii) relate prioritized probes to interpretable survival separation at the "
        "same clinical horizon (Figure 4; Table 6); and (iv) situate these epigenetic patterns alongside exploratory "
        "somatic analyses in which histology may modify mutation–outcome associations and downstream expression "
        "modules (Figures 5–6; Table 7). For readability, the empirical blocks below follow this same order under "
        "four subsection headings, with each table and figure placed in the substudy to which it primarily belongs."
    )
    p_results.text = roadmap
    p_results.style = "Normal"

    titles = [
        (
            "We trained penalized Cox models on high-dimensional expression–methylation features",
            "Histology-specific prognostic models and cross-lineage transfer",
        ),
        (
            "If multi-omic risk scores fail to generalize across histologies",
            "Fast-death methylation programs by histology",
        ),
        (
            "Discovery among decedents does not, by itself, establish that the same probes stratify survival",
            "From methylation discovery to landmark survival separation",
        ),
        (
            "Finally, we examined somatic alterations in relation to OS60",
            "Exploratory somatic context: expression modules and mutation×histology interactions",
        ),
    ]

    for anchor_prefix, title in titles:
        if _has_exact_title(doc, title):
            continue
        anchor = _find_startswith(doc, anchor_prefix)
        _insert_paragraph_before(anchor, title, style="Heading 3")

    # Body paragraphs incorrectly tagged as Heading 2 (legacy styling).
    for prefix in (
        "Genome-wide contrasts revealed divergent fast-death methylation",
        "In EAC, univariate Cox models on OS60 identified",
    ):
        try:
            p = _find_startswith(doc, prefix)
        except ValueError:
            continue
        if p.style and p.style.name == "Heading 2":
            p.style = "Normal"

    doc.save(str(path))
    print("Structured Results subsections:", path)


if __name__ == "__main__":
    main()
