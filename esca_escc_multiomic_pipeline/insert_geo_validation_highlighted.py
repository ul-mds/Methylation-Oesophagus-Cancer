"""
Insert GEO external-validation text and tables into the manuscript .docx.
All new content is highlighted in yellow. No other edits.
"""
from __future__ import annotations

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


DOC_PATH = (
    "/Users/senol/Desktop/AML/Esophageal Carcinoma (TCGA, Firehose Legacy)/"
    "TCGA_ESCA_ESCC_landmark60_full_scientific_manuscript.docx"
)


def yellow_run(paragraph, text: str) -> None:
    r = paragraph.add_run(text)
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def insert_paragraph_after(paragraph: Paragraph, style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        try:
            new_para.style = style
        except Exception:
            pass
    return new_para


def paragraph_after_element(oxml_element, body_parent) -> Paragraph:
    new_p = OxmlElement("w:p")
    oxml_element.addnext(new_p)
    return Paragraph(new_p, body_parent)


def insert_table_after(document: Document, paragraph: Paragraph, rows: int, cols: int):
    table = document.add_table(rows=rows, cols=cols)
    try:
        table.style = "Table Grid"
    except Exception:
        pass
    paragraph._p.addnext(table._tbl)
    return table


def fill_table_yellow(table, data: list[list[str]]) -> None:
    for i, row_data in enumerate(data):
        row = table.rows[i]
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            yellow_run(p, val)


def highlight_table(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def find_paragraph(doc: Document, prefix: str):
    for p in doc.paragraphs:
        if (p.text or "").startswith(prefix):
            return p
    return None


def find_paragraph_index(doc: Document, prefix: str) -> int:
    for i, p in enumerate(doc.paragraphs):
        if (p.text or "").startswith(prefix):
            return i
    return -1


def main() -> None:
    doc = Document(DOC_PATH)
    body_parent = doc.paragraphs[0]._parent

    h2 = "Heading 2"
    normal = "Normal"

    # --- Abstract ---
    abs_p = find_paragraph(doc, "Esophageal adenocarcinoma (EAC) and esophageal squamous")
    if not abs_p:
        raise SystemExit("Abstract body paragraph not found")
    p = insert_paragraph_after(abs_p, normal)
    yellow_run(
        p,
        "As modular external checks, we additionally examined public GEO data: paired tumor–normal "
        "expression in GSE161533 (n = 28 ESCC patients) against a pre-specified union of TCGA ESCC "
        "genes (60 Cox expression features ∪ top-80 mediation-priority genes; 110 unique genes), "
        "and 60-month landmark penalized Cox models on lncRNA arrays in GSE53624 (training) versus "
        "GSE53622 (independent). In GSE161533, two union genes (BCAT1, SDF4) ranked within the top "
        "1200 probe-derived tumor–normal contrasts (mean log2 fold changes +1.53 and +1.33). "
        "In GSE53624→GSE53622, the trained model’s C-index was ~0.76 in-cohort versus ~0.51 on "
        "independent hold-out patients, indicating limited cross-cohort prognostic transport for "
        "this lncRNA feature space.",
    )

    # --- Methods (before Results heading) ---
    r_idx = find_paragraph_index(doc, "Results")
    if r_idx < 0:
        raise SystemExit("Results heading not found")
    while r_idx > 0 and not (doc.paragraphs[r_idx - 1].text or "").strip():
        r_idx -= 1
    anchor = doc.paragraphs[r_idx - 1]

    m0 = insert_paragraph_after(anchor, h2)
    yellow_run(m0, "External validation (GEO cohorts)")

    m1 = insert_paragraph_after(m0, normal)
    yellow_run(
        m1,
        "GSE161533: We downloaded the GEO series matrix (GPL570; 28 ESCC patients with matched normal, "
        "peritumoral, and tumor samples). For tumor-centric concordance, we computed paired tumor versus "
        "matched normal log2 ratios (pseudocount +1), averaged across the 28 pairs. Affymetrix probes "
        "were mapped to HGNC symbols using the mygene.info batch API (scopes: reporter; human). "
        "Gene-level summaries used the probe with the largest absolute mean log2 fold change per symbol. "
        "Overlap with TCGA was pre-specified without using GSE161533 outcomes: the union of (i) all "
        "protein-coding genes entering the ESCC-only penalized Cox model as expression features (n = 60) "
        "and (ii) the top 80 genes from the ESCC mediation scan output (deduplicated union n = 110). "
        "We report genes from this union appearing among genes linked to the top 1200 probes ranked by "
        "absolute mean tumor–normal log2 fold change.",
    )

    m2 = insert_paragraph_after(m1, normal)
    yellow_run(
        m2,
        "GSE53624 / GSE53622: From GEO supplementary tables and series matrices we assembled lncRNA "
        "expression and clinical overall survival. We applied the same 60-month landmark censoring as "
        "in the primary analysis. We selected the 60 highest-variance lncRNA features in GSE53624 only, "
        "fit a penalized Cox model (penalizer = 0.1, l1_ratio = 0.5; lifelines CoxPHFitter), and "
        "evaluated Harrell’s C-index in GSE53624 (in-sample) and when scoring GSE53622 without refitting "
        "coefficients (hold-out). We additionally report in-cohort C-index when the same model class is "
        "fit within GSE53622 alone for context.",
    )

    # --- Results (before Discussion) ---
    d_idx = find_paragraph_index(doc, "Discussion")
    if d_idx < 0:
        raise SystemExit("Discussion not found")
    while d_idx > 0 and not (doc.paragraphs[d_idx - 1].text or "").strip():
        d_idx -= 1
    anchor2 = doc.paragraphs[d_idx - 1]

    r0 = insert_paragraph_after(anchor2, h2)
    yellow_run(r0, "External GEO cohort analyses")

    r1 = insert_paragraph_after(r0, normal)
    yellow_run(
        r1,
        "GSE161533 (paired tumor–normal mRNA): Among the pre-specified 110-gene TCGA union, two genes "
        "overlapped the gene-level representation of the top 1200 ranked probes by |mean log2 "
        "(tumor/normal)|: BCAT1 (mean log2 fold change +1.534; probe 226517_at) and SDF4 (+1.329; "
        "221972_s_at). Partitioning by source list, 1/60 Cox expression genes and 1/80 mediation-priority "
        "genes were captured in this top-1200 screen. This modest overlap is expected because TCGA features "
        "were selected for prognosis and methylation–expression coupling, whereas GSE161533 contrasts "
        "tumor versus normal mucosa without longitudinal outcomes in the GEO matrix.",
    )

    r2 = insert_paragraph_after(r1, normal)
    yellow_run(
        r2,
        "GSE53624 / GSE53622 (lncRNA; 60-month landmark): Training in GSE53624 (n = 119; 72 deaths "
        "within 60 months) yielded C-index 0.764. Applying the GSE53624-trained model to independent "
        "GSE53622 (n = 60; 33 events ≤60 months) gave hold-out C-index 0.505. Fitting the same penalized "
        "Cox framework within GSE53622 alone yielded C-index 0.812, suggesting cohort-specific "
        "prognostic structure on this platform and limited signature transport across these series.",
    )

    tc = insert_paragraph_after(r2, normal)
    yellow_run(
        tc,
        "Supplementary Table S8. GSE161533: TCGA pre-specified ESCC genes overlapping top-1200 tumor–normal contrasts.",
    )

    t1 = insert_table_after(doc, tc, rows=3, cols=3)
    fill_table_yellow(
        t1,
        [
            ["Gene", "Mean log2(tumor/normal)", "Representative probe"],
            ["BCAT1", "+1.534", "226517_at"],
            ["SDF4", "+1.329", "221972_s_at"],
        ],
    )
    highlight_table(t1)

    n1 = paragraph_after_element(t1._tbl, body_parent)
    yellow_run(
        n1,
        "Note. Top N = 1200 probes by |mean paired log2 fold change|; n = 28 paired patients; "
        "mygene.info probe-to-symbol mapping with per-gene max-|Δ| probe retention. Union list: "
        "110 genes (60 Cox expr ∪ 80 mediation top genes, deduplicated).",
    )

    tcap = insert_paragraph_after(n1, normal)
    yellow_run(
        tcap,
        "Supplementary Table S9. GSE53624 / GSE53622: landmark (60-month) penalized Cox on lncRNA features.",
    )

    t2 = insert_table_after(doc, tcap, rows=4, cols=5)
    fill_table_yellow(
        t2,
        [
            ["Analysis", "Cohort", "n", "Events ≤60 mo", "C-index"],
            ["Training fit", "GSE53624", "119", "72", "0.764"],
            ["External scoring", "GSE53622", "60", "33", "0.505"],
            ["Internal fit (contextual)", "GSE53622", "60", "33", "0.812"],
        ],
    )
    highlight_table(t2)

    n2 = paragraph_after_element(t2._tbl, body_parent)
    yellow_run(
        n2,
        "Note. Sixty highest-variance lncRNAs selected in GSE53624 only; penalizer = 0.1, l1_ratio = 0.5. "
        "Platform GPL18109 (custom lncRNA array), not TCGA mRNA or HM450.",
    )

    # --- Discussion ---
    lim_p = None
    for p in doc.paragraphs:
        if (p.text or "").startswith("Limitations include reliance"):
            lim_p = p
            break
    if lim_p:
        d1 = insert_paragraph_after(lim_p, normal)
        yellow_run(
            d1,
            "External GEO modules complement but do not replicate the primary TCGA design: GSE161533 "
            "supports limited independent tumor–normal expression concordance for pre-specified TCGA genes "
            "without enabling mediation or multivariable prognosis replication (outcomes absent in the "
            "series matrix). GSE53624/GSE53622 show that the landmark penalized Cox workflow applies to "
            "external RNA data yet underscore weak cross-cohort transport for the chosen lncRNA feature "
            "space, reinforcing that histology-matched, assay-matched prospective validation remains the "
            "priority for clinical translation.",
        )

    doc.save(DOC_PATH)
    print("Saved:", DOC_PATH)


if __name__ == "__main__":
    main()
