"""Replace the Results roadmap paragraph with a concise, citation-forward lead (run once after reorganize)."""
from __future__ import annotations

from pathlib import Path

from docx import Document

DOC_NAME = "TCGA_ESCA_ESCC_landmark60_full_scientific_manuscript.docx"

NEW_LEAD = (
    "We analysed the TCGA esophageal carcinoma (ESCA) Firehose Legacy cohort as described in the Methods and "
    "organised empirical outputs in four blocks. First, we quantified cross-histology transfer of penalized Cox "
    "risk scores built from RNA and HM450 features (Fig. 1a; Table 1). Second, we contrasted genome-wide HM450 "
    "methylation between fast-death and other decedents within each histology (Fig. 2a,b; Fig. 3a,b; Tables 2–5). "
    "Third, we linked selected probes to Kaplan–Meier separation and EAC univariate methylation–outcome screening "
    "(Fig. 4a–f; Table 6). Fourth, we report exploratory mutation×histology interactions and mutation-associated "
    "expression-module shifts (Fig. 5a; Fig. 6a; Table 7). Subsections below follow this order. Recent integrative "
    "work continues to motivate histology-resolved methylation reporting [5,8–10]."
)


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    path = root / DOC_NAME
    doc = Document(str(path))
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith("The TCGA esophageal carcinoma cohort combines"):
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            p.add_run(NEW_LEAD)
            break
    else:
        raise SystemExit("Results lead paragraph not found.")
    doc.save(str(path))
    print("Patched Results opening:", path)


if __name__ == "__main__":
    main()
