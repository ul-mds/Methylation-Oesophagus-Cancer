"""Insert five Keywords after the abstract Conclusions line (idempotent)."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


def _txt(p) -> str:
    return "".join(t.text or "" for t in p._p.iter(qn("w:t"))).strip()


def main():
    path = Path(__file__).resolve().parents[1] / "TCGA_ESCA_ESCC_landmark60_full_scientific_manuscript.docx"
    doc = Document(str(path))

    for p in doc.paragraphs:
        if _txt(p).startswith("Keywords:"):
            print("Keywords already present:", path)
            return

    insert_after = None
    for p in doc.paragraphs:
        if _txt(p).startswith("Conclusions: Prognostic signal in TCGA esophageal carcinoma"):
            insert_after = p
            break
    if insert_after is None:
        raise RuntimeError("Could not find abstract Conclusions paragraph.")

    line = (
        "Keywords: esophageal carcinoma; esophageal adenocarcinoma; esophageal squamous cell carcinoma; "
        "DNA methylation; prognostic model transferability"
    )
    new_p = OxmlElement("w:p")
    insert_after._element.addnext(new_p)
    Paragraph(new_p, insert_after._parent).add_run(line)

    doc.save(str(path))
    print("Keywords inserted:", path)


if __name__ == "__main__":
    main()
