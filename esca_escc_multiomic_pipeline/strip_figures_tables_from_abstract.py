"""Remove Figure/Table pointers from Background/Methods/Results/Conclusions abstract blocks."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def _txt(p) -> str:
    return "".join(t.text or "" for t in p._p.iter(qn("w:t"))).strip()


def main():
    path = Path(__file__).resolve().parents[1] / "TCGA_ESCA_ESCC_landmark60_full_scientific_manuscript.docx"
    doc = Document(str(path))

    updates = {
        "Methods: We analyzed the TCGA ESCA Firehose Legacy study package, assembling a sample-level dataset with clinical outcomes, HM450 methylation, and RNA expression. Overall survival (OS) was evaluated using a 60-month landmark (follow-up capped at 60 months; deaths after 60 months censored). We trained histology-specific Cox models (EAC-only; ESCC-only) and assessed both within-histology performance and cross-histology transferability using the concordance index. To connect epigenetics to clinically aggressive disease, we defined a fast-death phenotype among decedents within each histology and identified differential methylation in fast-death versus others, followed by 60-month Kaplan–Meier visualizations for selected loci. We additionally performed exploratory mutation×histology interaction screening and mechanistic follow-up using expression-derived module scores. Primary numeric summaries appear in Tables 1–7; integrated graphics appear in Figures 1–6.": (
            "Methods: We analyzed the TCGA ESCA Firehose Legacy study package, assembling a sample-level dataset with clinical outcomes, HM450 methylation, and RNA expression. Overall survival (OS) was evaluated using a 60-month landmark (follow-up capped at 60 months; deaths after 60 months censored). We trained histology-specific Cox models (EAC-only; ESCC-only) and assessed both within-histology performance and cross-histology transferability using the concordance index. To connect epigenetics to clinically aggressive disease, we defined a fast-death phenotype among decedents within each histology and identified differential methylation in fast-death versus others, followed by 60-month Kaplan–Meier visualizations for selected loci. We additionally performed exploratory mutation×histology interaction screening and mechanistic follow-up using expression-derived module scores."
        ),
        "Results: Histology-specific models discriminated well within histology (EAC-only C-index=0.894; ESCC-only C-index=0.901) but failed to transfer across histology (EAC→ESCC C-index=0.469; ESCC→EAC C-index=0.450), as summarized graphically in Figure 1 and in full in Table 1. A pooled model with histology-stratified baseline improved overall performance (C-index=0.803; Table 1), showing that pooling requires explicit histology handling. Fast-death methylation differed sharply by histology: genome-wide contrasts (Figure 2) and ranked hypo-/hypermethylation bar panels (Figure 3) aligned with Tables 2–5—EAC fast-death featured hypomethylation at loci including SIGLEC14, EPHB1, ZNF214, and RAB11FIP3 (Table 2) with complementary hypermethylated ranks (Table 4), whereas ESCC fast-death was dominated by strong SHISA3 hypomethylation (Table 3) alongside C7orf13, PRKG1, and SEMA3E, with hypermethylated ranks in Table 5. Methylation–OS associations in EAC (Table 6) and survival separation by median methylation across histologies are shown in Figure 4. Finally, mutation-linked module shifts (Figure 5; Table 7) and the nominal mutation×histology interaction landscape (Figure 6) provide exploratory downstream context.": (
            "Results: Histology-specific models discriminated well within histology (EAC-only C-index=0.894; ESCC-only C-index=0.901) but failed to transfer across histology (EAC→ESCC C-index=0.469; ESCC→EAC C-index=0.450). A pooled model with histology-stratified baseline improved overall performance (C-index=0.803), showing that pooling requires explicit histology handling. Fast-death methylation differed sharply by histology: genome-wide contrasts and ranked hypo-/hypermethylation bar summaries were concordant with locus-level discovery—EAC fast-death featured hypomethylation at loci including SIGLEC14, EPHB1, ZNF214, and RAB11FIP3 with complementary hypermethylated ranks, whereas ESCC fast-death was dominated by strong SHISA3 hypomethylation alongside C7orf13, PRKG1, and SEMA3E, with distinct hypermethylated ranks. Methylation–OS associations in EAC and survival separation by median methylation across histologies supported interpretable risk strata. Mutation-linked module shifts and the nominal mutation×histology interaction landscape provide exploratory downstream context."
        ),
        "Conclusions: Prognostic signal in TCGA esophageal carcinoma is strongly histology-specific and largely non-transferable between EAC and ESCC (Figure 1; Table 1). Histology-resolved fast-death epigenetic programs (Figures 2–3; Tables 2–5) and 60-month survival stratification (Figure 4; Table 6) provide a clinically anchored view of aggressive biology within a bounded horizon, while mutation-layer follow-up (Figures 5–6; Table 7) motivates external validation.": (
            "Conclusions: Prognostic signal in TCGA esophageal carcinoma is strongly histology-specific and largely non-transferable between EAC and ESCC. Histology-resolved fast-death epigenetic programs and 60-month survival stratification provide a clinically anchored view of aggressive biology within a bounded horizon, while mutation-layer follow-up motivates external validation."
        ),
    }

    for old, new in updates.items():
        done = False
        for p in doc.paragraphs:
            if _txt(p) == old:
                p.text = new
                done = True
                break
        if not done:
            for p in doc.paragraphs:
                if _txt(p) == new:
                    done = True
                    break
            if not done:
                raise RuntimeError("Abstract paragraph not found or already manually edited.")

    doc.save(str(path))
    print("Stripped Figure/Table refs from abstract:", path)


if __name__ == "__main__":
    main()
