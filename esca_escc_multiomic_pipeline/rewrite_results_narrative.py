"""
Rewrite the Results section as a cohesive scientific narrative with smooth
transitions among findings, figures, and tables.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document


def main():
    study = Path(__file__).resolve().parents[1]
    path = study / "TCGA_ESCA_ESCC_landmark60_full_scientific_manuscript.docx"
    doc = Document(str(path))

    updates: dict[int, str] = {
        24: "Results",
        25: (
            "The TCGA esophageal carcinoma cohort combines esophageal adenocarcinoma (EAC) and esophageal squamous "
            "cell carcinoma (ESCC), yet integrated genomics has established that these histologies represent distinct "
            "molecular diseases rather than interchangeable subtypes of a single entity. A practical consequence—often "
            "unstated in prognostic omics studies—is that risk models optimized in one histology may lack external "
            "validity when applied to the other, even within the same data resource. We therefore structured the "
            "Results around a single logical arc: (i) quantify cross-histology transferability of multi-feature Cox "
            "scores under a 60-month landmark overall survival endpoint (Figure 1; Table 1); (ii) ask whether clinically "
            "extreme courses (early death among decedents) exhibit histology-specific DNA methylation programs "
            "(Figures 2–3; Tables 2–5); (iii) relate prioritized probes to interpretable survival separation at the "
            "same clinical horizon (Figure 4; Table 6); and (iv) situate these epigenetic patterns alongside exploratory "
            "somatic analyses in which histology may modify mutation–outcome associations and downstream expression "
            "modules (Figures 5–6; Table 7)."
        ),
        26: (
            "We trained penalized Cox models on high-dimensional expression–methylation features within each histology "
            "and evaluated the concordance (C-index) of predicted risk on 60-month landmark OS. Within histology, both "
            "lineage-specific models achieved strong discrimination (EAC-only model evaluated in EAC: C-index=0.894; "
            "ESCC-only model evaluated in ESCC: C-index=0.901), indicating that the chosen feature space captures "
            "prognostic structure when training and testing contexts match. Cross-application told a different story: "
            "the EAC-trained model yielded C-index=0.469 in ESCC, and the ESCC-trained model C-index=0.450 in EAC—"
            "both near the null value of 0.5 for rank-based survival concordance. Thus, the ordering of predicted risk "
            "learned in one histology does not translate into meaningful patient ranking in the other, despite shared "
            "assay platforms and harmonized preprocessing. A pooled Cox model with histology-stratified baseline hazards "
            "achieved C-index=0.803 on the combined cohort, demonstrating that prognostic information is not absent when "
            "histologies are combined, but that it must be embedded in a histology-aware baseline rather than assumed to "
            "be carried by a single shared coefficient vector. Figure 1 visualizes this contrast between within-lineage "
            "performance, cross-lineage collapse, and the stratified-pooled comparator; Table 1 reports the full "
            "model×evaluation matrix for transparency and secondary reanalysis."
        ),
        27: (
            "Table 1. Concordance index (C-index) for histology-specific and pooled models on 60-month landmark overall "
            "survival."
        ),
        28: (
            "If multi-omic risk scores fail to generalize across histologies, it is plausible that the epigenetic "
            "correlates of aggressive disease are likewise lineage-specific rather than pan-esophageal. We therefore "
            "examined DNA methylation in a phenotype anchored to clinical outcome: among patients who died, we defined "
            "“fast death” as membership in the shortest survival tail within each histology (top 20% of observed deaths, "
            "minimum n=10), and contrasted HM450 β-values against all other samples of the same histology. This design "
            "emphasizes epigenetic features enriched in the most rapidly fatal trajectories rather than generic "
            "tumor–normal hypermethylation, aligning statistical discovery with a severe clinical reference class."
        ),
        29: (
            "Genome-wide contrasts revealed divergent fast-death methylation architectures between EAC and ESCC. In EAC, "
            "the volcano plot (Figure 2, left) shows a broad excess of statistically supported negative Δβ (hypomethylation "
            "in fast-death versus others), suggesting widespread relaxation of methylation at many loci in the most "
            "aggressive adenocarcinoma decedents. In ESCC, the companion volcano (Figure 2, right) is dominated by a "
            "focal, extreme hypomethylation signal at SHISA3, with additional strong negative Δβ at loci annotated to "
            "C7orf13, PRKG1, and SEMA3E—indicating that the strongest single-probe epigenetic separation in squamous "
            "fast-death is both quantitatively sharper and different in gene emphasis than in adenocarcinoma."
        ),
        30: (
            "Figure 3 compresses these landscapes into side-by-side bar summaries of leading hypo- and hypermethylated "
            "probes, making directionality and locus composition directly comparable at a glance. Tables 2 and 3 "
            "enumerate ranked hypomethylated loci in EAC and ESCC, respectively (including SIGLEC14, EPHB1, ZNF214, and "
            "RAB11FIP3 among EAC leaders, and SHISA3-driven ranks in ESCC), whereas Tables 4 and 5 list the "
            "corresponding hypermethylated ranks that populate the upward-oriented bars in Figure 3. Together, Figures 2–3 "
            "and Tables 2–5 support a coherent inference: “fast-death” methylation is not a universal esophageal program "
            "but a histology-resolved epigenetic signature with distinct breadth and focal genes."
        ),
        31: (
            "Table 2. EAC/ESCA: top hypomethylated HM450 probes in fast-death versus other samples (Benjamini–Hochberg "
            "FDR, q_diff)."
        ),
        33: (
            "Table 3. ESCC: top hypomethylated HM450 probes in fast-death versus other samples (Benjamini–Hochberg FDR, "
            "q_diff)."
        ),
        35: (
            "Table 4. EAC/ESCA: top hypermethylated HM450 probes in fast-death versus other samples "
            "(Benjamini–Hochberg FDR, q_diff)."
        ),
        37: (
            "Table 5. ESCC: top hypermethylated HM450 probes in fast-death versus other samples (Benjamini–Hochberg FDR, "
            "q_diff)."
        ),
        38: (
            "Discovery among decedents does not, by itself, establish that the same probes stratify survival in the full "
            "case mix—including long survivors and censored observations—at a clinically reportable horizon. We therefore "
            "linked methylation to 60-month landmark OS using complementary selection rules by histology."
        ),
        39: (
            "In EAC, univariate Cox models on OS60 identified CDC14C, VNN3, and GXYLT2 among the strongest continuous "
            "methylation–outcome associations (Table 6). Kaplan–Meier curves after median dichotomization (Figure 4, upper "
            "row) show separated survival trajectories within the capped follow-up window, connecting the Cox ranking in "
            "Table 6 to clinically interpretable risk strata. For ESCC, we prioritized probes with the strongest "
            "fast-death differential evidence (q_diff) that remained amenable to stable Kaplan–Meier visualization—"
            "SHISA3, PRKG1, and C7orf13—so the lower row of Figure 4 extends the epigenetic narrative from Figures 2–3 "
            "into a population-level time-to-event display. The juxtaposition of rows underscores a methodological point "
            "with biological implications: optimal probes for survival graphics are not symmetric across histologies when "
            "one lineage is selected by Cox screening (EAC) and the other by fast-death extremity (ESCC), consistent with "
            "the non-transferable prognostic geometry documented in our opening analysis."
        ),
        40: "",
        41: (
            "Table 6. EAC/ESCA: univariate Cox associations between continuous probe methylation and 60-month landmark "
            "overall survival (selected top probes by q)."
        ),
        42: (
            "Finally, we examined somatic alterations in relation to OS60 and to histology-stratified transcriptional "
            "context—an exploratory layer meant to contextualize the epigenetic findings rather than to assert definitive "
            "causal interactions."
        ),
        43: "",
        44: (
            "Because mutation×histology interaction tests are event-limited in moderate samples, we paired "
            "histology-stratified transcriptional context with a nominal interaction discovery map. Figure 5 "
            "summarizes differences in expression-derived module scores between mutant and wild-type tumors "
            "(Δ=mutant−wild-type), including histology-partitioned movement on EMT and immune/interferon axes "
            "for selected alterations. Table 7 lists the strongest mutation-to-module or mutation-to-expression "
            "associations underlying that heatmap. Figure 6 then presents the interaction landscape (Cox β for "
            "the product term versus −log10 p): labeled genes highlight where histology appears to modify the "
            "mutation–survival slope and warrant independent replication under larger event counts. Read "
            "sequentially, Figure 5 and Table 7 offer testable transcriptional hypotheses integrable with the "
            "methylation programs above, while Figure 6 prioritizes genomic candidates for follow-up."
        ),
        45: (
            "Table 7. Exploratory mechanistic follow-up: strongest associations between somatic mutation and expression "
            "or module scores, stratified by histology (ranked by p-value)."
        ),
        47: (
            "Figures. The panels below follow the same sequence as the Results narrative: model transferability (Figure 1), "
            "fast-death methylation contrasts (Figures 2–3), landmark survival stratification (Figure 4), and exploratory "
            "mutation context (Figures 5–6)."
        ),
    }

    for idx, text in updates.items():
        if idx >= len(doc.paragraphs):
            raise IndexError(f"Paragraph {idx} out of range ({len(doc.paragraphs)})")
        doc.paragraphs[idx].text = text

    doc.save(str(path))
    print("Results narrative rewritten:", path)


if __name__ == "__main__":
    main()
